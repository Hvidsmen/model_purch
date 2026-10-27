from io import BytesIO
from unittest.mock import patch
from django.contrib.auth.models import User,Group
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook
from .access import ANALYST,LOGISTICIAN,role_for,ensure_roles
from .models import UserProjectAccess,GroupProjectAccess
from .purchases_access import READ_ROUTES
from model_purch import test_goods_editor
from model_purch.models import PGGoods,Purch,AlgorithmRun


class PurchaseAnalystTests(TestCase):
    def setUp(self):
        test_goods_editor.GoodsEditorTests.setUp(self)
        ensure_roles()
        self.user=User.objects.create_user('purchase-viewer')
        UserProjectAccess.objects.create(user=self.user,project='purchases',role=ANALYST)
        self.client.force_login(self.user)
        self.purchase=Purch.objects.create(name='Read supplier',lag_income=30,lage_make=5)
        self.run=AlgorithmRun.objects.create(scenario=self.scenario,status='completed')

    def test_personal_and_group_grants_are_independent_of_dashboards(self):
        self.assertEqual(role_for(self.user,'purchases'),ANALYST)
        self.assertIsNone(role_for(self.user,'dashboards'))
        UserProjectAccess.objects.filter(user=self.user).delete()
        self.user=User.objects.get(pk=self.user.pk)
        group=Group.objects.create(name='Purchase analysts')
        GroupProjectAccess.objects.create(group=group,project='purchases',role=ANALYST)
        self.user.groups.add(group)
        self.assertEqual(role_for(self.user,'purchases'),ANALYST)
        UserProjectAccess.objects.create(user=self.user,project='purchases',role='')
        self.assertIsNone(role_for(self.user,'purchases'))
        legacy=User.objects.create_user('legacy-dashboard-analyst')
        legacy.groups.add(Group.objects.get(name=ANALYST))
        self.assertEqual(role_for(legacy,'dashboards'),ANALYST)
        self.assertIsNone(role_for(legacy,'purchases'))

    def test_every_mutation_and_non_read_get_route_is_denied(self):
        from model_purch.urls import urlpatterns
        before=PGGoods.objects.values().get(pk=self.good.pk)
        with patch('model_purch.views.connect_database') as sql,patch('model_purch.views.sync_purch_data') as sync:
            for route in urlpatterns:
                name=route.name
                kwargs={key:self.run.pk if key=='run_id' else self.scenario.pk for key in route.pattern.converters}
                url=reverse(name,kwargs=kwargs)
                for method in ('post','put','patch','delete'):
                    with self.subTest(route=name,method=method):
                        self.assertEqual(getattr(self.client,method)(url,{'scenario':self.scenario.pk}).status_code,403)
                if name not in READ_ROUTES:
                    with self.subTest(route=name,method='get'):
                        self.assertEqual(self.client.get(url).status_code,403)
            sql.assert_not_called();sync.assert_not_called()
        self.assertEqual(PGGoods.objects.values().get(pk=self.good.pk),before)

    def test_read_pages_and_status_are_available(self):
        cases=[('scenario_list',[]),('scenario_archive',[]),('scenario_edit',[self.scenario.pk]),
            ('pggoods_list',[]),('edit_pggoods',[self.good.pk]),('purch_list',[]),('purch_edit',[self.purchase.pk]),
            ('goods_groups',[]),('goods_plan_parameters',[self.scenario.pk]),('results_page',[]),('stock_chart_page',[]),('get_algorithm_statuэ',[self.run.pk])]
        for name,args in cases:
            with self.subTest(route=name):
                page=self.client.get(reverse(name,args=args),{'scenario':self.scenario.pk})
                self.assertEqual(page.status_code,200)
        for name in ('portal_users','portal_user_groups','portal_updates','global_coeff_admin_motivation','dashboards_index'):
            self.assertEqual(self.client.get(reverse(name)).status_code,403)

    def test_excel_recalculates_in_memory_without_saving_any_business_data(self):
        PGGoods.objects.filter(pk=self.good.pk).update(exw_usd=120,ddp_usd=777,plan_quantity=100,stock_quantity=25,stock_ddp_usd=80)
        before=PGGoods.objects.values().get(pk=self.good.pk)
        with CaptureQueriesContext(connection) as queries,patch('model_purch.services.pricing.reprice_goods',side_effect=AssertionError('No writes')):
            response=self.client.get(reverse('export_pggoods'),{'scenario':self.scenario.pk})
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(response.content));self.addCleanup(book.close)
        headers=[cell.value for cell in book.active[1]]
        self.assertEqual(book.active.cell(2,headers.index('DDP USD')+1).value,120)
        self.assertEqual(book.active.cell(2,headers.index('Ср. взв. DDP, USD')+1).value,110)
        self.assertEqual(PGGoods.objects.values().get(pk=self.good.pk),before)
        business_writes=[query['sql'] for query in queries if query['sql'].lstrip().upper().startswith(('INSERT','UPDATE','DELETE')) and 'django_session' not in query['sql']]
        self.assertEqual(business_writes,[])
        self.assertEqual(self.client.get(reverse('export_purchases')).status_code,200)

    def test_role_is_assignable_in_both_user_and_group_forms(self):
        from .user_views import UserForm
        from .group_views import ReportGroupForm
        admin=User.objects.get(username='portal-test-admin')
        for form in (UserForm(actor=admin),ReportGroupForm(actor=admin)):
            self.assertIn((ANALYST,ANALYST),form.fields['purchases_role'].choices)

    def test_saving_user_and_group_assignments_uses_new_role(self):
        from .user_views import UserForm
        from .group_views import ReportGroupForm
        admin=User.objects.get(username='portal-test-admin')
        form=UserForm({'username':self.user.username,'is_active':'on','portal_role':'','motivation_role':'','purchases_role':ANALYST,'dashboards_role':''},instance=self.user,actor=admin)
        self.assertTrue(form.is_valid(),form.errors)
        saved=form.save()
        self.assertEqual(role_for(User.objects.get(pk=saved.pk),'purchases'),ANALYST)
        member=User.objects.create_user('group-purchase-analyst')
        form=ReportGroupForm({'name':'Read procurement group','portal_role':'','motivation_role':'','purchases_role':ANALYST,'dashboards_role':'','members':[member.pk]},actor=admin)
        self.assertTrue(form.is_valid(),form.errors);form.save()
        self.assertEqual(role_for(User.objects.get(pk=member.pk),'purchases'),ANALYST)

    def test_legacy_dashboard_group_does_not_gain_procurement_access(self):
        from .models import PortalReportGroup
        from .project_catalog import group_roles
        group=Group.objects.create(name='Existing dashboard group')
        PortalReportGroup.objects.create(group=group,portal_role=ANALYST)
        self.assertEqual(group_roles(group),{'dashboards':ANALYST})
