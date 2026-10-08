from base.testing import authorize_test_case
from django.test import TestCase
from django.urls import reverse
from .models import ScenarioModel, KindPurch, PGGoods


class GoodsEditorTests(TestCase):
    def setUp(self):
        authorize_test_case(self)
        self.scenario = ScenarioModel.objects.create(name='Plan', date_start_plan='2026-01-01', date_end_plan='2026-12-31')
        kind = KindPurch.objects.create(name='Purchased')
        self.good = PGGoods.objects.create(scenario_plan=self.scenario, planning_group='Group', planning_sales='Sales',
            group_goods='Goods', kind_purch=kind, brand='Brand', purch='Supplier', volume=.01, exw_usd=0,
            ddp_usd=117.326646, kddp=1, stock_cnt_day=1, percent_stock_end=0)
        self.url = reverse('edit_pggoods', args=[self.good.pk]) + f'?scenario={self.scenario.pk}'
        self.data = {'scenario': self.scenario.pk, 'volume': '.01', 'exw_usd': '0', 'ddp_usd': '117.326646',
                     'kddp': '', 'stock_cnt_day': '1', 'percent_stock_end': '0'}

    def test_numeric_edit_accepts_zero_exw_and_preserves_unsubmitted_metadata(self):
        response = self.client.post(self.url, self.data)
        self.assertRedirects(response, reverse('pggoods_list') + f'?scenario={self.scenario.pk}')
        self.good.refresh_from_db()
        self.assertEqual((self.good.brand, self.good.purch, self.good.kddp), ('Brand', 'Supplier', 1))
        self.assertEqual(self.good.ddp_usd, 0)  # Zero EXW and no freight parameters.

    def test_brand_and_purchase_are_displayed_and_can_be_changed_or_cleared(self):
        page = self.client.get(self.url)
        self.assertContains(page, 'name="brand"')
        self.assertContains(page, 'name="purch"')
        self.assertContains(page, 'Плановая группа')
        for brand, purchase in [('New brand', 'New supplier'), ('', '')]:
            self.client.post(self.url, dict(self.data, brand=brand, purch=purchase))
            self.good.refresh_from_db()
            self.assertEqual(self.good.brand or '', brand)
            self.assertEqual(self.good.purch or '', purchase)

    def test_coefficient_is_calculated_server_side(self):
        self.client.post(self.url, dict(self.data, exw_usd='10', ddp_usd='15', kddp='999'))
        self.good.refresh_from_db()
        self.assertEqual(self.good.kddp, 1)  # Submitted DDP is ignored; without additional costs DDP = EXW.

    def test_invalid_volume_explains_error_and_preserves_database_values(self):
        response = self.client.post(self.url, dict(self.data, volume='0'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Изменения не сохранены')
        self.assertContains(response, 'Объём должен быть больше нуля')
        self.assertContains(response, 'href="#id_volume"')
        self.good.refresh_from_db()
        self.assertEqual(self.good.volume, .01)

    def test_container_volume_default_edit_and_invalid_values(self):
        self.assertEqual(self.good.container_volume, 65)
        self.assertContains(self.client.get(self.url), 'name="container_volume"')
        response = self.client.post(self.url, dict(self.data, container_volume='72.5'))
        self.assertEqual(response.status_code, 302)
        self.good.refresh_from_db()
        self.assertEqual(self.good.container_volume, 72.5)
        for value in ('0', '-1', 'nan', 'inf', ''):
            response = self.client.post(self.url, dict(self.data, container_volume=value))
            self.assertEqual(response.status_code, 200)
            self.good.refresh_from_db()
            self.assertEqual(self.good.container_volume, 72.5)

    def test_container_volume_bulk_copy_and_excel_round_trip(self):
        import io
        from openpyxl import load_workbook
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .services.copying import copy_goods
        bulk_url = reverse('bulk_update_pggoods') + f'?scenario={self.scenario.pk}'
        response = self.client.post(bulk_url, {'updates': [{'id': self.good.pk, 'container_volume': 70}]}, content_type='application/json')
        self.assertEqual(response.json()['updated'], 1)
        self.good.refresh_from_db()
        self.assertEqual(self.good.container_volume, 70)
        self.client.post(bulk_url, {'updates': [{'id': self.good.pk, 'container_volume': 0}]}, content_type='application/json')
        self.good.refresh_from_db()
        self.assertEqual(self.good.container_volume, 70)
        target = ScenarioModel.objects.create(name='Copy', date_start_plan='2027-01-01', date_end_plan='2027-12-31')
        copy_goods(self.scenario, target)
        self.assertEqual(PGGoods.objects.get(scenario_plan=target).container_volume, 70)
        response = self.client.get(reverse('export_pggoods'), {'scenario': self.scenario.pk})
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet.cell(1, 15).value, 'Объём контейнера, м³')
        sheet.cell(2, 15).value = 80
        stream = io.BytesIO()
        workbook.save(stream)
        self.client.post(reverse('import_from_excel'), {'scenario': self.scenario.pk, 'excel_file': SimpleUploadedFile('goods.xlsx', stream.getvalue())})
        self.good.refresh_from_db()
        self.assertEqual(self.good.container_volume, 80)
        sheet.delete_cols(15, 6)
        stream = io.BytesIO()
        workbook.save(stream)
        self.client.post(reverse('import_from_excel'), {'scenario': self.scenario.pk, 'excel_file': SimpleUploadedFile('legacy.xlsx', stream.getvalue())})
        self.good.refresh_from_db()
        self.assertEqual(self.good.container_volume, 80)

    def test_sql_export_adds_container_column_and_sends_custom_value(self):
        from unittest.mock import Mock, patch
        from .models import Purch
        Purch.objects.create(name='Supplier', lag_income=90, lage_make=30)
        self.good.container_volume = 72.5
        from decimal import Decimal
        self.good.duty_rate = Decimal('7.25')
        self.good.save()
        connection, cursor = Mock(), Mock()
        connection.cursor.return_value = cursor
        cursor.description = [('id',), ('planning_group',), ('scenario_name',)]
        cursor.fetchall.return_value = []
        cursor.fetchone.side_effect = [(501,), (701,)]
        with patch('model_purch.views.MS_SQL_CONN_STR', 'test'), patch('model_purch.views.pyodbc.connect', return_value=connection):
            self.client.post(reverse('export_scenario_to_sql', args=[self.scenario.pk]))
        calls = cursor.execute.call_args_list
        self.assertTrue(any("COL_LENGTH(N'portal.PGGoods', N'container_volume')" in call.args[0] for call in calls))
        update = next(call for call in calls if call.args[0].startswith('UPDATE [portal].[PGGoods]'))
        self.assertIn('[container_volume] = ?', update.args[0])
        self.assertIn(72.5, update.args[1:])
        self.assertIn('[duty_rate] = ?', update.args[0])
        self.assertIn('7.25', update.args[1:])
        self.assertTrue(any("COL_LENGTH(N'portal.PGGoods', N'duty_rate')" in call.args[0] for call in calls))
        connection.rollback.assert_not_called()

    def test_duty_rate_edit_bulk_copy_and_validation(self):
        from decimal import Decimal
        from .services.copying import copy_goods
        self.assertEqual(self.good.duty_rate, 0)
        self.assertContains(self.client.get(self.url), 'name="duty_rate"')
        self.client.post(self.url, dict(self.data, duty_rate='7.25'))
        self.good.refresh_from_db()
        self.assertEqual(self.good.duty_rate, Decimal('7.25'))
        bulk_url = reverse('bulk_update_pggoods') + f'?scenario={self.scenario.pk}'
        for value in ('-1', '101', 'nan', '', '1.234'):
            response = self.client.post(self.url, dict(self.data, duty_rate=value))
            self.assertEqual(response.status_code, 200)
            self.assertIn('duty_rate', response.context['form'].errors)
            self.client.post(bulk_url, {'updates': [{'id': self.good.pk, 'duty_rate': value}]}, content_type='application/json')
            self.good.refresh_from_db()
            self.assertEqual(self.good.duty_rate, Decimal('7.25'))
        self.client.post(bulk_url, {'updates': [{'id': self.good.pk, 'duty_rate': '12.50'}]}, content_type='application/json')
        self.good.refresh_from_db()
        self.assertEqual(self.good.duty_rate, Decimal('12.50'))
        target = ScenarioModel.objects.create(name='Duty copy', date_start_plan='2027-01-01', date_end_plan='2027-12-31')
        copy_goods(self.scenario, target)
        self.assertEqual(PGGoods.objects.get(scenario_plan=target).duty_rate, Decimal('12.50'))
        self.client.post(self.url, self.data)
        self.good.refresh_from_db()
        self.assertEqual(self.good.duty_rate, Decimal('12.50'))

    def test_duty_rate_excel_round_trip_and_legacy_import_preserve_value(self):
        import io
        from decimal import Decimal
        from openpyxl import load_workbook
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.good.duty_rate = Decimal('7.25')
        self.good.save()
        response = self.client.get(reverse('export_pggoods'), {'scenario': self.scenario.pk})
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual((sheet.cell(1, 16).value, sheet.cell(2, 16).value), ('Пошлина, %', 7.25))
        sheet.cell(2, 16).value = 12.5
        for legacy in (False, True):
            if legacy:
                sheet.delete_cols(16, 5)
            stream = io.BytesIO()
            workbook.save(stream)
            self.client.post(reverse('import_from_excel'), {'scenario': self.scenario.pk, 'excel_file': SimpleUploadedFile('duty.xlsx', stream.getvalue())})
            self.good.refresh_from_db()
            self.assertEqual(self.good.duty_rate, Decimal('12.50'))

    def test_duty_database_range_constraint(self):
        from django.db import IntegrityError, transaction
        for value in (-1, 101):
            with self.assertRaises(IntegrityError), transaction.atomic():
                PGGoods.objects.filter(pk=self.good.pk).update(duty_rate=value)
