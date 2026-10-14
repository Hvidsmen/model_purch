from io import BytesIO
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from django.core.exceptions import ValidationError
from django.db import transaction
from ..models import Purch, PurchPay, KindLagPay, PGGoods
from ..goods_identity import planning_group_key
from .pricing import reprice_goods

PURCHASE_HEADERS = ['ID', 'Название закупки', 'Тип поставки', 'Лаг доставки, дни', 'Лаг производства, дни']
PAYMENT_HEADERS = ['ID закупки', 'Закупка', 'Название платежа', '% оплаты', 'Лаг платежа, дни', 'Точка отсчёта']


def export_workbook():
    book = Workbook()
    purchases = book.active
    purchases.title = 'Закупки'
    purchases.append(PURCHASE_HEADERS)
    payments = book.create_sheet('Платежи')
    payments.append(PAYMENT_HEADERS)
    for purchase in Purch.objects.order_by('name').prefetch_related('purchpay_set__kind_lag_pay'):
        purchases.append([purchase.pk, purchase.name, purchase.get_supply_type_display(), purchase.lag_income, purchase.lage_make or 0])
        for payment in purchase.purchpay_set.all():
            payments.append([purchase.pk, purchase.name, payment.name, payment.percent_pay, payment.lag_day_pay, str(payment.kind_lag_pay or '')])
    references = book.create_sheet('Точки отсчёта')
    references.append(['Название'])
    for name in KindLagPay.objects.order_by('name').values_list('name', flat=True):
        references.append([name])
    for sheet in book:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='E8EFF8')
            cell.font = Font(bold=True, color='1E3A5F')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(55, max(18, max(len(str(cell.value or '')) for cell in column) + 2))
    for sheet, numbers in [(purchases, [4, 5]), (payments, [4, 5])]:
        for index in numbers:
            sheet.cell(1, index).fill = PatternFill('solid', fgColor='D1FAE5')
    result = BytesIO()
    book.save(result)
    return result.getvalue()


def integer(value, label, minimum=0):
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value() or not minimum <= number <= 2147483647:
            raise ValueError
        return int(number)
    except (InvalidOperation, ValueError):
        raise ValidationError(f'{label}: требуется целое число от {minimum}.')


def import_workbook(upload):
    if not upload or not upload.name.lower().endswith('.xlsx') or upload.size > 5 * 1024 * 1024:
        raise ValidationError('Выберите файл .xlsx размером до 5 МБ.')
    content = upload.read()
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValidationError('Файл слишком большой после распаковки.')
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, ValueError, KeyError, OSError, ParseError, InvalidFileException):
        raise ValidationError('Не удалось прочитать файл Excel.')
    try:
        def rows(title, headers):
            if title not in book.sheetnames:
                raise ValidationError(f'Нет листа «{title}». Используйте выгрузку как шаблон.')
            iterator = book[title].iter_rows(values_only=True)
            if list(next(iterator, ()))[:len(headers)] != headers:
                raise ValidationError(f'Неверные колонки листа «{title}». Используйте выгрузку как шаблон.')
            for number, values in enumerate(iterator, 2):
                if number > 10001:
                    raise ValidationError('Слишком много строк в файле (максимум 10000 на лист).')
                if not any(value is not None for value in values):
                    continue
                yield number, list(values[:len(headers)]) + [None] * max(0, len(headers) - len(values))
        known = {item.pk: item for item in Purch.objects.all()}
        by_name = {item.name_key: item for item in known.values()}
        records, names, ids = {}, {}, set()
        types = {label: key for key, label in Purch.SupplyType.choices}
        types.update({key: key for key, label in Purch.SupplyType.choices})
        for line, row in rows('Закупки', PURCHASE_HEADERS):
            pk, name, supply, delivery, manufacture = row
            name = str(name or '').strip()
            if not name or len(name) > 255:
                raise ValidationError(f'Закупки, строка {line}: укажите название до 255 символов.')
            key = planning_group_key(name)
            pk = integer(pk, 'ID', 1) if pk not in (None, '') else None
            if pk and pk not in known:
                raise ValidationError(f'Закупки, строка {line}: ID не найден. Для новой закупки оставьте ID пустым.')
            target = known[pk] if pk else by_name.get(key)
            if key in names or (target and target.pk in ids):
                raise ValidationError(f'Закупки, строка {line}: закупка повторяется.')
            if key in by_name and target and by_name[key].pk != target.pk:
                raise ValidationError(f'Закупки, строка {line}: название занято другой закупкой.')
            if supply not in types:
                raise ValidationError(f'Закупки, строка {line}: тип поставки — Прямой контракт, P2 или Закупка РФ.')
            record = {'target': target, 'name': name, 'supply_type': types[supply], 'lag_income': integer(delivery, f'Строка {line}, лаг доставки'),
                'lage_make': integer(manufacture, f'Строка {line}, лаг производства'), 'payments': []}
            records[key] = record
            names[key] = record
            if target:
                ids.add(target.pk)
        if not records:
            raise ValidationError('В файле нет закупок.')
        kinds = {}
        for kind in KindLagPay.objects.all():
            kinds.setdefault(kind.name, []).append(kind)
        by_id = {record['target'].pk: record for record in records.values() if record['target']}
        for line, row in rows('Платежи', PAYMENT_HEADERS):
            pk, purchase_name, name, percent, lag, kind_name = row
            record = by_id.get(integer(pk, 'ID закупки', 1)) if pk not in (None, '') else names.get(planning_group_key(purchase_name))
            if record is None:
                raise ValidationError(f'Платежи, строка {line}: закупка отсутствует на листе «Закупки».')
            try:
                percent = Decimal(str(percent).replace(',', '.'))
                if not percent.is_finite() or not 0 <= percent <= 100:
                    raise ValueError
            except (InvalidOperation, ValueError):
                raise ValidationError(f'Платежи, строка {line}: процент должен быть от 0 до 100.')
            name = str(name or '').strip()
            if not name or len(name) > 255:
                raise ValidationError(f'Платежи, строка {line}: укажите название до 255 символов.')
            matches = kinds.get(str(kind_name), []) if kind_name not in (None, '') else []
            if kind_name not in (None, '') and len(matches) != 1:
                raise ValidationError(f'Платежи, строка {line}: неизвестная или неоднозначная точка отсчёта. См. лист «Точки отсчёта».')
            record['payments'].append({'name': name, 'percent_pay': float(percent), 'lag_day_pay': integer(lag, f'Платежи, строка {line}, лаг', -2147483648), 'kind_lag_pay': matches[0] if matches else None})
        with transaction.atomic():
            count = 0
            for record in records.values():
                target = record['target'] or Purch()
                old_name = target.name if target.pk else None
                for field in ['name', 'lag_income', 'lage_make', 'supply_type']:
                    setattr(target, field, record[field])
                target.full_clean()
                target.save()
                if old_name and old_name != target.name:
                    # Goods store the directory name, not a foreign key.
                    aliases = [name for name in PGGoods.objects.values_list('purch', flat=True).distinct() if planning_group_key(name) == planning_group_key(old_name)]
                    PGGoods.objects.filter(purch__in=aliases).update(purch=target.name)
                target.purchpay_set.all().delete()
                PurchPay.objects.bulk_create([PurchPay(purch=target, **values) for values in record['payments']])
                count += 1
            invalid = reprice_goods(strict=False)
        return count, len(invalid)
    finally:
        book.close()
