from io import BytesIO
from openpyxl import load_workbook
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from base.testing import authorize_admin
from .models import Purch, PurchPay, KindLagPay
from .services.purchase_excel import export_workbook, import_workbook


class PurchaseExcelTests(TestCase):
    def setUp(self):
        authorize_admin(self.client)
        self.purchase = Purch.objects.create(name='Поставка', lag_income=120, lage_make=35, supply_type='p2')
        self.kind = KindLagPay.objects.create(name='Отгрузка')
        PurchPay.objects.create(purch=self.purchase, name='Аванс', percent_pay=30, lag_day_pay=10, kind_lag_pay=self.kind)
        PurchPay.objects.create(purch=self.purchase, name='Остаток', percent_pay=70, lag_day_pay=20)
        self.other = Purch.objects.create(name='Вне файла', lag_income=15, lage_make=0)

    def upload(self, book):
        output = BytesIO()
        book.save(output)
        return SimpleUploadedFile('purchases.xlsx', output.getvalue())

    def test_all_parameters_visible_and_excel_roundtrip(self):
        self.purchase.purchpay_set.filter(name='Аванс').update(lag_day_pay=-10)
        page = self.client.get(reverse('purch_list'))
        for value in ['Производство, дни', 'Тип поставки', 'Точка отсчёта', 'Отгрузка', 'Аванс', 'Остаток']:
            self.assertContains(page, value)
        self.assertNotContains(page, '<summary class="payment-summary">')
        response = self.client.get(reverse('export_purchases'))
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.sheetnames, ['Закупки', 'Платежи', 'Точки отсчёта'])
        self.assertEqual(book['Платежи'].cell(2, 6).value, 'Отгрузка')
        count, invalid = import_workbook(self.upload(book))
        self.assertEqual((count, invalid), (2, 0))
        self.assertEqual(list(self.purchase.purchpay_set.values_list('percent_pay', flat=True)), [30, 70])
        self.assertEqual(self.purchase.purchpay_set.first().kind_lag_pay, self.kind)
        self.assertEqual(self.purchase.purchpay_set.first().lag_day_pay, -10)

    def test_import_updates_and_creates_without_touching_other_purchase(self):
        book = load_workbook(BytesIO(export_workbook()))
        sheet = book['Закупки']
        for row in range(sheet.max_row, 1, -1):
            if sheet.cell(row, 1).value == self.other.pk:
                sheet.delete_rows(row)
        sheet.cell(2, 3).value = 'Закупка РФ'
        sheet.cell(2, 4).value = 45
        sheet.append([None, 'Новая', 'Прямой контракт', 10, 12])
        book['Платежи'].append([None, 'Новая', 'Полная оплата', 100, 5, 'Отгрузка'])
        import_workbook(self.upload(book))
        self.purchase.refresh_from_db()
        self.other.refresh_from_db()
        self.assertEqual((self.purchase.supply_type, self.purchase.lag_income), ('russian', 45))
        self.assertEqual(self.other.lag_income, 15)
        self.assertEqual(Purch.objects.get(name='Новая').purchpay_set.get().percent_pay, 100)

    def test_invalid_payment_rolls_back_and_confirmation_required(self):
        book = load_workbook(BytesIO(export_workbook()))
        for row in book['Закупки'].iter_rows(min_row=2):
            if row[0].value == self.purchase.pk:
                row[3].value = 999
        book['Платежи'].cell(3, 4).value = 'NaN'
        with self.assertRaises(ValidationError):
            import_workbook(self.upload(book))
        self.purchase.refresh_from_db()
        self.assertEqual(self.purchase.lag_income, 120)
        self.assertEqual(self.purchase.purchpay_set.count(), 2)
        response = self.client.post(reverse('import_purchases'), {'file': self.upload(book)})
        self.assertRedirects(response, reverse('purch_list'))
        self.purchase.refresh_from_db()
        self.assertEqual(self.purchase.lag_income, 120)

