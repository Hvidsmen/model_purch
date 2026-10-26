from base.testing import authorize_test_case
import importlib.util
import os
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from unittest import skipUnless
from unittest.mock import patch
from django.contrib.staticfiles.testing import StaticLiveServerTestCase
from django.urls import reverse
from .models import GlobalCoeffVersion, GlobalCoeff, Goods, SegmentCoeff, TypeCoeff, VariationCalculate, Subdivision, SubdivisionCoeff, SubdivisionManagerCoeff, Chanel, KindManagerCoeff


@skipUnless(os.environ.get('RUN_BROWSER_TESTS') == '1' and importlib.util.find_spec('playwright'), 'Optional Chromium check')
class GlobalVersionBrowserTests(StaticLiveServerTestCase):
    def test_new_reports_compact_role_screens(self):
        from contextlib import nullcontext
        from uuid import uuid4
        from django.utils import timezone
        from django.contrib.auth.models import User, Group
        from base.models import MotivationReportAccess
        from base.access import POLITICIAN
        from .models import SalesPlanScenario, MotivationExecution
        from .services.manager_detail import manager_detail_report, period_motivation_summary
        from .test_manager_detail import detail_rows
        from playwright.sync_api import expect
        def prepare():
            channel = Chanel.objects.create(chanel_name='Report design')
            branch = Subdivision.objects.create(subdivision_key='A', subdivision_name='A', subdivision_global='A', chanel=channel)
            employee = User.objects.create_user('design-employee')
            employee.groups.add(Group.objects.get(name=POLITICIAN))
            MotivationReportAccess.objects.create(user=employee, subdivision=branch, deal_author='Иванов')
            SalesPlanScenario.objects.create(title='План 2026', source_version='Visual plan', loaded_at=timezone.now())
            MotivationExecution.objects.create(run_id=uuid4(), job_key=uuid4(), state='completed', started_at=timezone.now(), completed_at=timezone.now(), parameters={'fact_year': 2026, 'plan_version': 'Visual plan', 'quarter_share': '80', 'year_share': '20', 'payout_threshold': '.8'})
            return employee.pk
        employee = self.db(prepare)
        source = [dict(row, АвторСделки=name, КварталМотивации='4 Квартал(1-4 кв)' if row['КварталМотивации'] != 'Год' else 'Год') for name in ('Иванов', 'Петров', 'Сидоров') for row in detail_rows()]
        def load(kind, values, **kwargs):
            rows = [row for row in source if not values.get('DealAuthor') or row['АвторСделки'] in values['DealAuthor']]
            if kwargs.get('summary_only'):
                return {'currency': 'USD', 'period_summary': period_motivation_summary(rows)}
            return manager_detail_report(rows, performance_rows=detail_rows())
        budget = [{'name': label, 'global': 24000, 'local': 27000, 'difference': 3000, 'difference_percent': .125,
                   'children': [{'name': 'Политики', 'global': 24000, 'local': 27000, 'difference': 3000, 'children': []}]}
                  for label in ('Компания', 'Д-Москва')]
        with patch('admin_motivation.services.reports_next.current_run', return_value=nullcontext()), patch('admin_motivation.services.reports_next.load_report', side_effect=load), patch('admin_motivation.services.reports_next.planned_budget', return_value=budget):
            self.page.set_viewport_size({'width':1440,'height':1000})
            self.page.goto(self.live_server_url + reverse('index_admin_motivation'))
            expect(self.page.get_by_role('heading', name='Мотивация', exact=True, level=1)).to_be_visible()
            expect(self.page.locator('.ux-home-primary')).to_contain_text('Работа с моделью')
            self.page.screenshot(path='/tmp/motivation-home-ux.png', full_page=True)
            self.page.goto(self.live_server_url + reverse('motivation_reports_next'))
            expect(self.page.get_by_role('link', name='Моя премия', exact=False)).to_be_visible()
            self.page.screenshot(path='/tmp/new-reports-catalog-compact.png', full_page=True)
            self.page.goto(self.live_server_url + reverse('motivation_report_next', args=['mine']) + f'?run=1&preview_user={employee}&type=Политики')
            expect(self.page.locator('.nr-period:not(.nr-goal):not(.nr-note):not(.nr-summary-card)')).to_have_count(2)
            expect(self.page.locator('.next-total')).to_contain_text('180,00')
            expect(self.page.get_by_role('heading', name='Сколько осталось до 100%')).to_be_visible()
            expect(self.page.get_by_role('heading', name='Из каких товаров сложилась премия')).to_be_visible()
            expect(self.page.locator('.nr-target-reward').first).to_contain_text('Пока нельзя рассчитать')
            expect(self.page.locator('.nr-matrix-node').first).to_be_visible()
            expect(self.page.locator('.nr-matrix-header')).to_contain_text('Возможная премия')
            expect(self.page.locator('.nr-matrix-total')).to_be_visible()
            expect(self.page.locator('.nr-with-total').first.locator(':scope > article')).to_have_count(3)
            expect(self.page.locator('.nr-goal.nr-summary-card')).to_have_count(1)
            expect(self.page.locator('.nr-with-total').first.locator(':scope > article').nth(2)).to_contain_text('Итого за выбранные периоды')
            self.assertGreater(len(self.page.locator('.nr-matrix-header').evaluate('(el) => getComputedStyle(el).gridTemplateColumns').split()), 8)
            self.assertGreaterEqual(float(self.page.locator('.nr-matrix-header > span').first.evaluate('(el) => getComputedStyle(el).fontSize').replace('px','')), 13)
            expect(self.page.locator('.motivation-workflow')).to_have_count(0)
            expect(self.page.locator('.nr-kpis')).to_have_count(0)
            self.page.locator('[data-matrix-details]').uncheck()
            expect(self.page.locator('.nr-matrix')).to_have_attribute('data-mode', 'basic')
            expect(self.page.locator('.nr-matrix-header .nr-matrix-execution').first).to_be_hidden()
            self.page.locator('[data-matrix-details]').check()
            expect(self.page.locator('.nr-matrix-header .nr-matrix-execution').first).to_be_visible()

            self.page.screenshot(path='/tmp/new-reports-mine-compact.png', full_page=True)
            self.page.locator('.nr-matrix-node > summary').first.click()
            self.page.locator('.nr-matrix-children > .nr-matrix-node > summary').first.click()
            self.page.get_by_role('link', name='Ставки и отгрузки', exact=True).first.click()
            expect(self.page.locator('#next-deep-calculation')).to_have_attribute('open', '')
            expect(self.page.locator('#manager-detail-extra')).to_be_checked()
            expect(self.page.locator('.manager-detail-cards')).to_have_count(0)

            self.page.goto(self.live_server_url + reverse('motivation_report_next', args=['team']) + '?run=1&type=Политики')
            expect(self.page.locator('.nr-table tbody tr')).to_have_count(3)
            expect(self.page.get_by_role('heading', name='Итого по видам мотивации')).to_be_visible()
            expect(self.page.locator('.nr-table thead')).to_contain_text('4 кв.')
            expect(self.page.get_by_text('По периодам', exact=True)).to_have_count(0)
            search = self.page.get_by_role('searchbox', name='Найти сотрудника во всём отчёте')
            search.fill('Петров')
            with self.page.expect_navigation():
                self.page.get_by_role('button', name='Найти', exact=True).click()
            expect(self.page.locator('.nr-table tbody tr')).to_have_count(1)
            expect(self.page.locator('.nr-table tbody')).to_contain_text('Петров')
            self.page.screenshot(path='/tmp/new-reports-team-compact.png', full_page=True)
            self.page.goto(self.live_server_url + reverse('motivation_report_next', args=['registry']) + '?run=1&type=Политики')
            expect(self.page.locator('.nr-table tbody tr')).to_have_count(3)
            expect(self.page.locator('.nr-table thead')).to_contain_text('Год')
            self.page.screenshot(path='/tmp/new-reports-registry-pivot.png', full_page=True)

            self.page.goto(self.live_server_url + reverse('motivation_report_next', args=['executive']) + '?run=1')
            expect(self.page.locator('.nr-bars')).to_be_visible()
            self.page.locator('.nr-filter-panel > summary').click()
            self.page.locator('.next-types > summary').click()
            self.page.locator('[data-type-all]').click()
            expect(self.page.locator('.next-types input:checked')).to_have_count(3)
            self.page.locator('[data-type-clear]').click()
            expect(self.page.locator('.next-types input:checked')).to_have_count(0)
            self.page.locator('input[name=type][value="Политики"]').check()
            self.page.locator('input[name=type][value="Продажи"]').check()
            expect(self.page.locator('[data-type-caption]')).to_have_text('Политики, Продажи')
            self.page.locator('.next-periods > summary').click()
            self.page.locator('input[name=period][value="4"]').check()
            self.page.locator('input[name=period][value="5"]').check()
            self.page.locator('.next-periods > summary').click()
            with self.page.expect_navigation():
                self.page.get_by_role('button', name='Показать', exact=True).click()
            expect(self.page.locator('.next-types input:checked')).to_have_count(2)
            expect(self.page.locator('.next-periods input:checked')).to_have_count(2)

            self.page.screenshot(path='/tmp/new-reports-executive-compact.png', full_page=True)
            self.page.set_viewport_size({'width':390,'height':844})
            self.assertLessEqual(self.page.evaluate('document.documentElement.scrollWidth'), 390)
            self.page.screenshot(path='/tmp/new-reports-mobile-compact.png', full_page=True)
        self.assertEqual(self.errors, [])

    def test_new_report_tab_plan_choice_and_employee_drilldown(self):
        from contextlib import nullcontext
        from uuid import uuid4
        from django.utils import timezone
        from .models import SalesPlanScenario, MotivationExecution
        from .services.manager_detail import manager_detail_report
        from .test_manager_detail import detail_rows
        from playwright.sync_api import expect
        def prepare():
            plans = [SalesPlanScenario.objects.create(title=title, source_version=title, loaded_at=timezone.now()) for title in ('Run plan', 'Alternative plan')]
            execution = MotivationExecution.objects.create(run_id=uuid4(), job_key=uuid4(), state='completed', started_at=timezone.now(), completed_at=timezone.now(), parameters={'fact_year': 2026, 'plan_version': 'Run plan', 'quarter_share': '80', 'year_share': '20'})
            return plans[1].pk
        alternate = self.db(prepare)
        budget = [{'name': 'Компания', 'global': 100, 'local': 120, 'difference': 20, 'children': []}]
        annual = manager_detail_report([detail_rows()[1]])
        with patch('admin_motivation.services.reports_next.planned_budget', return_value=budget) as comparison, patch('admin_motivation.services.reports_next.current_run', return_value=nullcontext()), patch('admin_motivation.services.reports_next.load_report', return_value=annual):
            self.page.goto(self.live_server_url + reverse('motivation_reports'))
            self.page.get_by_role('link', name='Новые отчёты · пилот', exact=True).click()
            expect(self.page.get_by_role('heading', name='Новые отчёты мотивации')).to_be_visible()
            self.page.get_by_role('link', name='Сводный отчёт: план премий', exact=False).click()
            self.page.select_option('select[name=plan]', str(alternate))
            self.page.locator('.next-periods > summary').click()
            self.page.locator('[data-period-all]').click()
            expect(self.page.locator('.next-periods input:checked')).to_have_count(5)
            self.page.locator('[data-period-clear]').click()
            expect(self.page.locator('.next-periods input:checked')).to_have_count(0)
            self.page.locator('input[name=period][value="4"]').check()
            self.page.locator('input[name=period][value="5"]').check()
            self.page.locator('.next-periods > summary').click()

            self.page.get_by_role('button', name='Показать', exact=True).click()
            expect(self.page.locator('.nr-kpi').nth(2)).to_contain_text('20,00')
            self.assertEqual(comparison.call_args.args[1].pk, alternate)
            self.assertEqual(comparison.call_args.args[2].getlist('period'), ['4', '5'])
            self.page.goto(self.live_server_url + reverse('motivation_report_next', args=['team']) + '?run=1&type=Политики&period=5')
            expect(self.page.locator('.next-total')).to_contain_text('95,00 USD')
            self.page.get_by_role('link', name='Как рассчитано →', exact=True).first.click()
            expect(self.page.get_by_role('heading', name='Как получилась сумма')).to_be_visible()
            self.page.locator('#next-deep-calculation > summary').click()
            expect(self.page.locator('#manager-detail-panel')).to_be_visible()
            expect(self.page.locator('.manager-detail-cards')).to_have_count(0)
            expect(self.page.locator('#manager-detail-nodes')).to_contain_text('Иванов')
            self.page.screenshot(path='/tmp/new-motivation-reports-detail.png', full_page=True)
        self.assertEqual(self.errors, [])

    def test_coverage_bulk_button_includes_other_pages(self):
        from playwright.sync_api import expect
        from .models import SalesPlanScenario, SalesPlanLine
        from django.utils import timezone
        def prepare():
            scenario = SalesPlanScenario.objects.create(title='Bulk coverage', loaded_at=timezone.now())
            SalesPlanLine.objects.bulk_create([SalesPlanLine(scenario=scenario, plan_date=date(2026, 1, 1),
                subdivision='A', planning_group_sales='Bulk planning', group='Bulk ERP',
                brand=f'Bulk browser {i}', amount_usd=100) for i in range(105)])
            return scenario.pk
        self.db(lambda: SegmentCoeff.objects.filter(pk=self.coeff.segment_id).update(segment_name='K0'))
        plan = self.db(prepare)
        self.page.goto(self.live_server_url + reverse('motivation_coefficient_coverage') + f'?version={self.first.pk}&plan={plan}&scope=global&kind=Политики&run=1')
        self.page.get_by_role('button', name='Создать для всех найденных пропусков', exact=True).click()
        expect(self.page.get_by_role('heading', name='Создать коэффициенты для всех пропусков')).to_be_visible()
        expect(self.page.locator('tbody tr')).to_have_count(105)
        self.page.get_by_role('button', name='Создать все предложенные коэффициенты в черновике', exact=True).click()
        self.assertEqual(self.db(lambda: GlobalCoeff.objects.filter(version=self.first, goods__brand__startswith='Bulk browser').count()), 525)
        self.assertEqual(self.errors, [])

    def test_coverage_scope_and_create_missing_cells(self):
        from playwright.sync_api import expect
        from .models import SalesPlanScenario, SalesPlanLine
        from django.utils import timezone
        def prepare():
            scenario = SalesPlanScenario.objects.create(title='Coverage repair', loaded_at=timezone.now())
            SalesPlanLine.objects.create(scenario=scenario, plan_date=date(2026, 1, 1),
                subdivision='A', planning_group_sales='Repair planning', group='Repair ERP',
                brand='Repair brand', amount_usd=100)
            return scenario.pk
        plan = self.db(prepare)
        self.page.goto(self.live_server_url + reverse('motivation_coefficient_coverage') + f'?version={self.first.pk}&plan={plan}&scope=global&kind=Политики&run=1')
        expect(self.page.locator('select[name=scope]')).to_have_value('global')
        self.page.get_by_role('link', name='Создать коэффициенты', exact=True).click()
        expect(self.page.get_by_role('heading', name='Создать недостающие коэффициенты')).to_be_visible()
        for i in range(5):
            self.page.locator(f'input[name=K{i}]').fill('0.5')
        self.page.screenshot(path='/tmp/coverage-create-coefficients.png', full_page=True)
        self.page.get_by_role('button', name='Создать в черновике', exact=True).click()
        self.assertEqual(self.db(lambda: GlobalCoeff.objects.filter(version=self.first, goods__brand='Repair brand').count()), 5)
        self.assertEqual(self.errors, [])

    def test_discount_segment_form_copy_history_and_nulls(self):
        from playwright.sync_api import expect
        from .models import DiscountSegment
        from .services.approval import approve_review, review_approval
        from .services.versions import create_version
        self.page.goto(self.live_server_url + reverse('motivation_discount_segments'))
        expect(self.page.get_by_role('heading', name='Сегменты скидок', exact=True)).to_be_visible()
        self.page.fill('#id_segments-0-type_discount', 'MASS')
        for index, value in enumerate(('', '10', '17', '19', '24')):
            self.page.fill(f'#id_segments-0-o{index}', value)
        self.page.get_by_role('button', name='Сохранить сегменты', exact=True).click()
        expect(self.page.locator('.discount-notice')).to_contain_text('сохранены')
        self.assertIsNone(self.db(lambda: DiscountSegment.objects.get(version=self.first).o0))
        self.db(lambda: approve_review(review_approval(GlobalCoeffVersion.objects.get(pk=self.first.pk))['token']))
        latest = self.db(lambda: create_version(date(2026, 1, 1), source_version=self.first.pk).pk)
        self.page.goto(self.live_server_url + reverse('motivation_discount_segments') + f'?version={latest}')
        expect(self.page.locator('#id_segments-0-o1')).to_have_value('10.0000')
        self.page.fill('#id_segments-0-o4', '30')
        self.page.fill('#id_segments-1-type_discount', 'NO_DISCOUNT')
        self.page.get_by_role('button', name='Сохранить сегменты', exact=True).click()
        expect(self.page.locator('.discount-notice')).to_contain_text('сохранены')
        self.assertIsNone(self.db(lambda: DiscountSegment.objects.get(version_id=latest, type_discount='NO_DISCOUNT').o4))
        self.page.set_viewport_size({'width':1440,'height':1000})
        self.page.screenshot(path='/tmp/da-discount-segments.png',full_page=True)
        self.page.set_viewport_size({'width':390,'height':844})
        self.assertLessEqual(self.page.evaluate('document.documentElement.scrollWidth'),390)
        self.page.goto(self.live_server_url + reverse('motivation_discount_segments') + f'?version={self.first.pk}')
        expect(self.page.locator('#id_segments-0-o4')).to_be_disabled()
        expect(self.page.locator('#id_segments-0-o4')).to_have_value('24.0000')
        self.assertEqual(self.errors, [])

    def test_bulk_formula_reopen_and_unused_director_delete(self):
        from playwright.sync_api import expect
        from .models import DirectorEmployee, VersionReopenLog
        from .services.pilot_formula import GLOBAL_SUBDIVISION
        from .services.approval import review_approval, approve_review
        def setup_second():
            original = GlobalCoeff.objects.get(pk=self.coeff.pk)
            good = Goods.objects.create(goods_key='OtherKey', planning_group_sales='Other', group='Other ERP', brand='Other brand')
            row = GlobalCoeff.objects.create(version=self.first, goods=good, type_coeff=original.type_coeff, segment=original.segment, variation_calculate=original.variation_calculate, motivation_coeff=.2, manager_coeff=1)
            return good.pk, row.pk, VariationCalculate.objects.get_or_create(variation_name=GLOBAL_SUBDIVISION)[0].pk
        good, row, formula = self.db(setup_second)
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation'))
        self.page.get_by_role('button',name='Формула для товаров',exact=True).click()
        self.page.locator('#coefficient-search').fill('OtherKey')
        self.page.locator('[data-bulk-visible]').click()
        expect(self.page.locator('[data-bulk-count]')).to_have_text('Выбрано: 1')
        expect(self.page.locator('#editor-dirty-count')).to_have_text('Изменено: 0')
        self.page.select_option('#bulk-variation-form select[name=variation]', str(formula))
        self.page.get_by_role('button',name='Применить формулу',exact=True).click()
        self.page.get_by_role('button',name='Подтвердить формулу',exact=True).click()
        expect(self.page.locator('#version-panel')).to_contain_text('Формула применена')
        self.assertEqual(self.db(lambda: GlobalCoeff.objects.get(pk=row).variation_calculate_id),formula)
        self.assertNotEqual(self.db(lambda: GlobalCoeff.objects.get(pk=self.coeff.pk).variation_calculate_id),formula)
        self.db(lambda: approve_review(review_approval(GlobalCoeffVersion.objects.get(pk=self.first.pk))['token']))
        self.page.reload()
        self.page.locator('.version-tools button[aria-controls="version-reopen"]').click()
        self.page.fill('#version-reopen input[name=reason]','Browser correction')
        self.page.locator('#version-reopen input[name=confirm]').check()
        self.page.get_by_role('button',name='Вернуть в черновик',exact=True).click()
        expect(self.page.locator('#version-panel')).to_contain_text('Версия возвращена')
        self.assertTrue(self.db(lambda: VersionReopenLog.objects.filter(version=self.first).exists()))
        director = self.db(lambda: DirectorEmployee.objects.create(full_name='Delete browser director').pk)
        self.page.goto(self.live_server_url + reverse('motivation_director_directory'))
        record = self.page.locator('tbody tr').filter(has_text='Delete browser director')
        record.locator('.director-delete>summary').click()
        record.locator('input[name=confirm]').check()
        record.get_by_role('button',name='Подтвердить удаление',exact=True).click()
        expect(self.page.get_by_role('status')).to_contain_text('Директор удалён')
        self.assertFalse(self.db(lambda: DirectorEmployee.objects.filter(pk=director).exists()))
        self.assertEqual(self.errors, [])

    def test_version_tools_compact_layout_and_director_bulk_formula(self):
        from .models import GlobalDirectorCoeff
        from playwright.sync_api import expect
        self.good = self.db(lambda: self.coeff.goods)
        formula = self.db(lambda: VariationCalculate.objects.get_or_create(variation_name='Подразделение(Глобально)')[0].pk)
        row = self.db(lambda: GlobalDirectorCoeff.objects.create(version=self.first,goods=self.good,k0='.01',k4='.05').pk)
        self.page.set_viewport_size({'width':1920,'height':1080})
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation'))
        self.page.get_by_role('button',name='Добавить товар',exact=True).click()
        expect(self.page.locator('#coefficient-action-hub #add-product-details')).to_have_attribute('open','')
        self.page.get_by_role('button',name='Импорт и справочники',exact=True).click()
        expect(self.page.locator('#add-product-details')).not_to_have_attribute('open','')
        self.assertEqual(self.page.locator('#coefficient-tools input[type=file]').evaluate('(el) => el.form.id'),'mainCoeffForm')
        self.page.get_by_role('button',name='Импорт и справочники',exact=True).click()
        geometry = self.page.evaluate("""() => {
            const box = selector => document.querySelector(selector).getBoundingClientRect();
            return {tools:box('.version-tools').x,picker:box('.version-picker').right,
                    variants:box('.compare-variants').x,compare:box('#comparison-refresh').right};
        }""")
        self.assertGreater(geometry['tools'],geometry['picker'])
        self.assertGreater(geometry['variants'],geometry['compare'])
        self.page.locator('.version-tools button[aria-controls="version-history"]').click()
        expect(self.page.locator('#version-history table')).to_be_visible()
        self.page.locator('.version-tools button[aria-controls="create-version-details"]').click()
        expect(self.page.locator('#version-history')).not_to_have_attribute('open','')
        panel_box = self.page.locator('#version-panel').bounding_box()
        content_box = self.page.locator('#create-version-details .version-tool-content').bounding_box()
        self.assertGreaterEqual(panel_box['y'] + panel_box['height'], content_box['y'] + content_box['height'])
        self.page.screenshot(path='/tmp/portal-version-create-inline.png',full_page=True)
        self.page.keyboard.press('Escape')
        expect(self.page.locator('#create-version-details')).not_to_have_attribute('open','')
        self.page.get_by_role('button',name='Формула для товаров',exact=True).click()
        self.page.locator('#bulk-variation-form input[value="Политики"]').uncheck()
        self.page.locator('#bulk-variation-form input[value="Директора"]').check()
        self.page.locator('#bulk-variation-form select[name=variation]').select_option(str(formula))
        self.page.locator(f'.bulk-good[value="{self.good.pk}"]').check()
        self.page.get_by_role('button',name='Применить формулу',exact=True).click()
        with self.page.expect_navigation():
            self.page.get_by_role('button',name='Подтвердить формулу',exact=True).click()
        self.assertEqual(self.db(lambda: GlobalDirectorCoeff.objects.get(pk=row).variation_calculate_id),formula)
        self.assertEqual(str(self.db(lambda: GlobalDirectorCoeff.objects.get(pk=row).k0)),'0.010000')
        expect(self.page.get_by_text('Директор: Подразделение(Глобально)',exact=True)).to_be_visible()
        self.page.set_viewport_size({'width':390,'height':844})
        self.page.locator('.version-tools button[aria-controls="version-history"]').click()
        self.assertLessEqual(self.page.evaluate('document.documentElement.scrollWidth'),391)
        self.assertLessEqual(self.page.locator('#version-history .version-tool-content').bounding_box()['x'] + self.page.locator('#version-history .version-tool-content').bounding_box()['width'],391)
        self.page.screenshot(path='/tmp/portal-compact-version-tools.png',full_page=True)

    def test_workflow_grouping_and_inline_product_settings_preserve_edits(self):
        from playwright.sync_api import expect
        formula=self.db(lambda: VariationCalculate.objects.get_or_create(variation_name='Подразделение(Глобально)')[0].pk)
        self.page.set_viewport_size({'width':1920,'height':1080})
        self.page.goto(self.live_server_url+reverse('global_coeff_admin_motivation'))
        expect(self.page.locator('.workflow-steps')).to_contain_text('Утверждение')
        group=self.page.locator('.coefficient-group-heading button')
        group.click();expect(self.page.locator('.coefficient-row')).not_to_be_visible()
        group.click();expect(self.page.locator('.coefficient-row')).to_be_visible()
        self.page.locator('.brand-settings > summary').click()
        panel=self.page.locator('.coefficient-row .brand-settings')
        expect(self.page.locator('.motivation-side-panel')).to_have_count(0)
        expect(panel).to_be_visible()
        panel.locator('select').select_option(str(formula))
        expect(panel.locator('.formula-explanation')).to_contain_text('личный объём')
        self.page.screenshot(path='/tmp/motivation-ux-desktop.png',full_page=True)
        panel.locator('summary').click()
        expect(self.page.locator('#mainCoeffForm .select-var-calc')).to_have_value(str(formula))
        expect(self.page.locator('#editor-dirty-count')).not_to_have_text('Изменено: 0')
        self.page.set_viewport_size({'width':390,'height':844})
        self.page.locator('.brand-settings > summary').click()
        expect(panel).to_be_visible()
        self.assertLessEqual(panel.bounding_box()['width'],390)
        self.assertLessEqual(self.page.evaluate('document.documentElement.scrollWidth'),391)
        self.page.screenshot(path='/tmp/motivation-ux-mobile.png',full_page=True)
        self.page.keyboard.press('Escape')
        expect(panel).not_to_have_attribute('open','')
        self.assertEqual(self.errors,[])

    def test_completed_job_links_to_report_with_recorded_context(self):
        from playwright.sync_api import expect
        from django.utils import timezone
        from uuid import uuid4
        from .models import MotivationCalculationSchedule, MotivationExecution
        def setup_run():
            job=MotivationCalculationSchedule.objects.create(name='Report link',fact_year=2026,plan_version='Recorded plan',synced=True)
            run=MotivationExecution.objects.create(run_id=uuid4(),job_key=job.key,state='completed',started_at=timezone.now(),completed_at=timezone.now(),parameters={'fact_year':2026,'plan_version':'Recorded plan'},export_snapshot={'versions':[{'title':'Recorded coefficients','from':'2026-01-01'}]})
            return run
        run=self.db(setup_run)
        with patch('admin_motivation.operation_views.calculation_jobs.monitor',return_value={'status':'Завершено','history':[]}), patch('admin_motivation.services.execution_tracking.observe',return_value=run):
            self.page.goto(self.live_server_url+reverse('motivation_calculations'))
            link=self.page.locator('[data-job-report]');expect(link).to_be_visible()
            expect(link).to_have_attribute('href',__import__('re').compile('year=2026'))
            with patch('admin_motivation.report_views.choices',return_value=[]),patch('admin_motivation.report_views.load_report',return_value={'currency':'USD','nodes':[],'periods':[]}):
                with self.page.expect_navigation():link.click()
            expect(self.page.locator('.ux-report-receipt')).to_contain_text('Recorded plan')
            expect(self.page.locator('.ux-report-receipt')).to_contain_text('Recorded coefficients')
        self.assertEqual(self.errors,[])

    def test_segments_subdivision_search_and_navigation(self):
        from playwright.sync_api import expect
        def setup_branch():
            channel=Chanel.objects.create(chanel_name='Search nav')
            return Subdivision.objects.create(subdivision_key='NAV',subdivision_name='Д-Москва',subdivision_global='Д-Москва',chanel=channel).pk
        branch=self.db(setup_branch)
        self.page.goto(self.live_server_url + reverse('motivation_discount_segments'))
        labels = self.page.evaluate("""() => [...document.querySelector('#subdivision-nav').closest('ul').children].map(li => li.querySelector(':scope > a,:scope > details > summary').textContent.trim())""")
        self.assertEqual(labels, ['Планы продаж','Сегменты скидок','Глобальные коэффициенты','Подразделения','Отчёты','Выгрузка в DWH'])
        self.page.locator('#subdivision-nav > summary').click()
        self.page.locator('#subdivision-search').fill('Москва')
        link=self.page.locator('#subdivision-links a').filter(has_text='Москва')
        expect(link).to_be_visible()
        expect(self.page.locator('#subdivision-search-empty')).not_to_be_visible()
        with self.page.expect_navigation(): link.click()
        self.assertIn(reverse('coeff_subdivisions_admin_motivation',args=[branch]),self.page.url)
        self.assertEqual(self.errors, [])

    def test_verified_save_receipt_and_stale_tab_rejection(self):
        from playwright.sync_api import expect
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation'))
        cell = self.page.locator(f'input[name="global_coeff={self.coeff.pk}"]')
        cell.fill('25%');cell.press('Tab')
        with self.page.expect_navigation():
            self.page.locator('#editor-save').click()
        expect(self.page.locator('#version-panel')).to_contain_text('Сохранение проверено. Изменено полей:')
        self.assertEqual(self.db(lambda: GlobalCoeff.objects.get(pk=self.coeff.pk).motivation_coeff), .25)
        def another_save():
            row=GlobalCoeff.objects.get(pk=self.coeff.pk)
            row.motivation_coeff=.3;row.save()
        self.db(another_save)
        cell.fill('40%');cell.press('Tab')
        with self.page.expect_navigation():
            self.page.locator('#editor-save').click()
        expect(self.page.locator('#version-panel')).to_contain_text('Версия изменилась')
        self.assertEqual(self.db(lambda: GlobalCoeff.objects.get(pk=self.coeff.pk).motivation_coeff), .3)
        self.assertEqual(self.errors, [])

    def test_admin_calculation_job_create_and_run(self):
        from .models import SalesPlanScenario
        from django.utils import timezone
        for version in ('Selected SQL plan', 'Other SQL plan'):
            self.db(lambda version=version: SalesPlanScenario.objects.create(title=version, source_version=version, loaded_at=timezone.now()))
        from unittest.mock import patch
        from playwright.sync_api import expect
        with patch('admin_motivation.services.calculation_jobs.sync_job') as sync, patch('admin_motivation.services.calculation_jobs.start_job') as start, patch('admin_motivation.services.calculation_jobs.status', return_value='Выполняется'), patch('admin_motivation.services.calculation_jobs.monitor', return_value={'status': 'Выполняется', 'history': []}):
            self.page.goto(self.live_server_url + reverse('motivation_calculations'))
            self.page.locator('#id_name').fill('Browser calculation')
            self.page.locator('#id_fact_year').fill('2026')
            self.page.locator('#id_plan_version').select_option('Selected SQL plan')
            self.page.get_by_role('button', name='Сохранить задание в SQL Server', exact=True).click()
            expect(self.page.get_by_text('Browser calculation', exact=True)).to_be_visible()
            card = self.page.locator('.calculation-job')
            card.locator('.calculation-edit > summary').click()
            card.locator('[name=fact_year]').fill('2025')
            card.locator('[name=plan_version]').select_option('Other SQL plan')
            card.locator('[name=quarter_share]').fill('75')
            card.locator('[name=year_share]').fill('25')
            card.locator('[name=payout_threshold]').fill('0.85')
            card.get_by_role('button', name='Сохранить параметры и расписание', exact=True).click()
            expect(self.page.locator('.calculation-job')).to_contain_text('Год факта: 2025')
            self.page.get_by_role('button', name='Запустить сейчас', exact=True).click()
            expect(self.page.locator('[data-job-status]')).to_have_text('Выполняется')
            self.assertEqual(sync.call_count, 2)
            start.assert_called_once()
            self.assertEqual(self.errors, [])

    def test_approved_manager_correction_has_reason_confirmation_and_audit(self):
        from playwright.sync_api import expect
        from .models import ManagerCorrectionLog
        from .services.approval import approve_review, review_approval
        def prepare():
            channel = Chanel.objects.create(chanel_name='Correction UI')
            sub = Subdivision.objects.create(subdivision_key='Correction',subdivision_name='Correction',subdivision_global='Correction',chanel=channel)
            kind = KindManagerCoeff.objects.create(name='Год')
            row = SubdivisionManagerCoeff.objects.create(version=self.first,subdivision=sub,kind=kind,coeff=.8)
            approve_review(review_approval(GlobalCoeffVersion.objects.get(pk=self.first.pk))['token'])
            return sub.pk,kind.pk,row.pk
        sub,kind,row = self.db(prepare)
        self.page.goto(self.live_server_url + reverse('coeff_subdivisions_admin_motivation',args=[sub]))
        panel = self.page.locator('#manager-details')
        panel.locator(':scope > summary').click()
        field = panel.locator(f'input[name="kind_coeff={kind}"]')
        expect(field).to_be_enabled()
        field.fill('90%')
        panel.locator('[name=manager_reason]').fill('Итоги квартала')
        panel.locator('[name=confirm_manager_change]').check()
        with self.page.expect_navigation():
            panel.get_by_role('button',name='Сохранить руководителя',exact=True).click()
        self.assertEqual(self.db(lambda: SubdivisionManagerCoeff.objects.get(pk=row).coeff),.9)
        self.assertEqual(self.db(lambda: ManagerCorrectionLog.objects.get().reason),'Итоги квартала')
        self.assertEqual(self.db(lambda: GlobalCoeffVersion.objects.get(pk=self.first.pk).status),'approved')
        panel.locator(':scope > summary').click()
        panel.get_by_text('Журнал изменений поправок',exact=True).click()
        expect(panel.get_by_text('Итоги квартала',exact=False)).to_be_visible()
        expect(self.page.locator('#subCoeffForm fieldset')).to_have_attribute('disabled', '')

    def test_calculation_status_and_history_refresh_without_losing_form_edits(self):
        from unittest.mock import patch
        from playwright.sync_api import expect
        from .models import MotivationCalculationSchedule
        self.db(lambda: MotivationCalculationSchedule.objects.create(name='Live monitor', synced=True))
        snapshot = {'status': 'Выполняется', 'history': []}
        with patch('admin_motivation.services.calculation_jobs.status', return_value='Выполняется'), patch('admin_motivation.services.calculation_jobs.monitor', side_effect=lambda job: snapshot.copy()):
            self.page.clock.install()
            self.page.goto(self.live_server_url + reverse('motivation_calculations'))
            expect(self.page.locator('[data-job-monitor-note]')).to_contain_text('Проверено в')
            self.page.locator('#id_name').fill('Несохранённый текст')
            self.page.locator('.calculation-history summary').click()
            snapshot.update(status='Завершено', history=[{'id': 1, 'started': '24.10.2026 12:00:00',
                'duration': '00:05:42', 'status': 'Завершено', 'message': '<script>history message</script>'}])
            self.page.clock.fast_forward(10000)
            expect(self.page.locator('[data-job-status]')).to_have_text('Завершено')
            expect(self.page.locator('[data-job-history]')).to_contain_text('00:05:42')
            expect(self.page.locator('[data-job-history]')).to_contain_text('<script>history message</script>')
            expect(self.page.locator('[data-job-history] script')).to_have_count(0)
            expect(self.page.locator('#id_name')).to_have_value('Несохранённый текст')
            expect(self.page.locator('.calculation-history')).to_have_attribute('open', '')

    def test_manager_author_dropdown_is_visible_below_report_card(self):
        from unittest.mock import patch
        from playwright.sync_api import expect
        from .services.reports import FILTERS
        options = [{'name': name, 'label': caption, 'multiple': multiple,
                    'options': [{'value': f'Автор {i}', 'selected': False} for i in range(40)] if name == 'author' else []}
                   for name, parameter, caption, dataset, multiple in FILTERS['manager']]
        with patch('admin_motivation.report_views.choices', return_value=options):
            self.page.goto(self.live_server_url + reverse('motivation_report', args=['manager']))
            self.page.locator('select[name=author] + .searchable-filter > button').click()
            popup = self.page.locator('select[name=author] + .searchable-filter .searchable-filter-menu')
            expect(popup).to_be_visible()
            self.assertTrue(popup.evaluate('(el) => {const r=el.getBoundingClientRect(); return el.contains(document.elementFromPoint(r.left+20,r.bottom-15));}'))
            self.assertEqual(self.errors, [])

    def test_director_matrix_upload_uses_selected_version_and_existing_bindings(self):
        from playwright.sync_api import expect
        from .models import DirectorEmployee, SubdivisionDirectorCoeff
        from .test_director_coeff_import import workbook
        def setup():
            channel = Chanel.objects.create(chanel_name='Import channel')
            branch = Subdivision.objects.create(subdivision_key='Import', subdivision_name='Import', subdivision_global='Import', chanel=channel)
            for name in ('Первый директор', 'Второй директор'):
                employee = DirectorEmployee.objects.create(full_name=name)
                employee.subdivisions.add(branch)
            Goods.objects.create(goods_key='Import', planning_group_sales='1. RAC', group='Group', brand='*')
            return branch.pk
        branch = self.db(setup)
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation') + f'?version={self.first.pk}')
        self.page.locator('#coefficient-tools').evaluate('(el) => el.open = true')
        self.page.get_by_role('link', name='Загрузить коэффициенты директоров', exact=True).click()
        expect(self.page.get_by_text('Выбранная версия:', exact=False)).to_contain_text('01.01.2001')
        self.page.locator('[name=file]').set_input_files({'name': 'directors.xlsx', 'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'buffer': workbook().read()})
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Загрузить в выбранную версию', exact=True).click()
        expect(self.page.locator('.alert-success')).to_contain_text('создано: 2')
        self.assertEqual(self.db(lambda: SubdivisionDirectorCoeff.objects.filter(version=self.first, subdivision_id=branch).count()), 2)
        self.assertEqual(self.errors, [])

    def test_subdivision_summary_expands_types_and_keeps_company_total_first(self):
        from unittest.mock import patch
        from playwright.sync_api import expect
        from .services.reports import subdivision_report
        from .test_reports import NativeMotivationReportTests
        rows = [dict(NativeMotivationReportTests.division_rows()[0], ТипКоэф=kind, Подразделение=branch)
                for branch, kind in [('A', 'Политики'), ('A', 'Продажи'), ('B', 'Директора')]]
        with patch('admin_motivation.report_views.choices', return_value=[]), patch('admin_motivation.report_views.load_report', return_value=subdivision_report(rows)):
            self.page.goto(self.live_server_url + reverse('motivation_report', args=['subdivisions']) + '?run=1')
            table = self.page.locator('.subdivision-summary')
            expect(table.locator('tbody tr').first.locator('th').first).to_have_text('Итого')
            expect(table.get_by_role('columnheader', name='Тип коэффициента', exact=True)).to_be_visible()
            branches = table.locator('.subdivision-summary-group')
            expect(branches).to_have_count(2)
            expect(table.locator('.subdivision-summary-type:visible')).to_have_count(0)
            button = branches.nth(0).locator('button')
            button.click()
            expect(button).to_have_attribute('aria-expanded', 'true')
            expect(branches.nth(0).locator('.subdivision-summary-type:visible')).to_have_count(2)
            expect(branches.nth(1).locator('.subdivision-summary-type:visible')).to_have_count(0)
            expect(branches.nth(0).get_by_text('Политики', exact=True)).to_be_visible()
            colors = table.locator('tbody .motivation-potential, tbody .motivation-earned').evaluate_all('(cells) => cells.map(cell => getComputedStyle(cell).backgroundColor)')
            self.assertEqual(colors[0], colors[1])
            headers = table.locator('thead .motivation-potential, thead .motivation-earned').evaluate_all('(cells) => cells.map(cell => getComputedStyle(cell).backgroundColor)')
            self.assertNotEqual(headers[0], headers[1])
            toggle = self.page.locator('#subdivision-summary-potential')
            expect(toggle).to_be_checked()
            toggle.uncheck()
            expect(table.locator('.motivation-potential:visible')).to_have_count(0)
            expect(table.locator('.motivation-earned:visible').first).to_be_visible()
            expect(table.locator('[data-summary-period]').first).to_have_attribute('colspan', '1')
            expect(table.locator('[data-summary-overall]')).to_have_attribute('colspan', '3')
            expect(self.page.locator('#native-division-report .motivation-potential').first).to_be_visible()
            branches.nth(1).locator('button').click()
            expect(table.locator('.motivation-potential:visible')).to_have_count(0)
            toggle.check()
            expect(table.locator('.motivation-potential:visible').first).to_be_visible()
            expect(table.locator('[data-summary-period]').first).to_have_attribute('colspan', '2')
            expect(table.locator('[data-summary-overall]')).to_have_attribute('colspan', '4')
            button.press('Enter')
            expect(branches.nth(0).locator('.subdivision-summary-type:visible')).to_have_count(0)
            self.assertEqual(self.errors, [])

    def test_subdivision_report_checkbox_filters_submit_multiple_values(self):
        from unittest.mock import patch
        from playwright.sync_api import expect
        from .services.reports import FILTERS, subdivision_report
        def options(kind, values, **kwargs):
            return [{'name': name, 'label': caption, 'multiple': multiple,
                     'options': [{'value': value, 'selected': value in values.get(parameter, [])}
                                 for value in ('A', 'B', 'C')]}
                    for name, parameter, caption, dataset, multiple in FILTERS[kind]]
        with patch('admin_motivation.report_views.choices', side_effect=options), patch('admin_motivation.report_views.load_report', return_value=subdivision_report([])) as load:
            self.page.goto(self.live_server_url + reverse('motivation_report', args=['subdivisions']))
            expect(self.page.locator('.manager-multi-field')).to_have_count(5)
            for field_name in ('type', 'planning', 'goods', 'brand', 'subdivision'):
                field = self.page.locator(f'.manager-multi-filter[data-field="{field_name}"]')
                field.locator('summary').click()
                field.locator('[data-multi-all]').click()
                expect(field.locator('input[type=checkbox]:checked')).to_have_count(3)
                field.locator('[data-multi-clear]').click()
                expect(field.locator('input[type=checkbox]:checked')).to_have_count(0)
                field.locator('input[type=search]').fill('B')
                expect(field.locator('.manager-multi-options label:visible')).to_have_count(1)
                field.locator('input[type=search]').fill('')
                field.locator('input[type=checkbox]').nth(0).check()
                field.locator('input[type=checkbox]').nth(1).check()
                if field_name != 'subdivision':
                    field.locator('[data-multi-apply]').click()
            with self.page.expect_navigation():
                self.page.get_by_role('button', name='Сформировать отчёт', exact=True).click()
            values = load.call_args.args[1]
            for key in ('type_coeff', 'Subdivision', 'gp_sales', 'GPGoods', 'Brend'):
                self.assertEqual(values[key], ['A', 'B'])
            self.assertEqual(self.errors, [])

    def test_deals_scrollbar_is_available_above_long_table(self):
        from unittest.mock import patch
        from playwright.sync_api import expect
        payload = {'summary': [], 'totals': {'count': 50, 'shipments': 1000, 'quarterly': 100, 'annual': 20, 'possible': 120},
                   'details': [{'АвторСделки': 'Автор сделки', 'Номер': f'Deal-{index}',
                                'ДатаЗакрытияСделки': date(2025, 1, 10), 'СуммаОтгрузокUSD': 20}
                               for index in range(50)], 'page': 1, 'pages': 1, 'offset': 0}
        with patch('admin_motivation.report_views.choices', return_value=[]), patch('admin_motivation.report_views.load_report', return_value=payload):
            self.page.goto(self.live_server_url + reverse('motivation_report', args=['deals']) + '?run=1&start=2025-01-01&end=2025-12-31&type=Политики&subdivision=A')
            top = self.page.locator('#deal-scroll-top')
            scroll = self.page.locator('#deal-scroll')
            expect(top).to_be_visible()
            self.assertTrue(scroll.evaluate('(el) => el.scrollHeight > el.clientHeight'))
            top.evaluate('(el) => {el.scrollLeft = 200}')
            self.page.wait_for_function('document.getElementById("deal-scroll").scrollLeft > 100')
            top.evaluate('(el) => {el.scrollLeft = 100000}')
            self.page.wait_for_function('Math.abs(document.getElementById("deal-scroll").scrollLeft + document.getElementById("deal-scroll").clientWidth - document.getElementById("deal-scroll").scrollWidth) < 2', timeout=2000)
            scroll.evaluate('(el) => {el.scrollLeft = 50; el.scrollTop = 250}')
            self.page.wait_for_function('Math.abs(document.getElementById("deal-scroll-top").scrollLeft - 50) < 2')
            self.assertEqual(self.page.locator('.deal-table thead th').first.evaluate('(el) => getComputedStyle(el).position'), 'sticky')
            self.page.locator('#deal-detail-toggle').check()
            self.page.wait_for_function('Math.abs(document.querySelector("#deal-scroll-top > div").getBoundingClientRect().width - document.getElementById("deal-scroll").scrollWidth - document.getElementById("deal-scroll-top").clientWidth + document.getElementById("deal-scroll").clientWidth) < 2')
            self.assertEqual(self.errors, [])

    def test_detailed_manager_report_tree_currencies_and_filter_refresh(self):
        from unittest.mock import patch
        from playwright.sync_api import expect
        from .services.reports import FILTERS
        from .services.manager_detail import manager_detail_report
        from .test_manager_detail import detail_rows
        def options(kind, values, **kwargs):
            items = {'type': ['Политики'], 'subdivision': ['A', 'B'], 'author': ['Иванов'],
                     'planning': ['1. RAC'], 'goods': ['Кондиционеры'], 'brand': ['<Brand>'], 'year': ['2026']}
            return [{'name': name, 'label': caption, 'multiple': multiple,
                     'options': [{'value': value, 'selected': value in (values.get(parameter) if isinstance(values.get(parameter), list) else [values.get(parameter)])}
                                 for value in items[name]]}
                    for name, parameter, caption, dataset, multiple in FILTERS[kind]]
        def report(kind, values, *args):
            return manager_detail_report(detail_rows(), values.get('Currency', 'USD'), detail_rows())
        with patch('admin_motivation.report_views.choices', side_effect=options), patch('admin_motivation.report_views.load_report', side_effect=report):
            self.page.goto(self.live_server_url + reverse('motivation_report', args=['manager-detail']) + '?type=Политики&run=1')
            expect(self.page.locator('.report-title')).to_have_text('Мотивация. Детально · 2026 · USD')
            with self.page.expect_download() as downloaded:
                self.page.get_by_role('link', name='Скачать весь отчёт Excel', exact=True).click()
            from openpyxl import load_workbook
            self.assertEqual(downloaded.value.suggested_filename, 'motivation-manager-USD-all.xlsx')
            from io import BytesIO
            with open(downloaded.value.path(), 'rb') as export_file:
                book = load_workbook(BytesIO(export_file.read()))
            self.assertEqual(book.sheetnames, ['Итого по мотивации', 'Расширенная мотивация', 'Параметры отчёта'])
            self.assertEqual(book['Итого по мотивации']['B5'].value, 'Иванов')
            book.close()
            expect(self.page.locator('.report-filter-summary')).to_contain_text('Вид работника: Политики')
            expect(self.page.locator('.report-filter-summary')).to_contain_text('Подразделение: Все')
            expect(self.page.locator('.report-filter-summary')).to_contain_text('Валюта мотивации: USD')
            summary = self.page.locator('.manager-period-summary')
            expect(summary).to_be_visible()
            expect(summary.locator('tbody tr')).to_have_count(3)
            summary_scroll = summary.locator('.report-scroll')
            summary_scroll.evaluate('(el) => {el.style.maxHeight="90px";el.scrollTop=60;}')
            self.assertGreater(summary_scroll.evaluate('(el) => el.scrollTop'), 0)
            self.assertAlmostEqual(summary.locator('thead th').first.bounding_box()['y'], summary_scroll.bounding_box()['y'], delta=2)
            summary_scroll.evaluate('(el) => {el.scrollTop=0;el.style.maxHeight="";}')

            expect(summary.get_by_role('columnheader', name='4 кв.', exact=True)).to_be_visible()
            expect(summary.locator('tbody tr').first.locator('td').last).to_have_text('190')
            expect(self.page.locator('[data-manager-open=USD]')).to_have_count(1)
            expect(self.page.locator('[data-manager-open=RUB]')).to_have_count(0)
            self.assertEqual(self.page.locator('.report-title').evaluate('(heading) => getComputedStyle(heading).backgroundColor'), 'rgb(245, 243, 255)')
            self.page.screenshot(path='/tmp/manager-period-summary.png', full_page=True)
            panel = self.page.locator('#manager-detail-panel')
            self.assertTrue(panel.evaluate('(panel) => panel.open'))
            self.page.locator('[data-manager-open=USD]').click()
            self.assertTrue(panel.evaluate('(panel) => panel.open'))
            execution = self.page.locator('.manager-division-execution')
            expect(execution).to_be_visible()
            self.assertEqual(execution.evaluate('(row) => row.nextElementSibling.classList.contains("report-total-row")'), True)
            expect(execution.locator('.motivation-execution').first).to_have_text('100,00%')
            root = self.page.locator('#manager-detail-report')
            self.assertEqual(root.locator('.report-column-header > .motivation-execution').first.evaluate('(el) => getComputedStyle(el).backgroundColor'), 'rgb(237, 233, 254)')
            expect(root.locator('[data-manager-period-input]').first).to_be_hidden()
            root.locator('[data-manager-reveal="period-input"]').first.click()
            expect(root.locator('[data-manager-period-input]').first).to_be_visible()
            expect(root.locator('.report-column-header > .manager-period-bonus').first).to_be_hidden()
            expect(root.locator('.report-column-header > span').filter(has_text='Доплата за кварталы').last).to_be_visible()
            motivation_totals = root.locator('.report-column-header > span').filter(has_text='Итого мотивация')
            expect(motivation_totals.last).to_be_visible()
            self.assertEqual(root.locator('.report-column-header > span:visible').filter(has_text='Итого мотивация').count(), 1)

            totals_before = root.locator('.report-total-row').inner_text()
            self.page.locator('[data-manager-period-toggle="0"]').uncheck()
            expect(root.locator('.report-period-header > [data-manager-period="0"]')).to_be_hidden()
            expect(root.locator('.report-period-header > strong').last).to_be_visible()
            self.page.locator('[data-manager-period-toggle="0"]').check()
            self.assertEqual(root.locator('.report-total-row').inner_text(), totals_before)
            expect(root.locator('[data-manager-motivation]').first).to_be_hidden()
            root.locator('[data-manager-reveal=motivation]').first.click()
            expect(root.locator('[data-manager-motivation]').first).to_be_visible()
            expect(root.locator('[data-manager-segments]').first).to_be_hidden()
            root.locator('[data-manager-reveal=segments]').first.click()
            expect(root.locator('[data-manager-segments]').first).to_be_visible()
            root.locator('[data-manager-reveal=motivation]').first.click()
            expect(root.locator('[data-manager-motivation]').first).to_be_hidden()
            expect(self.page.locator('#manager-detail-nodes').get_by_text('Иванов', exact=True)).to_be_visible()
            expect(self.page.locator('#manager-detail-nodes').get_by_text('1. RAC', exact=True)).to_be_visible()
            expect(self.page.locator('#manager-detail-nodes').get_by_text('Кондиционеры', exact=True)).to_have_count(0)
            self.assertEqual(self.page.locator('#manager-detail-nodes > details[open]').count(), 1)
            self.assertGreater(self.page.locator('#manager-detail-nodes details:not([open])').count(), 0)
            self.page.locator('#manager-detail-collapse').click()
            expect(self.page.locator('#manager-detail-nodes').get_by_text('Иванов', exact=True)).to_be_hidden()
            self.page.locator('#manager-detail-expand').click()
            expect(self.page.locator('#manager-detail-nodes').get_by_text('Иванов', exact=True)).to_be_visible()
            expect(self.page.locator('#manager-detail-nodes').get_by_text('<Brand>', exact=True)).to_be_visible()
            viewport = root.locator('..')
            top_scroll = self.page.locator('#manager-detail-scroll-top')
            expect(top_scroll).to_be_visible()
            top_scroll.evaluate('(el) => el.scrollLeft = 300')
            self.page.wait_for_function('Math.abs(document.getElementById("manager-detail-scroll").scrollLeft - 300) < 2')
            viewport.evaluate('(el) => el.scrollLeft = 150')
            self.page.wait_for_function('Math.abs(document.getElementById("manager-detail-scroll-top").scrollLeft - 150) < 2')
            viewport.evaluate('(el) => el.scrollTop = 150')
            self.assertAlmostEqual(root.locator('.manager-sticky-header').bounding_box()['y'], viewport.bounding_box()['y'], delta=2)
            viewport.evaluate('(el) => el.scrollTop = 0')
            viewport.evaluate('(el) => el.scrollLeft = 0')
            label = root.locator('#manager-detail-nodes > details > summary > .native-label').first
            label.scroll_into_view_if_needed()
            viewport.evaluate('(el) => el.scrollLeft = 0')
            initial_left = label.bounding_box()['x']
            initial_metric_left = root.locator('.report-total-row > span:not([hidden])').first.bounding_box()['x']
            viewport.evaluate('(el) => el.scrollLeft = 550')
            self.assertAlmostEqual(label.bounding_box()['x'], initial_left, delta=2)
            self.assertLess(root.locator('.report-total-row > span:not([hidden])').first.bounding_box()['x'], initial_metric_left - 400)
            self.assertTrue(label.evaluate('(el) => {const r=el.getBoundingClientRect(); return el.contains(document.elementFromPoint(r.left+30,r.top+r.height/2));}'))
            viewport.evaluate('(el) => el.scrollLeft = 0')

            expect(self.page.locator('#manager-detail-report [data-manager-extra]').first).to_be_hidden()
            self.page.locator('#manager-detail-extra').check()
            expect(self.page.locator('#manager-detail-report [data-manager-extra]').first).to_be_visible()
            expect(self.page.locator('#manager-detail-nodes > details > summary > span').nth(1)).to_have_text('1,50%')
            self.page.locator('#manager-detail-collapse').click()
            expect(self.page.locator('#manager-detail-nodes').get_by_text('Иванов', exact=True)).to_be_hidden()
            self.page.locator('.report-filter-panel > summary').click()
            with self.page.expect_navigation():
                self.choose_filter('select[name=currency]', 'RUB')
            expect(self.page.locator('.report-title')).to_have_text('Мотивация. Детально · 2026 · RUB')
            expect(self.page.locator('[data-manager-open=USD]')).to_have_count(0)
            expect(self.page.locator('[data-manager-open=RUB]')).to_have_count(1)
            expect(summary.locator('tbody tr')).to_have_count(3)
            self.assertEqual(summary.locator('tbody tr td').last.text_content().replace('\xa0', '').replace(' ', ''), '19000')
            self.assertTrue(panel.evaluate('(panel) => panel.open'))
            panel.locator(':scope > summary').click()
            self.assertFalse(panel.evaluate('(panel) => panel.open'))
            self.page.locator('.report-filter-panel > summary').click()
            authors = self.page.locator('.manager-multi-filter[data-field=author]')
            authors.locator('summary').click()
            authors.locator('input[type=checkbox]').check()
            authors.locator('[data-multi-apply]').click()
            branches = self.page.locator('.manager-multi-filter[data-field=subdivision]')
            branches.locator('summary').click()
            expect(branches.locator('input[type=search]')).to_be_focused()
            branches.locator('input[type=search]').fill('A')
            expect(branches.locator('.manager-multi-options label:visible')).to_have_count(1)
            branches.locator('[data-multi-all]').click()
            expect(branches.locator('input[type=checkbox]:checked')).to_have_count(2)
            branches.locator('[data-multi-clear]').click()
            expect(branches.locator('input[type=checkbox]:checked')).to_have_count(0)
            branches.locator('[data-multi-all]').click()
            with self.page.expect_navigation():
                branches.locator('[data-multi-apply]').click()
            self.assertEqual(self.page.locator('select[name=author]').evaluate('(select) => [...select.selectedOptions].map(option => option.value)'), [])
            with self.page.expect_navigation():
                self.choose_filter('select[name=currency]', 'USD')
            expect(self.page.locator('.report-title')).to_have_text('Мотивация. Детально · 2026 · USD')
        self.assertEqual(self.errors, [])

    def test_manager_authors_refresh_automatically_with_type_and_branch(self):
        from unittest.mock import patch
        from playwright.sync_api import expect
        from .test_reports import NativeMotivationReportTests
        from .services.reports import manager_report
        def options(kind, values, **kwargs):
            groups = [('type', 'TypeCoeff_', ['Политики', 'Директора']),
                      ('subdivision', 'Subdivision', ['A', 'B']),
                      ('author', 'DealAuthor', [values['TypeCoeff_'] + '-' + values['Subdivision']] if values.get('TypeCoeff_') and values.get('Subdivision') else []),
                      ('quarter', 'qm', ['1 Квартал'])]
            return [{'name': name, 'label': name, 'multiple': False, 'options': [{'value': value, 'selected': value == values.get(key)} for value in items]} for name, key, items in groups]
        data = manager_report(NativeMotivationReportTests.manager_rows())
        data['years'] = ['2024']
        with patch('admin_motivation.report_views.choices', side_effect=options), patch('admin_motivation.report_views.load_report', return_value=data):
            self.page.goto(self.live_server_url + reverse('motivation_report', args=['manager']))
            with self.page.expect_navigation():
                self.choose_filter('select[name=type]', 'Директора')
            with self.page.expect_navigation():
                self.choose_filter('select[name=subdivision]', 'A')
            expect(self.page.locator('select[name=author] option')).to_have_count(2)
            self.choose_filter('select[name=author]', 'Директора-A')
            with self.page.expect_navigation():
                self.choose_filter('select[name=subdivision]', 'B')
            expect(self.page.locator('select[name=author]')).to_have_value('')
            expect(self.page.locator('select[name=author] option').last).to_have_text('Директора-B')
            self.choose_filter('select[name=author]', 'Директора-B')
            self.choose_filter('select[name=quarter]', '1 Квартал')
            with self.page.expect_navigation():
                self.page.get_by_role('button', name='Сформировать отчёт', exact=True).click()
            expect(self.page.locator('.report-title')).to_have_text('Мотивация. Дашборд · 2024 · USD')
            dashboard = self.page.locator('#manager-dashboard')
            for key in ['year-possible', 'year-available']:
                expect(dashboard.locator('[data-dashboard-column="' + key + '"]').first).to_be_hidden()
                dashboard.locator('[data-dashboard-reveal="' + key + '"]').first.click()
                expect(dashboard.locator('[data-dashboard-column="' + key + '"]').first).to_be_visible()
                self.assertEqual(dashboard.locator('.report-total-row > [data-dashboard-column="' + key + '"]').count(), 5)
                dashboard.locator('[data-dashboard-reveal="' + key + '"]').first.click()
            total_before = dashboard.locator('.report-total-row').inner_text()
            expect(dashboard.locator('[data-dashboard-column="quarter-motivation"]').first).to_be_hidden()
            dashboard.locator('[data-dashboard-reveal="quarter-motivation"]').first.click()
            expect(dashboard.locator('[data-dashboard-column="quarter-motivation"]').first).to_be_visible()
            expect(dashboard.locator('[data-dashboard-column="quarter-period"]').first).to_be_hidden()
            dashboard.locator('[data-dashboard-reveal="quarter-period"]').first.click()
            expect(dashboard.locator('[data-dashboard-column="quarter-period"]').first).to_be_visible()
            expect(dashboard.locator('[data-dashboard-column="year-period"]').first).to_be_hidden()
            dashboard.locator('[data-dashboard-reveal="quarter-motivation"]').first.click()
            dashboard.locator('[data-dashboard-reveal="quarter-period"]').first.click()
            self.assertEqual(dashboard.locator('.report-total-row').inner_text(), total_before)

            if self.page.locator('.report-filter-panel').get_attribute('open') is None:
                self.page.locator('.report-filter-panel > summary').click()
            self.choose_filter('select[name=currency]', 'RUB')
            with self.page.expect_navigation():
                self.page.get_by_role('button', name='Сформировать отчёт', exact=True).click()
            expect(self.page.locator('.report-title')).to_have_text('Мотивация. Дашборд · 2024 · RUB')
        self.assertEqual(self.errors, [])

    def test_current_coefficient_templates_download_for_global_and_subdivision(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from .test_sales_plans import SalesPlanTests
        fixture = SalesPlanTests()
        self.db(fixture.setUp)
        for route, args, filename in [('global_coeff_admin_motivation', [], 'global'),
                ('coeff_subdivisions_admin_motivation', [fixture.subs[0].pk], f'subdivision-{fixture.subs[0].pk}')]:
            self.page.goto(self.live_server_url + reverse(route, args=args) + f'?version={fixture.version.pk}')
            self.page.get_by_role('button',name='Импорт и справочники',exact=True).click()
            with self.page.expect_download() as download_info:
                self.page.get_by_role('link', name='Скачать шаблон', exact=False).click()
            download = download_info.value
            self.assertEqual(download.suggested_filename, f'motivation-coefficients-{filename}.xlsx')
            with open(download.path(), 'rb') as stream:
                book = load_workbook(BytesIO(stream.read()))
            rows = list(book.active.iter_rows(min_row=3, values_only=True))
            target = next(row for row in rows if row[1:4] == ('Sales', '*', '*'))
            self.assertIn(.1 if not args else .2, target[4:])
            book.close()
        self.assertEqual(self.errors, [])

    def test_plan_summary_expansion_and_independent_column_toggles(self):
        from playwright.sync_api import expect
        from .test_sales_plans import SalesPlanTests
        fixture = SalesPlanTests()
        self.db(fixture.setUp)
        self.db(fixture.load)
        from .models import DirectorEmployee, GlobalDirectorCoeff, SubdivisionDirectorCoeff
        def directors():
            fixture.scenario.lines.filter(subdivision='A').update(amount_ddp_usd=600)
            fixture.scenario.lines.filter(subdivision='B').update(amount_ddp_usd=300)
            employee = DirectorEmployee.objects.create(full_name='Иванов Иван')
            GlobalDirectorCoeff.objects.create(version=fixture.version, goods=fixture.good,
                **{f'k{i}': '.01' for i in range(5)})
            SubdivisionDirectorCoeff.objects.create(version=fixture.version, goods=fixture.good,
                subdivision=fixture.subs[0], employee=employee, **{f'k{i}': '.02' for i in range(5)})
            return employee.pk
        employee_id = self.db(directors)
        self.page.goto(self.live_server_url + reverse('motivation_reports'))
        self.page.get_by_role('link', name='Сводный отчёт по плану продаж', exact=False).click()
        self.page.locator('#summary-plan').select_option(str(fixture.scenario.pk))
        self.page.get_by_role('button', name='Показать', exact=True).click()
        report = self.page.locator('#plan-summary')
        self.assertEqual(report.locator('details[open]').count(), 0)
        expect(report.locator('.summary-company-total > summary .summary-difference')).to_have_text('+130')
        expect(report.locator('.summary-company-total > summary .summary-global.summary-markup')).to_have_text('1,32')
        expect(report.locator('.summary-company-total > summary .summary-local.summary-markup')).to_have_text('1,19')
        expect(report.locator('.summary-company-total > summary .summary-markup-difference')).to_have_text('-0,14')
        report.locator('summary').first.click()
        expect(report.get_by_text('Политика', exact=True).first).to_be_visible()
        report.locator('[data-summary-side=global]').uncheck()
        expect(report.locator('.summary-heading .summary-global').first).to_be_hidden()
        expect(report.locator('.summary-heading .summary-local').first).to_be_visible()
        report.locator('[data-summary-side=local]').uncheck()
        expect(report.locator('.summary-heading .summary-local').first).to_be_hidden()
        expect(report.locator('.summary-company-total > summary .summary-difference')).to_be_visible()
        expect(report.locator('.summary-heading > span').nth(1)).to_be_visible()
        report.locator('[data-summary-side=global]').check()
        expect(report.locator('.summary-heading .summary-global').first).to_be_visible()
        report.locator('[data-summary-side=local]').check()
        report.locator('[data-summary-expand]').click()
        self.assertEqual(report.locator('details[open]').count(), report.locator('details').count())
        expect(report.get_by_text('Иванов Иван', exact=True).first).to_be_visible()
        self.page.set_viewport_size({'width': 1280, 'height': 600})
        viewport = report.locator('.summary-scroll')
        for company_open in [True, False]:
            report.locator('.summary-company-total').evaluate('(node, open) => node.open = open', company_open)
            viewport.evaluate('node => node.scrollTop = 200')
            self.page.evaluate('() => new Promise(requestAnimationFrame)')
            self.assertGreater(viewport.evaluate('node => node.scrollTop'), 0)
            heading_box = report.locator('.summary-heading').bounding_box()
            total_box = report.locator('.summary-company-total > summary').bounding_box()
            self.assertAlmostEqual(heading_box['y'], viewport.bounding_box()['y'], delta=1)
            self.assertAlmostEqual(total_box['y'], heading_box['y'] + heading_box['height'], delta=1)
        report.locator('.summary-company-total').evaluate('node => node.open = true')
        viewport.evaluate('node => {node.scrollTop = 0; node.scrollLeft = 150;}')
        self.assertAlmostEqual(report.locator('.summary-heading').bounding_box()['x'], report.locator('.summary-company-total > summary').bounding_box()['x'], delta=1)
        viewport.evaluate('node => node.scrollLeft = 0')
        self.page.screenshot(path='/tmp/plan-summary-desktop.png', full_page=True)
        self.page.set_viewport_size({'width': 480, 'height': 850})
        self.page.screenshot(path='/tmp/plan-summary-mobile.png', full_page=True)
        report.locator('[data-summary-collapse]').click()
        self.assertEqual(report.locator('details[open]').count(), 0)
        self.page.set_viewport_size({'width': 1280, 'height': 900})
        for field, values in [('subdivision', ['A', 'B']), ('type', ['policies', 'directors']),
                              ('planning_group_sales', ['Sales']), ('director', [str(employee_id)])]:
            dropdown = self.page.locator(f'[data-summary-filter="{field}"]')
            dropdown.locator('summary').click()
            for value in values:
                dropdown.locator(f'input[type=checkbox][value="{value}"]').check()
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Показать', exact=True).click()
        total = report.locator('.summary-company-total > summary')
        expect(total.locator('.summary-global').first).to_have_text('110')
        expect(total.locator('.summary-local').first).to_have_text('220')
        self.assertEqual(report.locator('.summary-table > details').count(), 2)
        expect(self.page.locator('[data-summary-filter=type] [data-filter-caption]')).to_have_text('Выбрано: 2')
        with self.page.expect_navigation():
            self.page.get_by_role('link', name='Сбросить фильтры', exact=True).click()
        self.assertEqual(report.locator('.summary-table > details').count(), 3)
        self.assertEqual(self.errors, [])

    def test_comparison_directors_percent_and_collapsed_breakdown(self):
        from playwright.sync_api import expect
        from .test_sales_plans import SalesPlanTests
        from .models import GlobalDirectorCoeff
        fixture = SalesPlanTests()
        def setup():
            fixture.setUp()
            fixture.load()
            GlobalDirectorCoeff.objects.create(version=fixture.version, goods=fixture.good,
                **{f'k{i}': '.01' for i in range(5)})
            return fixture.scenario.pk, fixture.good.pk
        plan, good = self.db(setup)
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation'))
        self.page.locator('#comparison-plan').select_option(str(plan))
        expect(self.page.locator('[data-amount="current.total"]')).to_have_text('240')
        expect(self.page.locator('[data-amount="current.plan_percent"]')).to_have_text('16%')
        expect(self.page.locator('[data-motivation-breakdown]').first).not_to_be_visible()
        self.assertEqual(self.page.locator('[data-amount="current.ddp"]').count(), 0)
        self.page.locator('#comparison-breakdown-toggle').click()
        expect(self.page.locator('[data-amount="current.directors"]')).to_have_text('15')
        self.page.locator(f'input[name="director_coeff={good}=0"]').fill('2%')
        self.page.locator(f'input[name="director_coeff={good}=0"]').press('Tab')
        expect(self.page.locator('[data-amount="current.directors"]')).to_have_text('17')
        expect(self.page.locator('[data-amount="current.total"]')).to_have_text('242')
        self.page.locator('#comparison-breakdown-toggle').click()
        expect(self.page.locator('[data-amount="current.directors"]')).not_to_be_visible()
        self.assertEqual(self.errors, [])

    def test_global_copy_selects_employees_per_branch(self):
        from .models import DirectorEmployee, GlobalDirectorCoeff, SubdivisionDirectorCoeff
        from playwright.sync_api import expect
        def fixture():
            channel = Chanel.objects.create(chanel_name='Recipients')
            branches = [Subdivision.objects.create(subdivision_key=key, subdivision_name=key, subdivision_global=key, chanel=channel) for key in ['Recipient A', 'Recipient B']]
            employees = [DirectorEmployee.objects.create(full_name=name) for name in ['Первый', 'Второй']]
            for employee in employees:
                employee.subdivisions.add(*branches)
            branches.append(Subdivision.objects.create(subdivision_key='Recipient C empty', subdivision_name='Recipient C empty', subdivision_global='Recipient C empty', chanel=channel))
            GlobalDirectorCoeff.objects.create(version=self.first, goods=self.coeff.goods, k0='.01')
            from django.db import models
            models.QuerySet(model=GlobalCoeffVersion).filter(pk=self.first.pk).update(status='approved')
            return [item.pk for item in branches], [item.pk for item in employees]
        branches, employees = self.db(fixture)
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation'))
        self.page.locator('#subdivisionDropdownBtn').click()
        self.page.locator('[name=copy_groups][value=policies]').uncheck()
        self.page.locator('[name=copy_groups][value=sales]').uncheck()
        self.page.locator('[name=copy_groups][value=directors]').check()
        self.page.locator('#copy-branch-search').fill('Recipient A')
        expect(self.page.locator('.copy-branch-row').nth(1)).not_to_be_visible()
        row = self.page.locator('.copy-branch-row').nth(0)
        row.locator('[name=selected_subdivisions]').check()
        row.locator('summary').click()
        expect(row.locator('.copy-employee').nth(0)).to_be_checked()
        expect(row.locator('.copy-employee').nth(1)).to_be_checked()
        row.locator('.copy-all-employees').uncheck()
        expect(row.locator('.copy-employee').nth(0)).not_to_be_checked()
        row.locator('.copy-all-employees').check()
        row.locator(f'.copy-employee[value="{employees[1]}"]').uncheck()
        self.page.locator('#copy-branch-search').fill('')
        self.page.locator(f'[name=selected_subdivisions][value="{branches[2]}"]').check()
        self.page.on('dialog', lambda dialog: dialog.accept())
        with self.page.expect_navigation():
            self.page.locator('[name=action_button][value=apply_sub]').click()
        self.assertEqual(self.db(lambda: list(SubdivisionDirectorCoeff.objects.values_list('subdivision_id', 'employee_id'))), [(branches[0], employees[0])])
        self.assertEqual(self.errors, [])

    def setUp(self):
        authorize_test_case(self)
        self.first, _ = GlobalCoeffVersion.objects.get_or_create(effective_from=date(2001, 1, 1))
        good = Goods.objects.create(goods_key='Key', planning_group_sales='Sales', group='Group', brand='Brand')
        segment = SegmentCoeff.objects.create(segment_name='Segment')
        kind = TypeCoeff.objects.create(type_coeff_name='Политики')
        variation = VariationCalculate.objects.create(variation_name='Level')
        self.coeff = GlobalCoeff.objects.create(version=self.first, goods=good, type_coeff=kind, segment=segment,
            motivation_coeff=0.123456, manager_coeff=1, variation_calculate=variation)
        from playwright.sync_api import sync_playwright
        self.driver = sync_playwright().start()
        self.addCleanup(self.driver.stop)
        self.browser = self.driver.chromium.launch(executable_path=os.environ.get('PLAYWRIGHT_CHROMIUM_EXECUTABLE'), headless=True)
        self.addCleanup(self.browser.close)
        self.page = self.browser.new_page()
        from django.conf import settings
        self.page.context.add_cookies([{'name': settings.SESSION_COOKIE_NAME, 'value': self.client.cookies[settings.SESSION_COOKIE_NAME].value, 'url': self.live_server_url}])
        self.errors = []
        self.page.on('pageerror', lambda error: self.errors.append(str(error)))
        self.pool = ThreadPoolExecutor(max_workers=1)
        self.addCleanup(self.pool.shutdown)

    def db(self, action):
        return self.pool.submit(action).result()

    def test_approval_dialog_opens_before_slow_review_and_blocks_duplicate_requests(self):
        from playwright.sync_api import expect
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation'))
        pending = []
        self.page.route('**' + reverse('motivation_coefficient_comparison'), lambda route: pending.append(route))
        self.page.locator('#approval-review').click()
        expect(self.page.locator('#approval-dialog')).to_be_visible()
        expect(self.page.locator('#approval-description')).to_contain_text('Проверяем версию')
        expect(self.page.locator('#approval-commit')).to_be_disabled()
        self.page.locator('#approval-review').evaluate('(button) => button.click()')
        self.assertEqual(len(pending), 1)
        pending[0].fulfill(status=400, content_type='application/json', body='{"error":"Проверка не завершена"}')
        expect(self.page.locator('#approval-error')).to_have_text('Проверка не завершена')
        expect(self.page.locator('#approval-commit')).to_be_disabled()
        self.page.locator('#approval-cancel').click()
        expect(self.page.locator('#approval-dialog')).not_to_be_visible()
        self.assertEqual(self.errors, [])

    def test_coefficient_table_fits_desktop_working_area(self):
        self.db(lambda: [SegmentCoeff.objects.get_or_create(segment_name=f'K{i}') for i in range(5)])
        for width in (1280, 1536):
            self.page.set_viewport_size({'width': width, 'height': 900})
            self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation'))
            dimensions = self.page.locator('#coefficient-table').evaluate(
                '(table) => ({table: table.getBoundingClientRect().width, wrapper: table.parentElement.clientWidth, scroll: table.parentElement.scrollWidth})')
            self.assertLessEqual(dimensions['table'], dimensions['wrapper'] + 1)
            self.assertLessEqual(dimensions['scroll'], dimensions['wrapper'] + 1)
        self.assertEqual(self.errors, [])

    def test_delete_draft_confirmation_and_return_to_previous_version(self):
        from playwright.sync_api import expect
        from .services.versions import create_version
        draft = self.db(lambda: create_version(date(2026, 1, 1), source_version=self.first.pk).pk)
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation') + f'?version={draft}')
        self.page.once('dialog', lambda dialog: dialog.dismiss())
        self.page.locator('button[value="delete_version"]').click()
        self.assertTrue(self.db(lambda: GlobalCoeffVersion.objects.filter(pk=draft).exists()))
        self.page.once('dialog', lambda dialog: dialog.accept())
        with self.page.expect_navigation():
            self.page.locator('button[value="delete_version"]').click()
        self.assertFalse(self.db(lambda: GlobalCoeffVersion.objects.filter(pk=draft).exists()))
        expect(self.page.locator('#coefficient-version')).to_have_value(str(self.first.pk))
        expect(self.page.locator('button[value="delete_version"]')).to_have_count(0)
        self.assertEqual(self.errors, [])

    def test_comparison_settings_collapse_keeps_results_and_remembers_state(self):
        from playwright.sync_api import expect
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation'))
        self.page.locator('#comparison-settings > summary').click()
        expect(self.page.locator('#comparison-plan')).not_to_be_visible()
        expect(self.page.locator('#comparison-panel .compare-table')).to_be_visible()
        self.page.reload()
        expect(self.page.locator('#comparison-plan')).not_to_be_visible()
        expect(self.page.locator('#comparison-panel .compare-table')).to_be_visible()
        self.page.locator('#comparison-settings > summary').click()
        expect(self.page.locator('#comparison-plan')).to_be_visible()
        self.assertEqual(self.errors, [])

    def choose_filter(self, selector, value):
        select = self.page.locator(selector)
        text = select.locator('option').evaluate_all('(options, value) => options.find(option => option.value === value).textContent', value)
        root = select.locator('xpath=following-sibling::div[1]')
        root.locator(':scope > button').click()
        root.locator('input[type=search]').fill(text)
        root.locator('.searchable-filter-options button').get_by_text(text, exact=True).click()

    def test_create_edit_and_read_history_without_losing_percentage_precision(self):
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation'))
        self.assertEqual(self.page.locator('#mainCoeffForm input[name^="global_coeff="]').input_value(), '12.3456%')
        self.page.locator('.version-tools button[aria-controls="create-version-details"]').click()
        self.page.locator('#id_effective_from').fill('2027-01-01')
        self.page.locator('#id_title').fill('January')
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Создать новую версию', exact=True).click()
        latest = self.db(lambda: GlobalCoeffVersion.objects.order_by('-effective_from').first())
        self.assertEqual(self.page.locator('#coefficient-version').input_value(), str(latest.pk))
        self.page.locator('#mainCoeffForm input[name^="global_coeff="]').fill('0.25')
        self.assertEqual(self.page.locator('#mainCoeffForm input[name^="global_coeff="]').input_value(), '0.25')
        self.page.locator('#mainCoeffForm input[name^="global_coeff="]').press('Tab')
        self.assertEqual(self.page.locator('#mainCoeffForm input[name^="global_coeff="]').input_value(), '25%')
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Сохранить изменения', exact=True).click()
        self.assertEqual(self.db(lambda: latest.coefficients.get().motivation_coeff), 0.25, self.page.locator('#version-panel').inner_text())
        self.assertEqual(self.db(lambda: GlobalCoeff.objects.get(pk=self.coeff.pk).motivation_coeff), 0.123456)
        self.page.locator('#coefficient-version').select_option(str(self.first.pk))
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Открыть', exact=True).click()
        self.assertTrue(self.page.get_by_role('button', name='Сохранить изменения', exact=True).is_disabled())
        self.assertEqual(self.page.locator('#mainCoeffForm input[name^="global_coeff="]').input_value(), '12.3456%')
        self.page.locator('.version-tools button[aria-controls="version-history"]').click()
        self.assertIn('Не утверждена', self.page.locator('#version-history').inner_text())
        self.assertEqual(self.errors, [])

    def test_shared_version_manager_and_subdivision_edits_leave_history_intact(self):
        def setup_subdivision():
            channel = Chanel.objects.create(chanel_name='Channel')
            sub = Subdivision.objects.create(subdivision_key='A', subdivision_name='A', subdivision_global='A', chanel=channel)
            kind = KindManagerCoeff.objects.create(name='Year')
            global_coeff = GlobalCoeff.objects.get(pk=self.coeff.pk)
            SubdivisionCoeff.objects.create(version=self.first, subdivision=sub, goods=global_coeff.goods,
                type_coeff=global_coeff.type_coeff, segment=global_coeff.segment, variation_calculate=global_coeff.variation_calculate,
                motivation_coeff=0.15, manager_coeff=1)
            SubdivisionManagerCoeff.objects.create(version=self.first, subdivision=sub, kind=kind, coeff=0.8)
            return sub.pk
        sub_id = self.db(setup_subdivision)
        self.page.goto(self.live_server_url + reverse('coeff_subdivisions_admin_motivation', args=[sub_id]))
        self.page.locator('.version-tools button[aria-controls="create-version-details"]').click()
        self.page.locator('#id_effective_from').fill('2027-01-01')
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Создать новую версию', exact=True).click()
        latest = self.db(lambda: GlobalCoeffVersion.objects.order_by('-effective_from').first())
        self.page.goto(self.live_server_url + reverse('coeff_subdivisions_admin_motivation', args=[sub_id]))
        self.assertEqual(self.page.locator('#coefficient-version').input_value(), str(latest.pk))
        self.page.locator('#manager-details summary').click()
        self.page.locator('input[name^="kind_coeff="]').fill('0.9')
        self.page.locator('input[name^="kind_coeff="]').press('Tab')
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Сохранить руководителя', exact=True).click()
        self.page.locator('input[name^="sub_coeff="]').fill('0.25')
        self.page.locator('input[name^="sub_coeff="]').press('Tab')
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Сохранить изменения', exact=True).click()
        self.assertEqual(self.db(lambda: latest.manager_coefficients.get().coeff), 0.9)
        self.assertEqual(self.db(lambda: latest.subdivision_coefficients.get().motivation_coeff), 0.25)
        self.assertEqual(self.db(lambda: self.first.manager_coefficients.get().coeff), 0.8)
        self.assertEqual(self.db(lambda: self.first.subdivision_coefficients.get().motivation_coeff), 0.15)
        self.page.locator('#coefficient-version').select_option(str(self.first.pk))
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Открыть', exact=True).click()
        self.page.locator('#manager-details summary').click()
        self.assertTrue(self.page.get_by_role('button', name='Сохранить руководителя', exact=True).is_disabled())
        self.assertTrue(self.page.get_by_role('button', name='Сохранить изменения', exact=True).is_disabled())
        self.assertEqual(self.errors, [])

    def test_sales_plan_create_load_calculate_and_change_scope(self):
        from unittest.mock import patch
        from .test_sales_plans import SalesPlanTests
        from .models import SalesPlanScenario
        fixtures = SalesPlanTests()
        self.db(fixtures.setUp)
        self.page.goto(self.live_server_url + reverse('motivation_sales_plans'))
        self.page.get_by_text('Создать сценарий', exact=True).click()
        self.page.locator('#id_title').fill('Browser plan')
        self.page.locator('#id_source_version').fill('2027')
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Создать', exact=True).click()
        scenario = self.db(lambda: SalesPlanScenario.objects.get(title='Browser plan'))
        with patch('admin_motivation.services.sales_plans.connect_database', fixtures.connector):
            with self.page.expect_navigation():
                self.page.get_by_role('button', name='Загрузить план из MS SQL', exact=True).click()
            with self.page.expect_navigation():
                self.page.get_by_role('button', name='Рассчитать мотивацию', exact=True).click()
        self.assertEqual(self.db(lambda: scenario.lines.get(subdivision='').total_usd), 225)
        self.assertEqual(self.page.locator('.report-node[data-level="0"]').count(), 1)
        self.page.get_by_role('button', name='Развернуть всё', exact=True).click()
        self.assertEqual(self.page.locator('.report-leaf').count(), 1)
        with self.page.expect_navigation():
            self.page.get_by_role('link', name='По подразделениям', exact=True).click()
        self.assertEqual(self.page.locator('.report-node[data-level="0"]').count(), 2)
        self.assertEqual(self.page.locator('.report-leaf').count(), 0)
        self.page.get_by_role('button', name='Развернуть всё', exact=True).click()
        self.assertEqual(self.page.locator('.report-leaf').count(), 2)
        with self.page.expect_download() as download:
            self.page.get_by_role('link', name='Скачать CSV', exact=True).click()
        self.assertIn('subdivisions.csv', download.value.suggested_filename)
        self.assertEqual(self.errors, [])

    def test_calculation_progress_is_visible_and_duplicate_click_is_blocked(self):
        from threading import Event
        from unittest.mock import patch
        from playwright.sync_api import expect
        from .test_sales_plans import SalesPlanTests
        from .services.sales_plans import calculate_plan
        fixtures = SalesPlanTests()
        self.db(fixtures.setUp)
        self.db(fixtures.load)
        release = Event()
        def slow_calculation(scenario_id, progress):
            progress({'stage': 'Расчёт тестового плана', 'percent': 42, 'processed': 1, 'total': 3})
            if not release.wait(10):
                raise RuntimeError('Progress test timed out')
            return calculate_plan(scenario_id, progress=progress)
        url = self.live_server_url + reverse('motivation_sales_plan', args=[fixtures.scenario.pk])
        self.page.goto(url)
        with patch('admin_motivation.plan_views.calculate_plan', side_effect=slow_calculation):
            try:
                self.page.get_by_role('button', name='Рассчитать мотивацию', exact=True).click()
                expect(self.page.locator('#plan-progress')).to_be_visible()
                expect(self.page.locator('#plan-progress-bar')).to_have_attribute('aria-valuenow', '42')
                expect(self.page.locator('#plan-progress-count')).to_have_text('Обработано 1 из 3 строк')
                self.assertTrue(self.page.get_by_role('button', name='Рассчитать мотивацию', exact=True).is_disabled())
                self.assertTrue(self.page.get_by_role('button', name='Загрузить план из MS SQL', exact=True).is_disabled())
                response = self.page.request.post(url, form={'action': 'calculate', 'csrfmiddlewaretoken': self.page.locator('input[name=csrfmiddlewaretoken]').first.input_value()}, headers={'Accept': 'application/x-ndjson', 'Referer': url})
                self.assertIn('уже выполняется', response.text())
                with self.page.expect_navigation():
                    release.set()
                expect(self.page.locator('#plan-progress-bar')).to_have_attribute('aria-valuenow', '100')
                self.assertEqual(self.db(lambda: fixtures.scenario.lines.get(subdivision='').total_usd), 225)
            finally:
                release.set()
        self.assertEqual(self.errors, [])

    def test_calculation_error_is_visible_and_buttons_are_reenabled(self):
        from playwright.sync_api import expect
        from .test_sales_plans import SalesPlanTests
        from .services.sales_plans import calculate_plan
        fixtures = SalesPlanTests()
        self.db(fixtures.setUp)
        self.db(fixtures.load)
        self.db(lambda: calculate_plan(fixtures.scenario.pk))
        self.db(lambda: fixtures.version.coefficients.filter(type_coeff__type_coeff_name='Политики').delete())
        self.page.goto(self.live_server_url + reverse('motivation_sales_plan', args=[fixtures.scenario.pk]))
        self.page.get_by_role('button', name='Рассчитать мотивацию', exact=True).click()
        expect(self.page.locator('#plan-progress-error')).to_be_visible()
        expect(self.page.locator('#plan-progress-error')).to_contain_text('Нет коэффициента')
        expect(self.page.get_by_role('button', name='Рассчитать мотивацию', exact=True)).to_be_enabled()
        self.assertEqual(self.db(lambda: fixtures.scenario.lines.get(subdivision='').total_usd), 225)
        self.assertEqual(self.errors, [])

    def test_period_tree_default_expansion_subtotals_filters_and_csv(self):
        from playwright.sync_api import expect
        from .test_sales_plans import SalesPlanTests
        from .services.sales_plans import calculate_plan
        fixtures = SalesPlanTests()
        self.db(fixtures.setUp)
        self.db(fixtures.load)
        self.db(lambda: fixtures.scenario.lines.filter(subdivision='').update(amount_ddp_usd=900))
        self.db(lambda: calculate_plan(fixtures.scenario.pk))
        url = self.live_server_url + reverse('motivation_sales_plan', args=[fixtures.scenario.pk])
        self.page.goto(url)
        expect(self.page.locator('.report-node[data-level="0"]')).to_have_attribute('open', '')
        self.assertEqual(self.page.locator('.report-node[data-level="1"][open]').count(), 0)
        self.assertFalse(self.page.locator('.report-leaf').is_visible())
        self.page.locator('.report-node[data-level="1"] > summary').click()
        self.page.locator('.report-node[data-level="2"] > summary').click()
        expect(self.page.locator('.report-leaf')).to_be_visible()
        expect(self.page.locator('.report-leaf')).to_contain_text('225')
        for selector in ['.report-node[data-level="0"] > summary', '.report-node[data-level="1"] > summary', '.report-node[data-level="2"] > summary', '.report-leaf']:
            cells = self.page.locator(selector).locator('span')
            self.assertEqual(cells.count(), 8)
            expect(cells.nth(4)).to_have_text('0')
            expect(cells.nth(6)).to_have_text('900')
            expect(cells.nth(7)).to_have_text('1,33')
        self.page.get_by_role('button', name='Свернуть всё', exact=True).click()
        self.assertEqual(self.page.locator('#period-report details[open]').count(), 0)
        self.page.get_by_role('button', name='Развернуть всё', exact=True).click()
        self.assertEqual(self.page.locator('#period-report details[open]').count(), 3)
        with self.page.expect_navigation():
            self.page.get_by_role('link', name='По подразделениям', exact=True).click()
        self.assertEqual(self.page.locator('.report-node[data-level="0"]').count(), 2)
        self.assertEqual(self.page.locator('.report-leaf').count(), 0)
        self.page.get_by_role('button', name='Развернуть всё', exact=True).click()
        self.assertEqual(self.page.locator('.report-leaf').count(), 2)
        self.page.get_by_role('button', name='Свернуть всё', exact=True).click()
        self.assertEqual(self.page.locator('#period-report details[open]').count(), 0)
        self.choose_filter('select[name=subdivision]', 'A')
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Применить', exact=True).click()
        self.assertEqual(self.page.locator('.report-node[data-level="0"]').count(), 1)
        expect(self.page.locator('.report-total')).to_contain_text('230')
        with self.page.expect_download() as download:
            self.page.get_by_role('link', name='Скачать CSV', exact=True).click()
        import csv
        with open(download.value.path(), encoding='utf-8-sig') as stream:
            rows = list(csv.reader(stream, delimiter=';'))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], 'A')
        self.assertNotIn('Дата', rows[0])
        with self.page.expect_navigation():
            self.page.get_by_role('link', name='Сбросить', exact=True).click()
        self.assertEqual(self.page.locator('.report-node[data-level="0"]').count(), 2)
        self.assertEqual(self.page.locator('.report-leaf').count(), 0)
        self.page.get_by_role('button', name='Развернуть всё', exact=True).click()
        self.assertEqual(self.page.locator('.report-leaf').count(), 2)
        self.assertEqual(self.errors, [])

    def test_comparison_preview_visibility_and_confirmed_replacement(self):
        from playwright.sync_api import expect
        from .test_sales_plans import SalesPlanTests
        from .services.approval import review_approval, approve_review
        from .services.versions import create_version
        def setup():
            fixture = SalesPlanTests(); fixture.setUp(); fixture.load()
            first = fixture.version
            approve_review(review_approval(first)['token'])
            future = create_version(date(2026, 7, 1), source_version=first.pk)
            future.coefficients.update(motivation_coeff=.2)
            approve_review(review_approval(future)['token'])
            current = create_version(date(2026, 7, 1), source_version=future.pk, allow_overwrite=True)
            row = current.coefficients.get(goods=fixture.good, type_coeff__type_coeff_name='Политики', segment__segment_name='K0')
            return fixture.scenario.pk, fixture.subs[0].pk, first.pk, future.pk, current.pk, row.pk
        plan, sub, first, future, current, row = self.db(setup)
        url = self.live_server_url + reverse('global_coeff_admin_motivation') + f'?version={current}'
        self.page.goto(url)
        self.assertTrue(self.page.locator('#comparison-show-approved').is_checked())
        self.assertFalse(self.page.locator('#comparison-show-baseline').is_checked())
        self.page.locator('#comparison-plan').select_option(str(plan))
        self.page.locator('#comparison-baseline').select_option(str(first))
        expect(self.page.locator('[data-amount="approved.total"]')).to_have_text('600')
        self.assertFalse(self.page.locator('[data-amount="baseline.total"]').is_visible())
        self.page.locator('#comparison-show-baseline').check()
        expect(self.page.locator('[data-amount="baseline.total"]')).to_have_text('225')
        self.page.locator(f'input[name="global_coeff={row}"]').fill('0.5')
        self.page.locator(f'input[name="global_coeff={row}"]').press('Tab')
        expect(self.page.locator('[data-amount="current.total"]')).to_have_text('645')
        self.assertEqual(self.db(lambda: GlobalCoeff.objects.get(pk=row).motivation_coeff), .2)
        self.page.locator('#approval-review').click()
        expect(self.page.locator('#approval-dialog')).to_be_visible()
        self.page.locator('#approval-reason').fill('Пилот: проверка администратором')
        self.page.locator('#approval-commit').click()
        expect(self.page.locator('#approval-error')).to_contain_text('Подтвердите замену')
        self.page.locator('#approval-confirm').check()
        with self.page.expect_navigation(): self.page.locator('#approval-commit').click()
        self.assertEqual(self.db(lambda: GlobalCoeffVersion.objects.get(pk=future).status), 'superseded')
        self.assertEqual(self.db(lambda: GlobalCoeff.objects.get(pk=row).motivation_coeff), .5)
        self.assertTrue(self.page.locator(f'input[name="global_coeff={row}"]').is_disabled())
        expect(self.page.locator('[data-amount="approved.total"]')).to_have_text('645')
        self.page.goto(self.live_server_url + reverse('coeff_subdivisions_admin_motivation', args=[sub]) + f'?version={current}')
        self.assertTrue(self.page.locator('#comparison-show-baseline').is_checked())
        self.assertEqual(self.page.locator('#comparison-plan').input_value(), str(plan))
        expect(self.page.locator('[data-amount="current.total"]')).to_have_text('230')
        self.page.locator('#comparison-plan').select_option('')
        expect(self.page.locator('[data-amount="current.total"]')).to_have_text('—')
        self.assertEqual(self.errors, [])

    def test_editor_filters_dirty_state_and_guard_preserve_hidden_edits(self):
        from playwright.sync_api import expect
        def setup_second():
            original = GlobalCoeff.objects.get(pk=self.coeff.pk)
            good = Goods.objects.create(goods_key='OtherKey', planning_group_sales='Other', group='Other ERP', brand='Other brand')
            return GlobalCoeff.objects.create(version=self.first, goods=good, type_coeff=original.type_coeff, segment=original.segment,
                variation_calculate=original.variation_calculate, motivation_coeff=.2, manager_coeff=1).pk
        second = self.db(setup_second)
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation'))
        self.assertFalse(self.page.locator('#create-version-details').evaluate('(element) => element.open'))
        self.assertFalse(self.page.locator('#add-product-details').evaluate('(element) => element.open'))
        expect(self.page.locator('#editor-save')).to_be_disabled()
        first_input = self.page.locator(f'input[name="global_coeff={self.coeff.pk}"]')
        first_input.focus(); first_input.press('Tab')
        expect(self.page.locator('#editor-dirty-count')).to_have_text('Изменено: 0')
        first_input.fill('0.25'); first_input.press('Tab')
        expect(first_input).to_have_class(__import__('re').compile('editor-changed'))
        expect(self.page.locator('#editor-dirty-count')).to_have_text('Изменено: 1')
        self.page.locator('#coefficient-only-changed').check()
        expect(self.page.locator('#coefficient-visible-count')).to_have_text('Показано: 1 из 2')
        dialogs = []
        def cancel_dialog(dialog): dialogs.append(dialog.message); dialog.dismiss()
        self.page.on('dialog', cancel_dialog)
        self.page.locator('.header-brand').click()
        self.assertEqual(len(dialogs), 1)
        self.assertIn('несохранённые', dialogs[0])
        expect(first_input).to_have_value('25%')
        self.page.remove_listener('dialog', cancel_dialog)
        first_input.fill('12.3456%'); first_input.press('Tab')
        expect(self.page.locator('#editor-dirty-count')).to_have_text('Изменено: 0')
        expect(self.page.locator('#coefficient-visible-count')).to_have_text('Показано: 0 из 2')
        self.page.locator('#coefficient-reset').click()
        first_input.fill('0.25'); first_input.press('Tab')
        second_input = self.page.locator(f'input[name="global_coeff={second}"]')
        second_input.fill('0.35'); second_input.press('Tab')
        self.choose_filter('#coefficient-pg', 'Other')
        expect(first_input).not_to_be_visible()
        expect(self.page.locator('#coefficient-visible-count')).to_have_text('Показано: 1 из 2')
        with self.page.expect_navigation(): self.page.locator('#editor-save').click()
        self.assertEqual(self.db(lambda: GlobalCoeff.objects.get(pk=self.coeff.pk).motivation_coeff), .25)
        self.assertEqual(self.db(lambda: GlobalCoeff.objects.get(pk=second).motivation_coeff), .35)
        expect(self.page.locator('#editor-dirty-count')).to_have_text('Изменено: 0')
        self.page.locator('#coefficient-search').fill('OtherKey')
        expect(self.page.locator('#coefficient-visible-count')).to_have_text('Показано: 1 из 2')
        self.page.locator('#coefficient-search').fill('missing')
        expect(self.page.locator('#coefficient-filter-empty')).to_be_visible()
        self.page.locator('#coefficient-reset').click()
        self.choose_filter('#coefficient-group', 'Other ERP')
        expect(self.page.locator('#coefficient-visible-count')).to_have_text('Показано: 1 из 2')
        self.page.locator('#coefficient-reset').click()
        self.choose_filter('#coefficient-brand', 'Other brand')
        expect(self.page.locator('#coefficient-visible-count')).to_have_text('Показано: 1 из 2')
        self.assertEqual(self.errors, [])

    def test_subdivision_navigation_and_copy_confirmation_and_manager_guard(self):
        from playwright.sync_api import expect
        def setup_sub():
            channel = Chanel.objects.create(chanel_name='Dealer')
            subdivisions = [Subdivision.objects.create(subdivision_key=name, subdivision_name=name, subdivision_global=name, chanel=channel) for name in ['A', 'B']]
            original = GlobalCoeff.objects.get(pk=self.coeff.pk)
            row = SubdivisionCoeff.objects.create(version=self.first, subdivision=subdivisions[0], goods=original.goods,
                type_coeff=original.type_coeff, segment=original.segment, variation_calculate=original.variation_calculate, motivation_coeff=.15, manager_coeff=1)
            manager = SubdivisionManagerCoeff.objects.create(version=self.first, subdivision=subdivisions[0], kind=KindManagerCoeff.objects.create(name='Year'), coeff=.8)
            return subdivisions[0].pk, row.pk, manager.pk
        sub, row, manager = self.db(setup_sub)
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation'))
        self.page.locator('#subdivision-nav summary').focus()
        self.page.locator('#subdivision-nav summary').press('Enter')
        self.page.locator('#subdivision-search').fill('A')
        expect(self.page.locator('#subdivision-links a:visible')).to_have_count(1)
        with self.page.expect_navigation(): self.page.locator('#subdivision-links a:visible').click()
        self.assertIn(f'/{sub}', self.page.url)
        self.assertFalse(self.page.locator('#manager-details').evaluate('(element) => element.open'))
        self.page.locator('#manager-details summary').click()
        self.page.locator('input[name^="kind_coeff="]').fill('0.9')
        self.page.locator('input[name^="kind_coeff="]').press('Tab')
        expect(self.page.locator('#editor-save')).to_be_disabled()
        self.page.locator('#approval-review').click()
        expect(self.page.locator('#editor-error')).to_contain_text('Сначала сохраните')
        expect(self.page.locator('#approval-dialog')).not_to_be_visible()
        with self.page.expect_navigation(): self.page.get_by_role('button', name='Сохранить руководителя', exact=True).click()
        self.assertEqual(self.db(lambda: SubdivisionManagerCoeff.objects.get(pk=manager).coeff), .9)
        local_input = self.page.locator(f'input[name="sub_coeff={row}"]')
        local_input.fill('.25'); local_input.press('Tab')
        self.page.get_by_role('button',name='Импорт и справочники',exact=True).click()
        dialogs = []
        def cancel(dialog): dialogs.append(dialog.message); dialog.dismiss()
        self.page.on('dialog', cancel)
        self.page.get_by_role('button', name='Копировать из глобальных', exact=True).click()
        self.page.get_by_role('button', name='Скопировать выбранную группу', exact=True).click()
        self.assertEqual(len(dialogs), 1)
        self.assertIn('Заменить коэффициенты группы', dialogs[0])
        self.assertEqual(self.db(lambda: SubdivisionCoeff.objects.get(pk=row).motivation_coeff), .15)
        self.page.remove_listener('dialog', cancel)
        self.page.once('dialog', lambda dialog: dialog.accept())
        with self.page.expect_navigation(): self.page.get_by_role('button', name='Скопировать выбранную группу', exact=True).click()
        self.assertEqual(self.db(lambda: SubdivisionCoeff.objects.get(version=self.first, subdivision_id=sub).motivation_coeff), .123456)
        expect(self.page.locator('#editor-dirty-count')).to_have_text('Изменено: 0')
        self.assertEqual(self.errors, [])

    def test_director_forms_copy_global_and_save_individual_percentages(self):
        from playwright.sync_api import expect
        from decimal import Decimal
        from .models import GlobalDirectorCoeff, SubdivisionDirectorCoeff, DirectorEmployee
        def setup():
            channel = Chanel.objects.create(chanel_name='Director channel')
            sub = Subdivision.objects.create(subdivision_key='Director branch', subdivision_name='Director branch', subdivision_global='Director branch', chanel=channel)
            employee = DirectorEmployee.objects.create(full_name='Директор один', employee_code='DIR-1')
            other = DirectorEmployee.objects.create(full_name='Директор два', employee_code='DIR-2')
            employee.subdivisions.add(sub)
            other.subdivisions.add(sub)
            GlobalDirectorCoeff.objects.create(version=self.first, goods=self.coeff.goods, k0=Decimal('.123456'), k1=Decimal('.02'))
            return sub.pk, employee.pk, other.pk
        sub, employee, other = self.db(setup)
        self.page.goto(self.live_server_url + reverse('motivation_directors_global'))
        expect(self.page.get_by_label('Директор K0', exact=True)).to_have_value('12.3456%')
        self.page.goto(self.live_server_url + reverse('motivation_directors_local') + f'?subdivision={sub}&employee={employee}')
        self.page.get_by_role('button',name='Импорт и справочники',exact=True).click()
        self.page.get_by_role('button', name='Копировать из глобальных', exact=True).click()
        self.page.locator('#copy-coefficient-group').select_option('directors')
        self.page.once('dialog', lambda dialog: dialog.accept())
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Скопировать выбранную группу', exact=True).click()
        expect(self.page.get_by_label('Директор K0', exact=True)).to_have_value('12.3456%')
        self.page.get_by_label('Директор K0', exact=True).fill('25%')
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Сохранить изменения', exact=True).click()
        expect(self.page.locator('#version-panel')).to_contain_text('Сохранение проверено')
        expect(self.page.get_by_label('Директор K0', exact=True)).to_have_value('25%')
        self.assertEqual(self.db(lambda: SubdivisionDirectorCoeff.objects.get(employee_id=employee).k0), Decimal('.25'))
        self.assertFalse(self.db(lambda: SubdivisionDirectorCoeff.objects.filter(employee_id=other).exists()))
        self.assertEqual(self.db(lambda: GlobalDirectorCoeff.objects.get().k0), Decimal('.123456'))
        self.assertEqual(self.errors, [])

    def test_shared_director_columns_visibility_and_copy_group_choice(self):
        from playwright.sync_api import expect
        from decimal import Decimal
        from .models import GlobalDirectorCoeff, SubdivisionDirectorCoeff, DirectorEmployee
        def setup():
            channel = Chanel.objects.create(chanel_name='Group copy channel')
            sub = Subdivision.objects.create(subdivision_key='Group copy branch', subdivision_name='Group copy branch', subdivision_global='Group copy branch', chanel=channel)
            employee = DirectorEmployee.objects.create(full_name='Директор группы', employee_code='GROUP-DIR')
            employee.subdivisions.add(sub)
            GlobalDirectorCoeff.objects.create(version=self.first, goods=self.coeff.goods, k0=Decimal('.03'))
            return sub.pk, employee.pk
        sub, employee = self.db(setup)
        url = self.live_server_url + reverse('global_coeff_admin_motivation')
        self.page.goto(url)
        expect(self.page.get_by_label('Директор K0', exact=True)).to_have_value('3%')
        self.page.locator('[data-coefficient-group-toggle="directors"]').uncheck()
        expect(self.page.get_by_label('Директор K0', exact=True)).not_to_be_visible()
        self.page.reload()
        expect(self.page.locator('[data-coefficient-group-toggle="directors"]')).not_to_be_checked()
        self.page.locator('[data-coefficient-group-toggle="directors"]').check()
        self.page.get_by_label('Директор K0', exact=True).fill('7%')
        self.page.get_by_label('Директор K0', exact=True).dispatch_event('change')
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Сохранить изменения', exact=True).click()
        self.assertEqual(self.db(lambda: GlobalDirectorCoeff.objects.get().k0), Decimal('.07'))
        self.page.goto(self.live_server_url + reverse('coeff_subdivisions_admin_motivation', args=[sub]))
        self.page.get_by_role('button',name='Импорт и справочники',exact=True).click()
        self.page.get_by_role('button', name='Копировать из глобальных', exact=True).click()
        expect(self.page.locator('#copy-global-dialog')).to_be_visible()
        self.page.locator('#copy-coefficient-group').select_option('directors')
        self.page.locator('#copy-director-select').select_option(str(employee))
        self.page.once('dialog', lambda dialog: dialog.accept())
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Скопировать выбранную группу', exact=True).click()
        self.assertEqual(self.db(lambda: SubdivisionDirectorCoeff.objects.get(employee_id=employee).k0), Decimal('.07'))
        self.assertIn(reverse('coeff_subdivisions_admin_motivation', args=[sub]), self.page.url)
        self.assertEqual(self.errors, [])

    def test_local_employee_change_warns_and_copy_to_another_director(self):
        from playwright.sync_api import expect
        from decimal import Decimal
        from .models import DirectorEmployee, SubdivisionDirectorCoeff
        def setup():
            channel = Chanel.objects.create(chanel_name='Switch director channel')
            sub = Subdivision.objects.create(subdivision_key='Switch branch', subdivision_name='Switch branch', subdivision_global='Switch branch', chanel=channel)
            first = DirectorEmployee.objects.create(full_name='Первый', employee_code='SW-1')
            second = DirectorEmployee.objects.create(full_name='Второй', employee_code='SW-2')
            for employee in (first, second):
                employee.subdivisions.add(sub)
                SubdivisionDirectorCoeff.objects.create(version=self.first, subdivision=sub, employee=employee,
                    goods=self.coeff.goods, k0=Decimal('.1') if employee.pk == first.pk else Decimal('.2'))
            return sub.pk, first.pk, second.pk
        sub, first, second = self.db(setup)
        url = self.live_server_url + reverse('coeff_subdivisions_admin_motivation', args=[sub]) + f'?employee={first}'
        self.page.goto(url)
        expect(self.page.get_by_label('Директор K0', exact=True)).to_have_value('10%')
        self.page.get_by_label('Директор K0', exact=True).fill('15%')
        self.page.once('dialog', lambda dialog: dialog.dismiss())
        self.choose_filter('#director-employee-choice', str(second))
        expect(self.page.locator('#director-employee-choice')).to_have_value(str(first))
        expect(self.page.get_by_label('Директор K0', exact=True)).to_have_value('15%')
        self.page.once('dialog', lambda dialog: dialog.accept())
        with self.page.expect_navigation():
            self.choose_filter('#director-employee-choice', str(second))
        expect(self.page.get_by_label('Директор K0', exact=True)).to_have_value('20%')
        self.page.get_by_role('button',name='Импорт и справочники',exact=True).click()
        self.page.get_by_role('button', name='Копировать другому директору', exact=True).click()
        self.page.locator('[name=copy_to_employee]').select_option(str(first))
        self.page.once('dialog', lambda dialog: dialog.accept())
        with self.page.expect_navigation():
            self.page.get_by_role('button', name='Скопировать коэффициенты', exact=True).click()
        self.assertEqual(self.db(lambda: SubdivisionDirectorCoeff.objects.get(employee_id=first).k0), Decimal('.2'))
        expect(self.page.locator('#director-employee-choice')).to_have_value(str(first))
        self.assertEqual(self.errors, [])

    def test_pilot_dashboard_expansion_keeps_headers_and_rows_aligned(self):
        from playwright.sync_api import expect
        from .services.reports import manager_report
        from .services.report_periods import PILOT_QUARTER
        from .test_reports import NativeMotivationReportTests
        rows = [dict(NativeMotivationReportTests.manager_rows()[0], КварталМотивации=p)
                for p in ('1 Квартал', PILOT_QUARTER, 'Год')]
        report = manager_report(rows, pilot=True)
        with patch('admin_motivation.report_views.choices', return_value=[]), patch('admin_motivation.report_views.load_report', return_value=report):
            self.page.goto(self.live_server_url + reverse('motivation_report', args=['manager']) + '?run=1&quarter=4+Квартал')
            dashboard = self.page.locator('#manager-dashboard')
            expect(dashboard).to_have_attribute('data-year-period-count', '2')
            for reveal in ('year-possible', 'year-available', 'year-period', 'year-motivation'):
                dashboard.locator('.report-column-header [data-dashboard-reveal="' + reveal + '"]').click()
            self.assertEqual(dashboard.evaluate('(el) => el.style.getPropertyValue("--dashboard-year")'), '14')
            counts = dashboard.locator('.report-column-header, .report-total-row, .division-percent').evaluate_all(
                '(rows) => rows.map(row => Array.from(row.children).filter(cell => !cell.hidden).length)')
            self.assertEqual(len(set(counts)), 1)
            expect(dashboard.locator('.report-column-header [data-dashboard-column="year-possible"]:visible')).to_have_count(2)
            self.page.screenshot(path='/tmp/pilot-dashboard.png', full_page=True)
        self.assertEqual(self.errors, [])

    def test_native_reports_render_manager_and_lazy_subdivision_tree(self):
        from unittest.mock import patch
        from playwright.sync_api import expect
        from .test_reports import NativeMotivationReportTests
        from .services.reports import manager_report, subdivision_report
        manager=manager_report(NativeMotivationReportTests.manager_rows())
        division=subdivision_report(NativeMotivationReportTests.division_rows())
        with patch('admin_motivation.report_views.choices',return_value=[]), patch('admin_motivation.report_views.load_report',return_value=manager):
            self.page.goto(self.live_server_url + reverse('motivation_report',args=['manager']) + '?run=1&type=Политики&subdivision=Филиал&author=Менеджер&quarter=1+Квартал')
            expect(self.page.locator('.report-title')).to_have_text('Мотивация. Дашборд · USD')
            expect(self.page.locator('.manager-section').nth(1)).not_to_have_attribute('open','')
            self.page.locator('.manager-section').nth(1).locator('summary').click()
            expect(self.page.get_by_text('4. VENT',exact=True)).to_be_visible()
            expect(self.page.locator('.manager-grid.report-total-row')).to_contain_text('360')
            self.page.screenshot(path='/tmp/native-manager-report.png',full_page=True)
        with patch('admin_motivation.report_views.choices',return_value=[]), patch('admin_motivation.report_views.load_report',return_value=division):
            self.page.goto(self.live_server_url + reverse('motivation_report',args=['subdivisions']) + '?run=1')
            self.assertEqual(self.page.locator('.native-division-node').count(),1)
            self.assertEqual(self.page.locator('.native-division-leaf').count(),0)
            self.page.get_by_role('button',name='Развернуть всё',exact=True).click()
            expect(self.page.locator('.native-division-leaf')).to_be_visible()
            expect(self.page.locator('.native-division-leaf')).to_contain_text('Brand')
            root = self.page.locator('#native-division-report')
            expect(root.locator('.report-column-header [data-division-bonus]')).to_have_count(1)
            expect(root.locator('.report-total-row [data-division-bonus]')).to_have_count(1)
            expect(root.locator('.native-division-leaf [data-division-bonus]')).to_have_count(1)
            expect(root.locator('[data-division-motivation]:visible')).to_have_count(0)
            expect(root.locator('[data-division-segments]:visible')).to_have_count(0)
            expect(root.get_by_text('План периода', exact=True).first).to_be_visible()
            expect(root.get_by_text('% вып. периода', exact=True).first).to_be_visible()
            root.locator('.report-column-header [data-division-reveal="motivation"]').first.click()
            expect(self.page.locator('.native-division-leaf [data-division-motivation]').first).to_be_visible()
            root.locator('.report-column-header [data-division-reveal="segments"]').first.click()
            expect(self.page.locator('.native-division-leaf [data-division-segments]').first).to_be_visible()
            self.assertEqual(root.evaluate('(el) => el.style.getPropertyValue("--report-columns")'), '15')
            widths = root.locator('.report-period-header > strong').evaluate_all('(els) => els.map(el => el.getBoundingClientRect().width)')
            self.assertEqual(widths[-1] - widths[0], 112)
            counts = root.locator('.report-column-header, .report-total-row, .native-division-leaf').evaluate_all('(els) => els.map(el => Array.from(el.children).filter(child => !child.hidden).length)')
            self.assertEqual(len(set(counts)), 1)
            root.locator('.native-division-node > summary [data-division-reveal="segments"]').first.click()
            expect(root.locator('[data-division-segments]:visible')).to_have_count(0)
            expect(root.locator('.native-division-node').first).to_have_attribute('open', '')
            self.page.get_by_role('button',name='Свернуть всё',exact=True).click()
            self.assertEqual(self.page.locator('.native-division-node[open]').count(),0)
        self.assertEqual(self.errors,[])

    def test_bulk_multiply_common_segments_validation_undo_and_save(self):
        from playwright.sync_api import expect
        from .models import GlobalDirectorCoeff
        def fixture():
            original=GlobalCoeff.objects.get(pk=self.coeff.pk)
            original.segment.segment_name='K0';original.segment.save()
            original.motivation_coeff=.01;original.save()
            rows=[original]
            for index in range(1,5):
                segment=SegmentCoeff.objects.create(segment_name=f'K{index}')
                rows.append(GlobalCoeff.objects.create(version=self.first,goods=original.goods,type_coeff=original.type_coeff,segment=segment,motivation_coeff=.01*(index+1),manager_coeff=1,variation_calculate=original.variation_calculate))
            director=GlobalDirectorCoeff.objects.create(version=self.first,goods=original.goods,**{f'k{i}':'.01' for i in range(5)})
            return original.goods_id,[row.pk for row in rows],director.pk
        good,ids,director=self.db(fixture)
        self.page.goto(self.live_server_url+reverse('global_coeff_admin_motivation'))
        self.page.get_by_role('button',name='Изменить ставки',exact=True).click()
        self.page.locator(f'.bulk-good[value="{good}"]').check()
        self.page.locator('#multiply-common').fill('1,1')
        self.page.locator('#multiply-preview').click()
        expect(self.page.locator('#multiply-review')).to_be_visible()
        self.assertEqual(self.db(lambda:GlobalCoeff.objects.get(pk=ids[0]).motivation_coeff),.01)
        self.page.locator('#multiply-apply').click()
        first=self.page.locator(f'[name="global_coeff={ids[0]}"]')
        expect(first).to_have_value('1.1%')
        self.page.locator('#multiply-undo').click();expect(first).to_have_value('1%')
        self.page.locator('#multiply-mode').select_option('segments')
        for i,value in enumerate(['2','1','0.5','1','0']):self.page.locator(f'[data-multiply-segment="{i}"]').fill(value)
        self.page.locator('#multiply-preview').click();self.page.locator('#multiply-apply').click()
        expect(first).to_have_value('2%')
        expect(self.page.locator(f'[name="global_coeff={ids[2]}"]')).to_have_value('1.5%')
        expect(self.page.locator(f'[name="global_coeff={ids[4]}"]')).to_have_value('0%')
        self.page.locator('#multiply-mode').select_option('common');self.page.locator('#multiply-common').fill('1000')
        self.page.locator('#multiply-preview').click();expect(self.page.locator('#multiply-error')).to_contain_text('диапазон')
        expect(first).to_have_value('2%')
        self.page.locator('[data-multiply-kind="policies"]').uncheck();self.page.locator('[data-multiply-kind="directors"]').check()
        self.page.locator('#multiply-common').fill('2');self.page.locator('#multiply-preview').click();self.page.locator('#multiply-apply').click()
        expect(self.page.locator(f'[name="director_coeff={good}=4"]')).to_have_value('2%')
        with self.page.expect_navigation():self.page.locator('#editor-save').click()
        self.assertEqual(self.db(lambda:GlobalCoeff.objects.get(pk=ids[0]).motivation_coeff),.02)
        self.assertEqual(self.db(lambda:GlobalCoeff.objects.get(pk=ids[4]).motivation_coeff),0)
        self.assertEqual(str(self.db(lambda:GlobalDirectorCoeff.objects.get(pk=director).k4)),'0.020000')
        self.assertEqual(self.errors,[])

    def test_local_bulk_multiply_changes_only_selected_director(self):
        from playwright.sync_api import expect
        from .models import DirectorEmployee, SubdivisionDirectorCoeff
        def fixture():
            original=GlobalCoeff.objects.get(pk=self.coeff.pk)
            original.segment.segment_name='K0';original.segment.save()
            channel=Chanel.objects.create(chanel_name='Multiply local')
            sub=Subdivision.objects.create(subdivision_key='Multiply branch',subdivision_name='Multiply branch',subdivision_global='Multiply branch',chanel=channel)
            SubdivisionCoeff.objects.create(version=self.first,goods=original.goods,type_coeff=original.type_coeff,segment=original.segment,motivation_coeff=.02,manager_coeff=1,variation_calculate=original.variation_calculate,subdivision=sub)
            employees=[DirectorEmployee.objects.create(full_name=name) for name in ['Multiply first','Multiply second']]
            for employee in employees:
                employee.subdivisions.add(sub)
                SubdivisionDirectorCoeff.objects.create(version=self.first,goods=original.goods,subdivision=sub,employee=employee,**{f'k{i}':'.01' for i in range(5)})
            return sub.pk,employees[0].pk,employees[1].pk,original.goods_id
        branch,first,second,good=self.db(fixture)
        self.page.goto(self.live_server_url+reverse('coeff_subdivisions_admin_motivation',args=[branch])+f'?employee={first}')
        self.page.get_by_role('button',name='Изменить ставки',exact=True).click()
        self.page.locator(f'.bulk-good[value="{good}"]').check()
        self.page.locator('[data-multiply-kind="policies"]').uncheck();self.page.locator('[data-multiply-kind="directors"]').check()
        self.page.locator('#multiply-common').fill('2');self.page.locator('#multiply-preview').click();self.page.locator('#multiply-apply').click()
        expect(self.page.locator(f'[name="director_coeff={good}=0"]')).to_have_value('2%')
        with self.page.expect_navigation():self.page.locator('#editor-save').click()
        self.assertEqual(str(self.db(lambda:SubdivisionDirectorCoeff.objects.get(employee_id=first).k0)),'0.020000')
        self.assertEqual(str(self.db(lambda:SubdivisionDirectorCoeff.objects.get(employee_id=second).k0)),'0.010000')
        self.assertEqual(self.db(lambda:SubdivisionCoeff.objects.get(subdivision_id=branch).motivation_coeff),.02)
        self.assertEqual(self.errors,[])

    def test_reapprove_after_reopen_preserves_prior_approval(self):
        from playwright.sync_api import expect
        from .services.approval import approve_review,review_approval
        from .models import MotivationApproval
        first_log=self.db(lambda:approve_review(review_approval(GlobalCoeffVersion.objects.get(pk=self.first.pk))['token']).approval.pk)
        self.page.goto(self.live_server_url+reverse('global_coeff_admin_motivation'))
        self.page.locator('.version-tools button[aria-controls="version-reopen"]').click()
        self.page.locator('#version-reopen [name=reason]').fill('Browser reapproval')
        self.page.locator('#version-reopen [name=confirm]').check()
        with self.page.expect_navigation():self.page.locator('#version-reopen button[type=submit]').click()
        self.page.locator('#approval-review').click()
        expect(self.page.locator('#approval-commit')).to_be_enabled()
        self.page.locator('#approval-reason').fill('Пилот: проверка администратором')
        with self.page.expect_navigation():self.page.locator('#approval-commit').click()
        expect(self.page.locator('#version-approval-actions')).to_contain_text('Утверждена')
        self.assertEqual(self.db(lambda:MotivationApproval.objects.filter(version_id=self.first.pk).count()),2)
        self.assertTrue(self.db(lambda:MotivationApproval.objects.filter(pk=first_log).exists()))
        self.assertEqual(self.errors,[])

    def test_director_import_template_button_downloads_xlsx(self):
        from playwright.sync_api import expect
        from openpyxl import load_workbook
        self.page.goto(self.live_server_url+reverse('motivation_director_coeff_import')+f'?version={self.first.pk}')
        button=self.page.get_by_role('link',name='Скачать шаблон Excel',exact=True)
        expect(button).to_be_visible()
        with self.page.expect_download() as event:button.click()
        download=event.value
        self.assertEqual(download.suggested_filename,f'director-coefficients-v{self.first.pk}.xlsx')
        download.save_as('/tmp/director-import-template-browser.xlsx')
        book=load_workbook('/tmp/director-import-template-browser.xlsx')
        self.assertIn('Коэффициенты',book.sheetnames)
        self.assertIn('Инструкция',book.sheetnames)
        self.assertEqual(book['Коэффициенты']['E5'].value,'K0')
        book.close()
        self.assertEqual(self.errors,[])

    def test_subdivision_premium_filter_year_header_and_excel_downloads(self):
        from playwright.sync_api import expect
        from openpyxl import load_workbook
        from .services.reports import PREMIUM_PERIODS, subdivision_report
        from .test_subdivision_report_export import SubdivisionReportExportTests
        base = SubdivisionReportExportTests.rows()[0]
        rows = [dict(base, КварталМотивации=f'{quarter} Квартал') for quarter in range(1, 5)] + [dict(base, КварталМотивации='Год')]
        def options(kind, values, **kwargs):
            return [{'name': 'period', 'label': 'Период премии', 'multiple': True,
                'options': [{'value': code, 'label': label, 'selected': code in values.get('PremiumPeriods', [])} for code, label in PREMIUM_PERIODS]}]
        def report(kind, values, page, scope, **kwargs):
            selected = values.get('PremiumPeriods')
            subset = [row for row in rows if not selected or ('5' if row['КварталМотивации'] == 'Год' else row['КварталМотивации'][0]) in selected]
            return subdivision_report(subset, scope)
        url = self.live_server_url + reverse('motivation_report', args=['subdivisions'])
        with patch('admin_motivation.report_views.choices', side_effect=options), patch('admin_motivation.report_views.load_report', side_effect=report):
            self.page.goto(url)
            dropdown = self.page.locator('details[data-field=period]')
            dropdown.locator('summary').click()
            dropdown.get_by_label('4 квартал', exact=True).check()
            dropdown.get_by_role('button', name='Применить', exact=True).click()
            with self.page.expect_navigation():
                self.page.get_by_role('button', name='Сформировать отчёт', exact=True).click()
            expect(self.page.locator('h1.report-title')).to_contain_text('Год расчёта: 2026')
            expect(self.page.locator('#subdivision-summary-table thead')).to_contain_text('4 Квартал')
            expect(self.page.locator('#subdivision-summary-table thead')).not_to_contain_text('1 Квартал')
            self.assertIn('period=4', self.page.url)
            for label, count in [('Скачать весь отчёт Excel', 3), ('Скачать итоговую таблицу Excel', 1), ('Скачать детализацию Excel', 1)]:
                with self.page.expect_download() as event:
                    self.page.get_by_role('link', name=label, exact=True).click()
                download = event.value
                path = '/tmp/' + download.suggested_filename
                download.save_as(path)
                book = load_workbook(path)
                self.assertEqual(len(book.sheetnames), count)
                self.assertIn('4 квартал', book.active['A2'].value)
                book.close()
            self.page.screenshot(path='/tmp/subdivision-report-premium-period.png', full_page=True)
            self.page.wait_for_load_state('networkidle')
        self.assertEqual(self.errors, [])

    def test_version_fallback_setting_and_plan_coverage(self):
        from playwright.sync_api import expect
        from django.utils import timezone
        from .models import SalesPlanScenario, SalesPlanLine
        from .services.pilot_formula import GLOBAL_SUBDIVISION
        form_url = reverse('motivation_fallback_formula')
        self.page.goto(self.live_server_url + reverse('global_coeff_admin_motivation'))
        form = self.page.locator(f'form[action="{form_url}"]')
        form.locator('select[name=fallback_variation]').select_option(GLOBAL_SUBDIVISION)
        with self.page.expect_navigation():
            form.get_by_role('button', name='Сохранить формулу по умолчанию').click()
        self.assertEqual(self.db(lambda: GlobalCoeffVersion.objects.get(pk=self.first.pk).fallback_variation), GLOBAL_SUBDIVISION)
        def create_plan():
            channel = Chanel.objects.create(chanel_name='Coverage channel')
            branch = Subdivision.objects.create(subdivision_key='Coverage-A', subdivision_name='Coverage-A', subdivision_global='Coverage-A', chanel=channel)
            plan = SalesPlanScenario.objects.create(title='Coverage plan', source_version='coverage', loaded_at=timezone.now())
            SalesPlanLine.objects.create(scenario=plan, plan_date=date(2026, 1, 1), subdivision=branch.subdivision_key,
                planning_group_sales='Missing planning group', group='Missing group', brand='Missing brand', amount_usd=1000, segment_amounts={'O0': '1000'})
            return plan.pk
        plan = self.db(create_plan)
        self.page.goto(self.live_server_url + reverse('motivation_coefficient_coverage') + f'?version={self.first.pk}&plan={plan}&run=1')
        expect(self.page.get_by_role('heading', name='Покрытие коэффициентами')).to_be_visible()
        expect(self.page.locator('tbody')).to_contain_text('Нет точного коэффициента')
        expect(self.page.locator('tbody')).to_contain_text(GLOBAL_SUBDIVISION)
        self.page.screenshot(path='/tmp/coefficient-fallback-coverage.png', full_page=True)
        self.page.wait_for_load_state('networkidle')
        self.assertEqual(self.errors, [])

    def test_unallocated_author_hidden_toggle_preserves_totals(self):
        from playwright.sync_api import expect
        from .test_manager_detail import detail_rows
        from .services.manager_detail import manager_detail_report
        rows = detail_rows()
        rows += [dict(row, АвторСделки=' ПОЛИТИКИ ') for row in detail_rows()]
        data = manager_detail_report(rows)
        with patch('admin_motivation.report_views.choices', return_value=[]), patch('admin_motivation.report_views.load_report', return_value=data):
            self.page.goto(self.live_server_url + reverse('motivation_report', args=['manager-detail']) + '?run=1&type=Политики&details=1')
            hidden = self.page.locator('[data-manager-unallocated]')
            expect(hidden).to_have_count(2)
            for row in hidden.all():
                expect(row).to_be_hidden()
            before = self.page.locator('.manager-period-summary .report-total-row').first.inner_text()
            self.page.locator('#manager-show-unallocated').check()
            for row in hidden.all():
                expect(row).to_be_visible()
            self.assertEqual(self.page.locator('.manager-period-summary .report-total-row').first.inner_text(), before)
            self.page.locator('#manager-show-unallocated').uncheck()
            self.page.locator('#manager-detail-expand').click()
            for row in hidden.all():
                expect(row).to_be_hidden()
        self.assertEqual(self.errors, [])

    def test_first_quarter_inputs_equal_period_fact(self):
        from playwright.sync_api import expect
        from .test_manager_detail import detail_rows
        from .services.manager_detail import manager_detail_report
        from .services.pilot_formula import GLOBAL_SUBDIVISION
        rows = [dict(row, VariationCalculate=GLOBAL_SUBDIVISION) for row in detail_rows()]
        performance = [dict(row, СуммаОтгрузокUSDКвартал=6469139, СуммаПланUSD=235214543, СуммаОтгрузокUSD=235214543) for row in rows]
        data = manager_detail_report(rows, performance_rows=performance)
        with patch('admin_motivation.report_views.choices', return_value=[]), patch('admin_motivation.report_views.load_report', return_value=data):
            self.page.goto(self.live_server_url + reverse('motivation_report', args=['manager-detail']) + '?run=1&type=Политики&details=1')
            department = self.page.locator('.manager-division-execution')
            expect(department.locator('[data-manager-reveal=motivation]').first).to_have_text('100,00%')
            department.locator('[data-manager-reveal=motivation]').first.click()
            for index in (0, 1):
                value = department.locator('[data-manager-motivation]').nth(index)
                expect(value).to_be_visible()
                self.assertEqual(value.inner_text().replace('\xa0', ' '), '6 469 139')
        self.assertEqual(self.errors, [])

    def test_global_subdivision_detail_execution_and_excel(self):
        from playwright.sync_api import expect
        from openpyxl import load_workbook
        from .test_manager_detail import detail_rows
        from .services.manager_detail import manager_detail_report
        from .services.pilot_formula import GLOBAL_SUBDIVISION
        rows = [dict(row, VariationCalculate=GLOBAL_SUBDIVISION,
                     КварталМотивации='2 Квартал(1-2 кв)' if row['КварталМотивации'] == '1 Квартал' else row['КварталМотивации'])
                for row in detail_rows()]
        performance = [dict(row, СуммаПланUSD=1000, СуммаОтгрузокUSD=628.4) for row in rows]
        def result(*args, **kwargs):
            return manager_detail_report(rows, performance_rows=performance)
        with patch('admin_motivation.report_views.choices', return_value=[]), patch('admin_motivation.report_views.load_report', side_effect=result):
            self.page.goto(self.live_server_url + reverse('motivation_report', args=['manager-detail']) + '?run=1&type=Политики&details=1')
            self.page.locator('#manager-detail-expand').click()
            expect(self.page.locator('#manager-detail-nodes .motivation-execution').first).to_contain_text('62,84%')
            first = self.page.locator('#manager-detail-nodes summary.division-report-grid').first
            expect(first.locator('[data-manager-period-input]').first).to_have_text('—')
            expect(first.locator('[data-manager-reveal=period-input]').first).to_have_text('—')
            department = self.page.locator('.manager-division-execution')
            department.locator('[data-manager-reveal=period-input]').first.click()
            expect(department.locator('[data-manager-period-input]').first).to_be_visible()
            self.assertEqual(department.locator('[data-manager-period-input]').first.inner_text().replace('\xa0', ' '), '1 000')
            department.locator('[data-manager-reveal=motivation]').first.click()
            expect(department.locator('[data-manager-motivation]').first).to_be_visible()
            self.assertEqual(department.locator('[data-manager-motivation]').first.inner_text().replace('\xa0', ' '), '1 000')
            with self.page.expect_download() as event:
                self.page.get_by_role('link', name='Скачать расширенную таблицу Excel').click()
            event.value.save_as('/tmp/global-subdivision-execution.xlsx')
            book = load_workbook('/tmp/global-subdivision-execution.xlsx')
            sheet = book.active
            # First period: plan absent; own fact and branch motivation inputs retained.
            self.assertEqual(sheet.cell(3, 12).value, 1000)
            self.assertEqual(sheet.cell(3, 13).value, 800)
            self.assertEqual(sheet.cell(3, 16).value, 1000)
            self.assertAlmostEqual(sheet.cell(3, 17).value, 628.4)
            self.assertIsNone(sheet.cell(4, 12).value)
            self.assertEqual(sheet.cell(4, 16).value, 1000)
            self.assertAlmostEqual(sheet.cell(4, 18).value, .6284)
            book.close()
            self.page.screenshot(path='/tmp/global-subdivision-detail-execution.png', full_page=True)
            self.page.wait_for_load_state('networkidle')
        self.assertEqual(self.errors, [])


    def test_rate_editor_set_relative_points_diff_reset_and_saved_journal(self):
        from playwright.sync_api import expect
        from .models import CoefficientEditLog, GlobalDirectorCoeff
        def fixture():
            row=GlobalCoeff.objects.get(pk=self.coeff.pk)
            row.segment.segment_name='K0';row.segment.save();row.motivation_coeff=.002;row.save()
            segment=SegmentCoeff.objects.create(segment_name='K1')
            other=GlobalCoeff.objects.create(version=self.first,goods=row.goods,type_coeff=row.type_coeff,segment=segment,motivation_coeff=.006,manager_coeff=1,variation_calculate=row.variation_calculate)
            return row.goods_id,other.pk
        good,other=self.db(fixture)
        self.page.goto(self.live_server_url+reverse('global_coeff_admin_motivation'))
        first=self.page.locator(f'[name="global_coeff={self.coeff.pk}"]')
        self.page.get_by_role('button',name='Изменить ставки',exact=True).click()
        self.page.locator(f'.bulk-good[value="{good}"]').check()
        for index in range(1,5): self.page.locator(f'[data-multiply-step="{index}"]').uncheck()
        def apply(operation,value):
            self.page.locator('#multiply-operation').select_option(operation)
            self.page.locator('#multiply-common').fill(value)
            self.page.locator('#multiply-preview').click()
            expect(self.page.locator('#multiply-review')).to_be_visible()
            self.page.locator('#multiply-apply').click()
        apply('set','0,40%');expect(first).to_have_value('0.4%')
        expect(self.page.locator(f'[name="global_coeff={other}"]')).to_have_value('0.6%')
        apply('relative','10');expect(first).to_have_value('0.44%')
        self.page.locator('#editor-undo').click();expect(first).to_have_value('0.4%')
        apply('points','0,10');expect(first).to_have_value('0.5%')
        self.assertEqual(self.db(lambda:GlobalCoeff.objects.get(pk=self.coeff.pk).motivation_coeff),.002)
        self.page.locator('#editor-review-changes').click()
        expect(self.page.locator('#editor-changes-rows tr')).to_have_count(1)
        expect(self.page.locator('#editor-changes-rows')).to_contain_text('0,2%')
        expect(self.page.locator('#editor-changes-rows')).to_contain_text('0,5%')
        self.page.get_by_role('button',name='Вернуть исходное',exact=True).click()
        expect(self.page.locator('#editor-changes-rows tr')).to_have_count(0)
        self.page.locator('#editor-changes-close').click();expect(first).to_have_value('0.2%')
        self.page.locator('#editor-undo').click();expect(first).to_have_value('0.5%')
        self.page.locator('#editor-reset-changes').click();self.page.locator('#editor-reset-confirm').click()
        expect(first).to_have_value('0.2%')
        expect(self.page.locator('#editor-save')).to_be_disabled()
        self.page.locator('#editor-undo').click();expect(first).to_have_value('0.5%')
        with self.page.expect_navigation():self.page.locator('#editor-save').click()
        self.assertEqual(self.db(lambda:GlobalCoeff.objects.get(pk=self.coeff.pk).motivation_coeff),.005)
        self.assertEqual(self.db(lambda:CoefficientEditLog.objects.count()),1)
        self.assertFalse(self.db(lambda:GlobalDirectorCoeff.objects.exists()))
        expect(self.page.locator('#editor-dirty-count')).to_have_text('Изменено: 0')
        self.page.screenshot(path='/tmp/coefficient-rate-editor.png',full_page=True)
        self.assertEqual(self.errors,[])


    def test_rate_editor_global_exact_copy_and_hidden_filter_selection(self):
        from playwright.sync_api import expect
        def fixture():
            original=GlobalCoeff.objects.get(pk=self.coeff.pk);original.segment.segment_name='K0';original.segment.save();original.motivation_coeff=.002;original.save()
            channel=Chanel.objects.create(chanel_name='Rates')
            sub=Subdivision.objects.create(subdivision_key='A',subdivision_name='A',subdivision_global='A',chanel=channel)
            first=SubdivisionCoeff.objects.create(version=self.first,subdivision=sub,goods=original.goods,type_coeff=original.type_coeff,segment=original.segment,variation_calculate=original.variation_calculate,motivation_coeff=.003,manager_coeff=1)
            other=Goods.objects.create(goods_key='Other',planning_group_sales='Other',group='Other',brand='Other')
            SubdivisionCoeff.objects.create(version=self.first,subdivision=sub,goods=other,type_coeff=original.type_coeff,segment=original.segment,variation_calculate=original.variation_calculate,motivation_coeff=.007,manager_coeff=1)
            return sub.pk,original.goods_id,first.pk,other.pk
        sub,good,first_id,other=self.db(fixture)
        self.page.goto(self.live_server_url+reverse('coeff_subdivisions_admin_motivation',args=[sub]))
        first=self.page.locator(f'[name="sub_coeff={first_id}"]')
        expect(first.locator('..')).to_contain_text('Глобально: 0,2%')
        self.page.get_by_role('button',name='Изменить ставки',exact=True).click()
        self.page.locator(f'.bulk-good[value="{good}"]').check()
        self.page.locator('#coefficient-search').fill('Other')
        self.page.locator('#multiply-filtered').click()
        expect(self.page.locator('.bulk-good:checked')).to_have_count(2)
        self.page.locator('#multiply-operation').select_option('global')
        self.page.locator('#multiply-preview').click()
        expect(self.page.locator('#multiply-error')).to_contain_text('нет точной глобальной ставки')
        expect(first).to_have_value('0.3%')
        self.page.locator(f'.bulk-good[value="{other}"]').uncheck()
        self.page.locator('#multiply-preview').click()
        expect(self.page.locator('#multiply-review-summary')).to_contain_text('скрыто сейчас: 1')
        self.page.locator('#multiply-apply').click()
        expect(first).to_have_value('0.2%')
        self.assertEqual(self.db(lambda:SubdivisionCoeff.objects.get(pk=first_id).motivation_coeff),.003)
        self.page.locator('#coefficient-reset').click()
        self.page.locator('#editor-undo').click();expect(first).to_have_value('0.3%')
        expect(first.locator('..')).to_contain_text('Своя ставка')
        self.assertEqual(self.errors,[])
