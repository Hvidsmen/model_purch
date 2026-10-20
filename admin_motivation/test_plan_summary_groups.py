from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from django.core.management import call_command
from django.db import connection, transaction
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook
from django.core.files.uploadedfile import SimpleUploadedFile
from . import test_sales_plans as plan_fixtures
from .models import SalesPlanLine, SalesPlanScenario, GlobalCoeff
from .services.plan_summary import summary_report, summary_filters
from .services.plan_summary_groups import ensure_summary_groups
from .services.versions import create_version



from django.test import TestCase
class PlanSummaryGroupsTests(TestCase):
    setUp = plan_fixtures.SalesPlanTests.setUp
    load = plan_fixtures.SalesPlanTests.load
    rows = plan_fixtures.SalesPlanTests.rows

    def test_warm_report_and_filters_never_read_daily_plan(self):
        self.load()
        self.scenario.lines.filter(subdivision='A').update(amount_ddp_usd=Decimal('600'))
        expected = summary_report(self.scenario, {})
        with CaptureQueriesContext(connection) as queries:
            actual = summary_report(self.scenario, {})
            summary_filters(self.scenario, {})
        self.assertEqual(actual, expected)
        self.assertFalse(any('FROM "admin_motivation_salesplanline"' in q['sql'] for q in queries))
        self.assertEqual(self.scenario.summary_groups.count(), 2)

    def test_coefficients_change_without_rebuilding_plan_inputs(self):
        self.load()
        before = summary_report(self.scenario, {})['summary_nodes'][0]
        GlobalCoeff.objects.filter(version=self.version).update(motivation_coeff=.3)
        with CaptureQueriesContext(connection) as queries:
            after = summary_report(self.scenario, {})['summary_nodes'][0]
        self.assertEqual(before['global'], Decimal('225'))
        self.assertEqual(after['global'], Decimal('900'))
        self.assertFalse(any('FROM "admin_motivation_salesplanline"' in q['sql'] for q in queries))

    def test_new_boundary_rebuilds_once_and_preserves_all_amounts(self):
        self.load()
        original = self.scenario.lines.get(subdivision='A')
        original.pk = None
        original.plan_date = date(2026, 10, 10)
        original.save()
        self.assertEqual(ensure_summary_groups(self.scenario).count(), 2)
        future = create_version(date(2026, 10, 8), 'New boundary', self.version.pk)
        GlobalCoeff.objects.filter(version=future).update(motivation_coeff=.3)
        with CaptureQueriesContext(connection) as queries:
            result = summary_report(self.scenario, {})
        self.assertTrue(any('FROM "admin_motivation_salesplanline"' in q['sql'] for q in queries))
        self.assertEqual(self.scenario.summary_groups.count(), 3)
        self.assertEqual(result['summary_nodes'][0]['plan'], Decimal('2500'))
        self.assertEqual(result['summary_nodes'][0]['global'], Decimal('825'))
        with CaptureQueriesContext(connection) as queries:
            self.assertEqual(summary_report(self.scenario, {}), result)
        self.assertFalse(any('FROM "admin_motivation_salesplanline"' in q['sql'] for q in queries))

    def test_all_orm_input_mutations_invalidate_and_outputs_do_not(self):
        self.load()
        def current():
            ensure_summary_groups(self.scenario)
            self.scenario.refresh_from_db()
            self.assertTrue(self.scenario.summary_signature)
        def invalid():
            self.scenario.refresh_from_db()
            self.assertEqual(self.scenario.summary_signature, '')
        current()
        line = self.scenario.lines.get(subdivision='A')
        line.amount_ddp_usd = Decimal('600')
        line.save(update_fields=['amount_ddp_usd'])
        invalid(); current()
        self.scenario.lines.filter(pk=line.pk).update(amount_ddp_usd=Decimal('700'))
        invalid(); current()
        line.amount_ddp_usd = Decimal('800')
        SalesPlanLine.objects.bulk_update([line], ['amount_ddp_usd'])
        invalid(); current()
        self.scenario.lines.update(total_usd=10, calculation={'audit': 'updated'})
        self.scenario.refresh_from_db()
        self.assertTrue(self.scenario.summary_signature)
        line.pk = None
        line.plan_date = date(2026, 10, 2)
        SalesPlanLine.objects.bulk_create([line])
        invalid(); current()
        line.delete()
        invalid(); current()
        self.scenario.lines.filter(subdivision='A').delete()
        invalid(); current()
        self.assertEqual(summary_report(self.scenario, {})['summary_nodes'][0]['plan'], Decimal('500'))

    def test_rollback_preserves_materialized_inputs(self):
        self.load()
        expected = summary_report(self.scenario, {})
        ids = list(self.scenario.summary_groups.values_list('pk', flat=True))
        try:
            with transaction.atomic():
                self.scenario.lines.filter(subdivision='A').delete()
                ensure_summary_groups(self.scenario)
                raise ValueError('Rollback')
        except ValueError:
            pass
        self.assertEqual(summary_report(self.scenario, {}), expected)
        self.assertEqual(list(self.scenario.summary_groups.values_list('pk', flat=True)), ids)

    def test_reimport_replaces_groups_and_command_can_prepare_old_plan(self):
        self.load()
        expected = summary_report(self.scenario, {})
        self.load()
        self.assertEqual(summary_report(self.scenario, {}), expected)
        SalesPlanScenario.objects.filter(pk=self.scenario.pk).update(summary_signature='')
        out = StringIO()
        call_command('rebuild_plan_summary', plan=self.scenario.pk, stdout=out)
        self.assertIn('групп 2', out.getvalue())
        self.scenario.refresh_from_db()
        self.assertTrue(self.scenario.summary_signature)

    def test_template_preserves_missing_cells_and_existing_settings(self):
        from .models import VariationCalculate
        from .views import import_current_coefficients
        from .services.coefficient_template import coefficient_template
        self.load()
        first = GlobalCoeff.objects.filter(version=self.version).order_by('pk').first()
        custom = VariationCalculate.objects.create(variation_name='Keep current level')
        GlobalCoeff.objects.filter(pk=first.pk).update(variation_calculate=custom, manager_coeff=.71)
        missing = GlobalCoeff.objects.filter(version=self.version).exclude(pk=first.pk).first()
        missing_pk = missing.pk
        missing.delete()
        before = list(GlobalCoeff.objects.filter(version=self.version).order_by('pk').values_list('pk', 'motivation_coeff', 'variation_calculate_id', 'manager_coeff'))
        payload = coefficient_template(self.version)
        import_current_coefficients(SimpleUploadedFile('current.xlsx', payload), self.version)
        self.assertFalse(GlobalCoeff.objects.filter(pk=missing_pk).exists())
        self.assertEqual(list(GlobalCoeff.objects.filter(version=self.version).order_by('pk').values_list('pk', 'motivation_coeff', 'variation_calculate_id', 'manager_coeff')), before)

    def test_templates_keep_current_director_values_and_missing_cells(self):
        from .models import DirectorEmployee, GlobalDirectorCoeff, SubdivisionDirectorCoeff, Goods
        self.load()
        employee = DirectorEmployee.objects.create(full_name='Template director')
        employee.subdivisions.add(self.subs[0])
        GlobalDirectorCoeff.objects.create(version=self.version, goods=self.good,
            **{f'k{i}': Decimal('.012345') for i in range(5)})
        SubdivisionDirectorCoeff.objects.create(version=self.version, goods=self.good,
            subdivision=self.subs[0], employee=employee,
            **{f'k{i}': Decimal('-.023456') for i in range(5)})
        only_director = Goods.objects.create(goods_key='only-director', planning_group_sales='Other', group='*', brand='*')
        GlobalDirectorCoeff.objects.create(version=self.version, goods=only_director,
            **{f'k{i}': Decimal('.076543') for i in range(5)})
        for route, args, employee_id, expected in [
                ('motivation_global_template', [], '', Decimal('.012345')),
                ('motivation_subdivision_template', [self.subs[0].pk], employee.pk, Decimal('-.023456'))]:
            response = self.client.get(reverse(route, args=args), {'version': self.version.pk, 'employee': employee_id})
            book = load_workbook(BytesIO(response.content))
            row = next(row for row in book.active.iter_rows(min_row=3, values_only=True) if row[1] == 'Sales')
            self.assertEqual([Decimal(str(value)) for value in row[14:]], [expected] * 5)
            book.close()
            before_global = list(GlobalCoeff.objects.filter(version=self.version).order_by('pk').values_list('pk', 'motivation_coeff'))
            posted = self.client.post(reverse('sub_act', args=args) if args else reverse('gb_act'),
                {'version': self.version.pk, 'action_button': 'file_load', 'employee': employee_id,
                 'myfile': SimpleUploadedFile('current.xlsx', response.content)})
            self.assertEqual(posted.status_code, 302)
            self.assertEqual(list(GlobalCoeff.objects.filter(version=self.version).order_by('pk').values_list('pk', 'motivation_coeff')), before_global)
            current = SubdivisionDirectorCoeff.objects.get(version=self.version, goods=self.good, employee=employee) if args else GlobalDirectorCoeff.objects.get(version=self.version, goods=self.good)
            self.assertEqual(current.k0, expected)
        self.assertFalse(GlobalCoeff.objects.filter(goods=only_director).exists())

    def test_template_downloads_are_generated_and_roundtrip_both_importers(self):
        self.load()
        for url_name, args, action_name in [('motivation_global_template', [], 'gb_act'),
                ('motivation_subdivision_template', [self.subs[0].pk], 'sub_act')]:
            url = reverse(url_name, args=args)
            response = self.client.get(url, {'version': self.version.pk})
            self.assertEqual(response.status_code, 200)
            self.assertIn('attachment;', response['Content-Disposition'])
            book = load_workbook(BytesIO(response.content), data_only=False)
            self.assertEqual(book.active.max_column, 19)
            self.assertIn('E1:I1', {str(region) for region in book.active.merged_cells.ranges})
            self.assertEqual(book.active['E1'].value, 'Политики')
            self.assertEqual(book.active['J1'].value, 'Продажи')
            self.assertEqual(book.active['O1'].value, 'Директор')
            self.assertEqual(book.active['E2'].value, 'K0')
            self.assertEqual(book.active.freeze_panes, 'E3')
            self.assertFalse(book.active.sheet_view.showGridLines)
            self.assertEqual(book.active['E3'].number_format, '0.00%')
            self.assertNotEqual(book.active['E1'].fill.fgColor, book.active['J1'].fill.fgColor)
            self.assertEqual(book.active.cell(1, 2).value, 'Группа планов продаж')
            self.assertEqual(book.active.cell(3, 2).value, 'Sales')
            self.assertEqual(book.active.cell(3, 5).value, .1 if not args else .2)
            book.close()
            upload = SimpleUploadedFile('coefficients.xlsx', response.content,
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            posted = self.client.post(reverse(action_name, args=args),
                {'version': self.version.pk, 'action_button': 'file_load', 'myfile': upload})
            self.assertEqual(posted.status_code, 302)
        global_page = self.client.get(reverse('global_coeff_admin_motivation'), {'version': self.version.pk})
        self.assertContains(global_page, reverse('motivation_global_template'))
        local_page = self.client.get(reverse('coeff_subdivisions_admin_motivation', args=[self.subs[0].pk]), {'version': self.version.pk})
        self.assertContains(local_page, reverse('motivation_subdivision_template', args=[self.subs[0].pk]))
        self.assertEqual(self.client.get(reverse('motivation_global_template'), {'version': 'bad'}).status_code, 404)
        self.assertEqual(self.client.post(reverse('motivation_global_template')).status_code, 405)
        self.client.logout()
        self.assertEqual(self.client.get(reverse('motivation_global_template')).status_code, 302)
