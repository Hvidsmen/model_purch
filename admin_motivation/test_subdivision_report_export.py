from decimal import Decimal
from io import BytesIO
from unittest.mock import Mock, patch
from urllib.parse import parse_qs, urlsplit
from django.core.exceptions import ValidationError
from django.http import QueryDict
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from base.testing import authorize_test_case
from .services.reports import parameters, load_report, subdivision_report, choices
from .services.subdivision_excel import subdivision_excel


class SubdivisionReportExportTests(TestCase):
    def setUp(self):
        authorize_test_case(self)

    @staticmethod
    def rows():
        from .test_reports import NativeMotivationReportTests
        base = dict(NativeMotivationReportTests.division_rows()[0], Год=2026, ТипКоэф='Политики', КварталМотивации='4 Квартал(1-4 кв)')
        return [base, dict(base, КварталМотивации='Год', СуммаМотивацииUSD=50,
                          СуммаМотивацииUSDУчетПроцента=40, Остаток=5)]

    def test_period_filter_is_validated_and_bound_before_aggregation(self):
        values = parameters('subdivisions', QueryDict('period=4&period=5&period=4'))
        self.assertEqual(values['PremiumPeriods'], ['4', '5'])
        with self.assertRaises(ValidationError):
            parameters('subdivisions', QueryDict('period=4);DROP TABLE x'))
        with self.assertRaises(ValidationError):
            parameters('subdivisions', QueryDict('period=1'))
        cursor, connection = Mock(), Mock()
        cursor.description = [('Подразделение',), ('Год',)]
        cursor.fetchmany.return_value = []
        with patch('admin_motivation.services.reports.connect_database', return_value=(connection, cursor)):
            load_report('subdivisions', values)
        sql, *args = cursor.execute.call_args.args
        self.assertEqual(args[-2:], ['4', '5'])
        self.assertIn("END IN (?,?)", sql)
        self.assertLess(sql.index('END IN (?,?)'), sql.index('GROUP BY'))
        self.assertNotIn('PremiumPeriods', sql)
        connection.close.assert_called_once()

    def test_lookup_periods_have_labels_and_preserve_selection(self):
        cursor, connection = Mock(), Mock()
        cursor.description = [('value',)]
        cursor.fetchmany.return_value = []
        with patch('admin_motivation.services.reports.connect_database', return_value=(connection, cursor)):
            filters = choices('subdivisions', parameters('subdivisions', QueryDict('period=4')), refresh=True)
        period = next(item for item in filters if item['name'] == 'period')
        self.assertEqual([item['label'] for item in period['options']], ['4 квартал', 'Год'])
        self.assertEqual([item['value'] for item in period['options'] if item['selected']], ['4'])

    def test_excel_tables_match_ui_totals_and_keep_numeric_and_literal_cells(self):
        rows = [dict(row, Подразделение='=1+1') for row in self.rows()]
        data = subdivision_report(rows)
        self.assertEqual(data['years'], ['2026'])
        values = parameters('subdivisions', QueryDict('period=4&period=5'))
        for scope, expected in [('summary', ['Мотивация подразделений']), ('detail', ['Детализация мотивации']),
                                ('all', ['Мотивация подразделений', 'Детализация мотивации', 'Параметры отчёта'])]:
            with subdivision_excel(data, values, scope) as output:
                book = load_workbook(output)
                self.assertEqual(book.sheetnames, expected)
                for sheet in book:
                    self.assertIn('Год расчёта: 2026', sheet['A2'].value)
                    self.assertIn('4 квартал, Год', sheet['A2'].value)
                if scope != 'detail':
                    sheet = book['Мотивация подразделений']
                    self.assertEqual([cell.value for cell in sheet[4]][2:], data['summary_totals'])
                    self.assertEqual(sheet['A5'].value, '=1+1')
                    self.assertEqual(sheet['A5'].data_type, 's')
                    self.assertEqual(sheet['C4'].data_type, 'n')
                if scope != 'summary':
                    sheet = book['Детализация мотивации']
                    headers = [cell.value for cell in sheet[3]]
                    index = headers.index('4 Квартал(1-4 кв)\n% выполнения периода') + 1
                    self.assertEqual(sheet.cell(4, index).value, .8)
                    self.assertEqual(sheet.cell(4, index).number_format, '0.00%')
                    self.assertGreater(sheet.max_row, 5)
                book.close()

    def test_report_header_links_and_all_three_downloads_preserve_filters(self):
        url = reverse('motivation_report', args=['subdivisions'])
        values = {'run': '1', 'period': ['4', '5'], 'subdivision': ['A', 'B'], 'scope': 'matrix'}
        with patch('admin_motivation.report_views.choices', return_value=[]), patch('admin_motivation.report_views.load_report', side_effect=lambda *args, **kwargs: subdivision_report(self.rows(), 'matrix')) as load:
            response = self.client.get(url, values)
            self.assertContains(response, 'Год расчёта: 2026')
            self.assertContains(response, 'Скачать весь отчёт Excel')
            self.assertContains(response, 'Скачать итоговую таблицу Excel')
            self.assertContains(response, 'Скачать детализацию Excel')
            for table, target in response.context['data']['excel_urls'].items():
                query = parse_qs(urlsplit(target).query)
                self.assertEqual(query['period'], ['4', '5'])
                self.assertEqual(query['subdivision'], ['A', 'B'])
                self.assertEqual(query['scope'], ['matrix'])
                export = self.client.get(url + target)
                self.assertEqual(export.status_code, 200)
                book = load_workbook(BytesIO(b''.join(export.streaming_content)))
                self.assertEqual(len(book.sheetnames), 3 if table == 'all' else 1)
                book.close()
                export.close()
            self.assertEqual(load.call_args.args[1]['PremiumPeriods'], ['4', '5'])
            self.assertEqual(load.call_args.args[3], 'matrix')

    def test_multi_year_header_uses_data_years_without_inventing_current_year(self):
        data = subdivision_report([dict(self.rows()[0], Год=2024), dict(self.rows()[0], Год=2025)])
        self.assertEqual(data['years'], ['2024', '2025'])
        self.assertTrue(all('2024' in value or '2025' in value for value in data['periods']))
