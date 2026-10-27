"""New report tab; no redirects or replacements of the legacy report pages."""
from tempfile import TemporaryFile
from zoneinfo import ZoneInfo
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.paginator import Paginator
from django.http import FileResponse, Http404
from django.shortcuts import render, get_object_or_404
from django.urls import reverse
from django.views.decorators.http import require_GET
from base.access import role_for, REPORT_ROLES, PERSONAL_ROLES, ADMIN
from .models import MotivationExecution, SalesPlanScenario, Subdivision
from .services import reports_next as reports


@require_GET
def index(request):
    return render(request, 'admin_motivation/reports_next/index.html', {'catalog': reports.available(request.user)})


def report_navigation(user, kind, query):
    """Carry shared conditions only; personal/product drilldowns stay on their report."""
    from django.http import QueryDict
    shared = QueryDict(mutable=True)
    for field in ('execution', 'year', 'period', 'type', 'subdivision', 'currency'):
        if query.getlist(field):
            shared.setlist(field, query.getlist(field))
    if query.get('run') == '1':
        shared['run'] = '1'
    captions = {'mine': 'Моя премия', 'team': 'Команда', 'registry': 'Реестр',
                'executive': 'Сравнение плана', 'detail': 'Детальный расчёт'}
    return [{'key': key, 'label': captions[key], 'active': key == kind,
             'url': reverse('motivation_report_next', args=[key]) + ('?' + shared.urlencode() if shared else '')}
            for key in reports.available(user)]


def workbook(kind, data, metadata, query=None):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    book = Workbook(write_only=True)
    sheet = book.create_sheet('Премии')
    sheet.freeze_panes = 'B3'
    def append(values, percent_columns=()):
        cells = []
        for position, value in enumerate(values):
            cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub('', value)[:32767] if isinstance(value, str) else value)
            if isinstance(value, str):
                cell.data_type = 's'
            elif value is not None:
                cell.number_format = '0.00%' if position in percent_columns or kind == 'executive' and position == 4 else '#,##0.00'
            cells.append(cell)
        sheet.append(cells)
    append([reports.CATALOG[kind][0]])
    if kind == 'executive':
        append(['Подразделение / вид мотивации / директор', 'План: глобальные коэффициенты, USD', 'План: коэффициенты подразделений, USD', 'Разница, USD', 'Разница, %'])
        def write(nodes, indent=0):
            for node in nodes:
                append(['  ' * indent + node['name'], node['global'], node['local'], node['difference'], node.get('difference_percent')])
                write(node.get('children', []), indent + 1)
        write(data)
    elif kind in ('team', 'registry'):
        from .services.reports_next_presenter import present
        display = present(kind, data, data, query.getlist('period') or ['all'] if query is not None else 'all')
        append(['Подразделение', 'Автор сделки', 'Вид мотивации', 'Валюта', 'Премия по расчёту',
                'Годовая доплата за кварталы', 'Всего по расчёту',
                *[column['label'] + ' · без годовой доплаты' for column in display['period_columns']], 'Статус'])
        for row in display['employees' if kind == 'team' else 'registry_rows']:
            append([row['subdivision'], row['author'], ', '.join(row['kinds']), data['currency'], row['amount'],
                    row['bonus'], row['total'], *row['period_values'],
                    'Не распределено' if row['unallocated'] else 'Рассчитано; выплата не подтверждена'])
        if kind == 'team':
            sheet = book.create_sheet('Итого по видам')
            append(['Вид мотивации', 'Премия по расчёту', 'Годовая доплата за кварталы', 'Всего по расчёту', 'Валюта'])
            for row in display['type_totals']:
                append([row['kind'], row['amount'], row['bonus'], row['total'], data['currency']])
    else:
        append(['Подразделение', 'Автор сделки', 'Вид мотивации', 'Период', 'Валюта', 'Премия по расчёту', 'Годовая доплата за кварталы', 'Всего по расчёту', 'Статус'])
        for row in data['rows']:
            append([row['subdivision'], row['author'], row['kind'], row['period'], row['currency'], row['amount'], row['bonus'], row['total'], 'Не распределено' if row['unallocated'] else 'Рассчитано; выплата не подтверждена'])
    if kind == 'mine' and query is not None:
        from .services.reports_next_presenter import personal_details
        personal = personal_details(data, query)
        sheet = book.create_sheet('Цели и потенциал')
        append(['Вид', 'Период', 'База цели', 'План мотивации, USD', 'Факт мотивации, USD', 'Выполнение', 'Осталось до 100%, USD', 'Премия при 100%', 'Почему не рассчитана', 'Возможная премия на текущих отгрузках', 'Валюта премии', 'Премия по расчёту', 'Пояснение итога'])
        for goal in personal['goals_with_total']:
            append([goal['kind'], goal['period'], 'Подразделение' if goal['global_scope'] else 'База мотивации',
                goal['plan'], goal['fact'], goal['ratio'], None if goal['first_quarter'] else goal['remaining'],
                goal['reward_at_target'], goal['reward_reason'], goal['possible'], data['currency'], goal['calculated'], goal.get('basis_note', '')], percent_columns=(5,))
        sheet = book.create_sheet('Товарная детализация')
        append(['Вид', 'Подразделение', 'Автор сделки', 'Группа планов', 'Группа ERP', 'Бренд', 'Период', 'Всего по расчёту', 'Валюта'])
        def products(nodes, section, group, path=()):
            for node in nodes:
                current = (*path, node['label'])
                if node['children']:
                    products(node['children'], section, group, current)
                else:
                    for period, value in zip(section['periods'], node['values'][:-1]):
                        append([section['kind'], group['branch'], group['author'], *current, period, value, data['currency']])
        for section in personal['product_sections']:
            for group in section['groups']:
                products(group['children'], section, group)
        sheet = book.create_sheet('Товарная матрица')
        all_periods = sorted({period for section in personal['product_sections'] for period in section['periods']}, key=lambda value: (reports.period_code(value), value))
        append(['Вид', 'Подразделение', 'Автор сделки', 'Уровень', 'Товарная матрица',
                *[period + ' · ' + metric for period in all_periods for metric in ('Возможная премия', 'Осталось до 100%, USD', 'Выполнение', 'Премия по расчёту')],
                'Итого · Возможная премия', 'Итого · Осталось до 100%, USD', 'Итого · Выполнение', 'Итого · Премия по расчёту',
                'Годовая доплата за кварталы', 'Всего по расчёту'])
        def matrix_row(section, group, row, label, level):
            lookup = dict(zip(section['periods'], row['metrics']))
            values = []
            for period in all_periods:
                metric = lookup.get(period)
                values.extend([metric['possible'], metric['remaining'], metric['ratio'], metric['calculated']] if metric else [None] * 4)
            append([section['kind'], group['branch'], group['author'], level, label, *values, row['total_metrics']['possible'], row['total_metrics']['remaining'], row['total_metrics']['ratio'], row['total_metrics']['calculated'], row['bonus'], row['total']], percent_columns=tuple(7 + index * 4 for index in range(len(all_periods) + 1)))
        def matrix_nodes(section, group, nodes):
            for row in nodes:
                matrix_row(section, group, row, row['label'], {2: 'Группа планов', 3: 'Группа ERP', 4: 'Бренд'}[row['level']])
                matrix_nodes(section, group, row['children'])
        for section in personal['product_sections']:
            for group in section['groups']:
                matrix_row(section, group, group, 'Всего по расчёту', 'Итог')
                matrix_nodes(section, group, group['children'])
    notes = book.create_sheet('Параметры')
    from .services.reports import PILOT_BONUS_NOTE
    metadata = [*metadata, ('Доплата за кварталы', PILOT_BONUS_NOTE)]
    for label, value in metadata:
        cells = [WriteOnlyCell(notes, value=ILLEGAL_CHARACTERS_RE.sub('', str(item))[:32767]) for item in (label, value)]
        for cell in cells:
            cell.data_type = 's'
        notes.append(cells)
    # Explicit strings also protect parameter values from Excel formula evaluation.
    # WriteOnlyCell is required because write-only sheets do not expose cells.
    output = TemporaryFile(suffix='.xlsx')
    try:
        book.save(output)
        output.seek(0)
        return output
    except Exception:
        output.close()
        raise
    finally:
        book.close()


@require_GET
def report(request, kind):
    if kind not in reports.CATALOG:
        raise Http404
    if kind not in reports.available(request.user):
        raise PermissionDenied
    query = request.GET.copy()
    report_user = request.user
    preview_accounts, preview_user = [], None
    admin_preview = role_for(request.user, 'motivation') == ADMIN and kind in ('mine', 'detail')
    if admin_preview:
        from django.contrib.auth import get_user_model
        from base.report_access import politician_scope
        candidates = get_user_model().objects.filter(is_active=True, motivation_report_access__isnull=False).select_related('motivation_report_access').prefetch_related('groups').order_by('last_name', 'first_name', 'username')
        preview_accounts = [user for user in candidates if role_for(user, 'motivation') in PERSONAL_ROLES and politician_scope(user)]
        if query.get('preview_user'):
            preview_user = next((user for user in preview_accounts if str(user.pk) == query['preview_user']), None)
            if preview_user is None:
                raise PermissionDenied
            report_user = preview_user
    else:
        query.pop('preview_user', None)
    try:
        scope, types, period = reports.restrict(report_user, query)
    except ValidationError as exc:
        from django.http import HttpResponseBadRequest
        return HttpResponseBadRequest(' '.join(exc.messages))
    runs = MotivationExecution.objects.filter(state='completed').order_by('-started_at', '-pk')
    selected = get_object_or_404(runs, pk=query['execution']) if query.get('execution', '').isdigit() else runs.first()
    plans = SalesPlanScenario.objects.filter(is_deleted=False, loaded_at__isnull=False).order_by('-loaded_at', '-pk')
    original_plan = reports.default_plan(selected) if selected else None
    plan = get_object_or_404(plans, pk=query['plan']) if query.get('plan', '').isdigit() else original_plan
    data, error, detail = None, '', None
    if selected:
        query.setdefault('year', str(selected.parameters.get('fact_year', '')))
    query.setdefault('director_scope', 'active')
    query.setdefault('currency', 'USD')
    query.setdefault('period', 'all')
    metadata = [('Расчёт', selected.completed_at.astimezone(ZoneInfo('Europe/Moscow')).strftime('%d.%m.%Y %H:%M МСК') if selected and selected.completed_at else 'Нет'), ('Год', query.get('year', '')),
        ('План расчёта DWH', selected.parameters.get('plan_version', '') if selected else ''),
        ('План сравнения' if kind == 'executive' else 'Загруженная копия плана расчёта', str(plan) if plan else 'Не загружена'), ('Период премии', reports.period_label(query)),
        ('Виды мотивации', ', '.join(types)),
        ('Валюта', 'USD' if kind == 'executive' else query['currency']),
        ('Порог выполнения для выплаты', selected.parameters.get('payout_threshold', 'Не сохранён') if selected else ''),
        ('Статус выплат', 'Данных о согласовании и фактической оплате нет. Показаны рассчитанные суммы.')]
    if query.get('search', '').strip() and kind in ('team', 'registry'):
        metadata.append(('Поиск сотрудника', query['search'].strip()))
    if kind == 'executive':
        metadata.append(('Состав директоров', 'Действующие привязки к подразделениям' if query['director_scope'] == 'active' else 'Все сохранённые коэффициенты, включая архивные'))
    if preview_user:
        metadata.insert(1, ('Личный отчёт сотрудника', preview_user.motivation_report_access.deal_author))
    try:
        if query.get('run') == '1' or query.get('export') == 'xlsx':
            if kind == 'mine' and admin_preview and preview_user is None:
                raise ValidationError('Выберите сотрудника для предпросмотра личной премии.')
            if not selected:
                raise ValidationError('Нет подтверждённых завершённых расчётов.')
            if kind == 'executive':
                if not plan:
                    raise ValidationError('План из расчёта не найден среди загруженных до запуска. Выберите загруженный план сравнения.')
                if query['year'] != str(selected.parameters.get('fact_year')):
                    raise ValidationError('Год должен совпадать с годом выбранного расчёта.')
                data = reports.planned_budget(selected, plan, query)
            else:
                with reports.current_run(selected):
                    data = reports.actual_reports(report_user, selected, query, detailed=kind in ('detail', 'mine'))
                if kind in ('team', 'registry') and query.get('search', '').strip():
                    from decimal import Decimal
                    term = query['search'].strip().casefold().replace('ё', 'е')
                    data['rows'] = [row for row in data['rows'] if term in row['author'].casefold().replace('ё', 'е')]
                    data['total'] = sum((row['total'] for row in data['rows']), Decimal(0))
                    data['has_unallocated'] = any(row['unallocated'] for row in data['rows'])
                for row in data['rows']:
                    target = query.copy()
                    target['type'], target['period'], target['run'] = row['kind'], row['code'], '1'
                    target.setlist('subdivision', [row['subdivision']])
                    target.setlist('author', [row['author']])
                    row['detail_url'] = reverse('motivation_report_next', args=['detail']) + '?' + target.urlencode()
                if kind == 'detail':
                    from .services.reports_next_presenter import new_detail_labels
                    from .services.manager_detail import combine_reports
                    detail = new_detail_labels(combine_reports(data['sections'], query['currency'])) if data['sections'] else None
            if query.get('export') == 'xlsx':
                if kind == 'detail' and detail:
                    from .services.report_excel import report_excel
                    output = report_excel(detail, {'Currency': query['currency'], 'Years': [query['year']], 'TypeCoeff_': query.getlist('type'), 'Subdivision': query.getlist('subdivision'), 'DealAuthor': query.getlist('author'), 'gp_sales': query.getlist('planning'), 'GPGoods': query.getlist('goods'), 'Brend': query.getlist('brand')}, query.get('table', 'all'), provenance=metadata)
                else:
                    output = workbook(kind, data, metadata, query=query)
                return FileResponse(output, as_attachment=True, filename=f'motivation-next-{kind}.xlsx', content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            if detail:
                detail['excel_urls'] = {}
                for table in ('all', 'summary', 'detail'):
                    target = query.copy()
                    target['export'], target['table'] = 'xlsx', table
                    detail['excel_urls'][table] = '?' + target.urlencode()
    except ValidationError as exc:
        error = ' '.join(exc.messages)
    except PermissionDenied:
        raise
    except Exception:
        import logging
        logging.getLogger(__name__).exception('New motivation report failed: %s', kind)
        error = 'Не удалось получить данные. Проверьте подключение к DWH и журнал сервера.'
    selected_subs = query.getlist('subdivision')
    branch_names = scope.get('subdivision', []) if scope else list(Subdivision.objects.values_list('subdivision_key', flat=True))
    branch_names = [branch_names] if isinstance(branch_names, str) else branch_names
    branches = [{'name': value, 'selected': value in selected_subs} for value in branch_names]
    period_options = [item for item in reports.PERIODS if not scope or not scope.get('pilot_periods') or item[0] in ('all', '4', '5')]
    export_query = query.copy()
    export_query['export'] = 'xlsx'
    reset_query = query.copy()
    for field in ('period', 'type', 'subdivision', 'author', 'planning', 'goods', 'brand', 'page', 'export', 'table', 'search'):
        reset_query.pop(field, None)
    branch_label = ', '.join(selected_subs) if len(selected_subs) <= 2 else f'Подразделений: {len(selected_subs)}'
    from .services.reports_next_presenter import present
    presentation = present(kind, data, data if kind != 'executive' else None, query.getlist('period') or ['all'])
    if kind == 'mine' and data:
        from .services.reports_next_presenter import personal_details
        presentation.update(personal_details(data, query))
        for card in presentation['period_cards']:
            goals = [goal for goal in presentation['goals'] if reports.period_code(goal['period']) == reports.period_code(card['label'])]
            if not card['amount'] and goals:
                card['explanation'] = ('Порог выплаты не достигнут. Остаток до порога показан в целях ниже.' if all(goal['under_threshold'] for goal in goals) else 'На выбранных отгрузках возможная премия равна 0.' if all(not goal['possible'] for goal in goals) else 'Результат расчёта — 0. Проверьте ставки и условия в детализации.')
    ledger_page = Paginator(presentation['registry_rows'], 100).get_page(query.get('page')) if data is not None and kind not in ('executive', 'detail') else None
    employee_page = Paginator(presentation['employees'], 50).get_page(query.get('page')) if kind == 'team' else None
    page_query = query.copy()
    page_query.pop('page', None)
    page_query.pop('export', None)
    all_types = scope.get('type', []) if scope else ['Политики', 'Продажи', 'Директора']
    all_types = [all_types] if isinstance(all_types, str) else all_types
    return render(request, 'admin_motivation/reports_next/report.html', {'kind': kind, 'title': reports.CATALOG[kind][0],
        'admin_preview': admin_preview, 'preview_accounts': preview_accounts, 'preview_user': preview_user, 'catalog': reports.available(request.user), 'data': detail if kind == 'detail' else data, 'ledger': data,
        'current_execution_id': runs.first().pk if runs.exists() else None, 'runs': runs, 'execution': selected, 'plans': plans, 'plan': plan, 'query': query, 'types': types,
        'reset_url': '?' + reset_query.urlencode(), 'branch_label': branch_label or 'Все доступные подразделения', 'selected_branches': selected_subs, 'all_types': all_types, 'type_choices': [{'value': value, 'selected': value in query.getlist('type')} for value in all_types], 'type_label': ', '.join(query.getlist('type')) or 'Все доступные виды',
        'product_filters': [(field, value) for field in ('planning', 'goods', 'brand') for value in query.getlist(field)], 'period_label': reports.period_label(query), 'period_choices': [{'value': code, 'label': label, 'selected': code in query.getlist('period')} for code, label in period_options if code != 'all'], 'periods': period_options, 'branches': branches, 'metadata': metadata, 'error': error, 'export_url': '?' + export_query.urlencode(),
        'report_navigation': report_navigation(request.user, kind, query),
        'presentation': presentation, 'employee_page': employee_page, 'ledger_page': ledger_page, 'page_query': page_query.urlencode(), 'selected_parameters': [], 'calculation_steps': [{'period': label, 'possible': block[13], 'ratio': block[16], 'calculated': block[17], 'bonus': block[18], 'total': block[19]} for label, block in zip(detail['periods'], detail['detail_totals'])] if detail else [],
        'is_personal': role_for(request.user, 'motivation') in REPORT_ROLES})
