from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from base.testing import authorize_test_case
from .models import GlobalCoeffVersion, Goods, Chanel, Subdivision, DirectorEmployee, SubdivisionDirectorCoeff, GlobalDirectorCoeff
from .services.director_coeff_import import import_director_coefficients, read_director_blocks
from .services.versions import create_version


def workbook(names=None, values=None):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Коэффициенты'
    sheet.cell(1, 5, 'Первый директор')
    sheet.cell(2, 5, 'Первый директор')  # Same name twice within a block is one employee.
    sheet.cell(1, 10, 'Второй директор')
    for row, column, name in names or []:
        sheet.cell(row, column, name)
    for column, text in enumerate(['№', 'Группа планов продаж', 'Группа ERP', 'Марка / бренд'], 1):
        sheet.cell(4, column, text)
    for start in (5, 10):
        for index in range(5):
            sheet.cell(5, start + index, f'K{index}')
    sheet.append([1, '1. RAC', 'Group', '*', *(values or [.001, .002, .003, .004, .005]), .01, .02, .03, .04, .05])
    stream = BytesIO()
    book.save(stream)
    return SimpleUploadedFile('directors.xlsx', stream.getvalue())


class DirectorCoeffImportTests(TestCase):
    def setUp(self):
        authorize_test_case(self)
        self.version = GlobalCoeffVersion.objects.get(effective_from=date(2001, 1, 1))
        channel = Chanel.objects.create(chanel_name='Channel')
        self.a = Subdivision.objects.create(subdivision_key='A', subdivision_name='A', subdivision_global='A', chanel=channel)
        self.b = Subdivision.objects.create(subdivision_key='B', subdivision_name='B', subdivision_global='B', chanel=channel)
        self.first = DirectorEmployee.objects.create(full_name='Первый директор')
        self.second = DirectorEmployee.objects.create(full_name='Второй директор')
        self.first.subdivisions.add(self.a, self.b)
        self.second.subdivisions.add(self.a)
        self.good = Goods.objects.create(goods_key='G', planning_group_sales='1. RAC', group='Group', brand='*')
        self.url = reverse('motivation_director_coeff_import')

    def test_import_all_branches_and_idempotent_update_preserves_other_scopes(self):
        GlobalDirectorCoeff.objects.create(version=self.version, goods=self.good, k0=Decimal('.8'))
        other = Goods.objects.create(goods_key='Other', planning_group_sales='Other', group='Other', brand='Other')
        preserved = SubdivisionDirectorCoeff.objects.create(version=self.version, subdivision=self.a, employee=self.first, goods=other, k0=Decimal('.9'))
        self.assertEqual(import_director_coefficients(workbook(), self.version.pk),
                         {'directors': 2, 'created': 3, 'updated': 0, 'unchanged': 0})
        for branch in (self.a, self.b):
            row = SubdivisionDirectorCoeff.objects.get(version=self.version, subdivision=branch, employee=self.first, goods=self.good)
            self.assertEqual(row.k0, Decimal('.001'))
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(version=self.version, employee=self.second, goods=self.good).k4, Decimal('.05'))
        preserved.refresh_from_db()
        self.assertEqual(preserved.k0, Decimal('.9'))
        self.assertEqual(GlobalDirectorCoeff.objects.get().k0, Decimal('.8'))
        self.assertEqual(import_director_coefficients(workbook(), self.version.pk)['unchanged'], 3)
        result = import_director_coefficients(workbook(values=['0,2%', '.2', '-.1', 0, '.123456789']), self.version.pk)
        self.assertEqual(result['updated'], 2)
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(subdivision=self.b, employee=self.first, goods=self.good).k4, Decimal('.123457'))

    def test_unknown_ambiguous_or_unbound_employee_is_atomic(self):
        for upload in [workbook(names=[(1, 5, 'Неизвестный директор'), (2, 5, '')])]:
            with self.assertRaisesMessage(ValidationError, 'не найден'):
                import_director_coefficients(upload, self.version.pk)
            self.assertFalse(SubdivisionDirectorCoeff.objects.exists())
        DirectorEmployee.objects.create(full_name='Первый директор')
        with self.assertRaisesMessage(ValidationError, 'несколько сотрудников'):
            import_director_coefficients(workbook(), self.version.pk)
        DirectorEmployee.objects.filter(full_name='Первый директор').exclude(pk=self.first.pk).delete()
        self.second.subdivisions.clear()
        with self.assertRaisesMessage(ValidationError, 'нет привязанных'):
            import_director_coefficients(workbook(), self.version.pk)
        self.assertFalse(SubdivisionDirectorCoeff.objects.exists())

    def test_bad_values_and_unmatched_matrix_do_not_write_anything(self):
        for values in ([None, .2, .3, .4, .5], ['=1/100', .2, .3, .4, .5], [2, .2, .3, .4, .5]):
            with self.assertRaises(ValidationError):
                import_director_coefficients(workbook(values=values), self.version.pk)
            self.assertFalse(SubdivisionDirectorCoeff.objects.exists())
        self.good.brand = 'Other'
        self.good.save()
        with self.assertRaisesMessage(ValidationError, 'товарная матрица'):
            import_director_coefficients(workbook(), self.version.pk)
        self.assertFalse(SubdivisionDirectorCoeff.objects.exists())

    def test_blank_block_preserves_coefficients_and_cross_block_names_are_rejected(self):
        import_director_coefficients(workbook(), self.version.pk)
        row = SubdivisionDirectorCoeff.objects.get(version=self.version, subdivision=self.a, employee=self.first, goods=self.good)
        upload = workbook(values=['', '', '', '', ''])
        self.assertEqual(import_director_coefficients(upload, self.version.pk)['unchanged'], 1)
        row.refresh_from_db()
        self.assertEqual(row.k0, Decimal('.001'))
        with self.assertRaisesMessage(ValidationError, 'нескольких блоках'):
            import_director_coefficients(workbook(names=[(1, 10, 'Первый директор')]), self.version.pk)
        self.assertEqual(SubdivisionDirectorCoeff.objects.filter(version=self.version).count(), 3)

    def test_historical_version_cannot_be_changed(self):
        latest = create_version(date(2025, 1, 1), title='Selected', source_version=self.version.pk)
        with self.assertRaisesMessage(ValidationError, 'последний черновик'):
            import_director_coefficients(workbook(), self.version.pk)
        self.assertEqual(import_director_coefficients(workbook(), latest.pk)['created'], 3)
        self.assertFalse(SubdivisionDirectorCoeff.objects.filter(version=self.version).exists())

    def test_database_failure_rolls_back_new_records(self):
        with patch('admin_motivation.services.director_coeff_import.SubdivisionDirectorCoeff.objects.bulk_update', side_effect=ValidationError('Write failed')):
            with self.assertRaises(ValidationError):
                import_director_coefficients(workbook(), self.version.pk)
        self.assertFalse(SubdivisionDirectorCoeff.objects.exists())

    def test_view_uses_selected_version_and_reports_errors_and_counts(self):
        page = self.client.get(self.url, {'version': self.version.pk})
        self.assertContains(page, 'Выбранная версия:')
        response = self.client.post(self.url, {'version': self.version.pk, 'file': workbook()})
        self.assertRedirects(response, reverse('global_coeff_admin_motivation') + f'?version={self.version.pk}', fetch_redirect_response=False)
        self.assertContains(self.client.get(response.url), 'создано: 3')
        response = self.client.post(self.url, {'version': self.version.pk, 'file': workbook(names=[(1, 5, 'Неизвестный'), (2, 5, '')])})
        self.assertContains(response, 'Файл не загружен')
        self.assertContains(response, 'Неизвестный')
        self.assertEqual(self.client.get(self.url).status_code, 404)
        self.assertEqual(self.client.get(reverse('global_coeff_admin_motivation')).status_code, 200)

    def test_invalid_workbook_is_a_readable_error(self):
        for upload in (None, SimpleUploadedFile('bad.xlsx', b'bad'), SimpleUploadedFile('bad.csv', b'bad')):
            with self.assertRaises(ValidationError):
                read_director_blocks(upload)
