from io import BytesIO
from unittest.mock import patch
from openpyxl import load_workbook
from django.test import TestCase
from django.urls import reverse
from django.core.exceptions import ValidationError
from base.testing import authorize_test_case
from .test_manager_detail import detail_rows
from .services.manager_detail import manager_detail_report
from .services.report_excel import report_excel


class ReportExcelTests(TestCase):
    def setUp(self):
        authorize_test_case(self)
        self.data = manager_detail_report(detail_rows(), 'RUB', detail_rows())
        self.parameters = {'Currency': 'RUB', 'Subdivision': ['A'], 'TypeCoeff_': 'Директора',
                           'DealAuthor': ['Иванов'], 'Years': ['2026']}

    def test_all_and_individual_tables_are_numeric_and_include_collapsed_nodes(self):
        for scope, names in [('all', ['Итого по мотивации', 'Расширенная мотивация', 'Параметры отчёта']),
                             ('summary', ['Итого по мотивации']), ('detail', ['Расширенная мотивация'])]:
            stream = report_excel(self.data, self.parameters, scope)
            try:
                book = load_workbook(stream)
                self.assertEqual(book.sheetnames, names)
                if scope != 'detail':
                    sheet = book['Итого по мотивации']
                    self.assertEqual(sheet['B5'].value, 'Иванов')
                    self.assertEqual(sheet['I3'].value, 19000)
                    self.assertEqual(sheet['I3'].data_type, 'n')
                    self.assertEqual(sheet['I4'].value, sheet['I3'].value)
                    self.assertEqual(sheet.freeze_panes, 'C3')
                if scope != 'summary':
                    sheet = book['Расширенная мотивация']
                    labels = [row[0].value for row in sheet.iter_rows(min_row=3)]
                    self.assertEqual(labels, ['% выполнения подразделений', 'Итого', 'A', 'Иванов', '1. RAC', 'Кондиционеры', '<Brand>'])
                    self.assertEqual(sheet['B5'].value, .015)
                    self.assertEqual(sheet['B5'].number_format, '0.00%')
                    self.assertEqual(sheet['S4'].value, 9000)
                    self.assertEqual(sheet.cell(4, 2 + len(self.data['periods']) * len(self.data['headers']) + 17).value, 18000)
                    self.assertEqual(sheet.freeze_panes, 'B3')
                if scope == 'all':
                    settings = list(book['Параметры отчёта'].values)
                    self.assertIn(('Валюта мотивации', 'RUB'), settings)
                    self.assertIn(('Автор сделки', 'Иванов'), settings)
                book.close()
            finally:
                stream.close()
        with self.assertRaises(ValidationError): report_excel(self.data, self.parameters, 'bad')

    def test_excel_text_is_not_a_formula_and_empty_report_is_valid(self):
        rows = detail_rows()
        for row in rows:
            row['АвторСделки'] = '=1+1'
            row['Марка'] = '=HYPERLINK("https://example.invalid")'
        data = manager_detail_report(rows)
        stream = report_excel(data, {'DealAuthor': ['=1+1']})
        try:
            book = load_workbook(stream, data_only=False)
            self.assertEqual(book['Итого по мотивации']['B5'].value, '=1+1')
            self.assertEqual(book['Итого по мотивации']['B5'].data_type, 's')
            for sheet in book:
                self.assertFalse(any(cell.data_type == 'f' for row in sheet for cell in row))
            book.close()
        finally:
            stream.close()
        stream = report_excel(manager_detail_report([]), {}, 'summary')
        try:
            book = load_workbook(stream)
            self.assertEqual(book.active['I3'].value, 0)
            book.close()
        finally:
            stream.close()

    def test_download_view_preserves_filters_and_skips_lookup_requests(self):
        url = reverse('motivation_report', args=['manager-detail'])
        with patch('admin_motivation.report_views.choices') as choices, \
             patch('admin_motivation.report_views.load_report', return_value=self.data) as load:
            response = self.client.get(url, {'export': 'xlsx', 'table': 'all', 'currency': 'RUB',
                'subdivision': ['A', 'B'], 'author': ['Иванов'], 'type': 'Директора', 'year': '2026'})
            self.assertEqual(response.status_code, 200)
            self.assertIn('motivation-manager-RUB-all.xlsx', response['Content-Disposition'])
            self.assertEqual(load.call_args.args[1]['Subdivision'], ['A', 'B'])
            self.assertEqual(load.call_args.args[1]['DealAuthor'], ['Иванов'])
            choices.assert_not_called()
            book = load_workbook(BytesIO(b''.join(response.streaming_content)))
            self.assertEqual(len(book.sheetnames), 3)
            book.close(); response.close()
        with patch('admin_motivation.report_views.choices', return_value=[]), \
             patch('admin_motivation.report_views.load_report', return_value=self.data):
            page = self.client.get(url, {'run': '1', 'type': 'Директора', 'author': 'Иванов', 'currency': 'RUB'})
        self.assertContains(page, 'Скачать весь отчёт Excel')
        self.assertContains(page, 'Скачать итоговую таблицу Excel')
        self.assertContains(page, 'Скачать расширенную таблицу Excel')
        self.assertContains(page, 'table=summary')
        self.client.logout()
        self.assertEqual(self.client.get(url, {'export': 'xlsx'}).status_code, 302)
