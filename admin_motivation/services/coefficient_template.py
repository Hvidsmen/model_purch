"""Generate the two-header XLSX layout accepted by coefficient file import."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from ..models import GlobalCoeff, SubdivisionCoeff, SegmentCoeff, GlobalDirectorCoeff, SubdivisionDirectorCoeff


def coefficient_template(version, subdivision=None, employee=None):
    segments = list(SegmentCoeff.objects.order_by('pk').values_list('pk', 'segment_name'))
    model = SubdivisionCoeff if subdivision else GlobalCoeff
    rows = model.objects.filter(version=version)
    if subdivision:
        rows = rows.filter(subdivision=subdivision)
    grouped = {}
    for goods, pg, group, brand, kind, segment, value in rows.order_by().values_list(
            'goods_id', 'goods__planning_group_sales', 'goods__group', 'goods__brand',
            'type_coeff__type_coeff_name', 'segment_id', 'motivation_coeff').iterator(chunk_size=2000):
        if kind not in ('Политики', 'Продажи'):
            continue
        record = grouped.setdefault((pg, group, brand, goods), {})
        record[kind, segment] = value
    directors = (SubdivisionDirectorCoeff.objects.filter(version=version, subdivision=subdivision, employee=employee)
                 if subdivision and employee else GlobalDirectorCoeff.objects.filter(version=version)
                 if not subdivision else SubdivisionDirectorCoeff.objects.none())
    for goods, pg, group, brand, *values in directors.order_by().values_list(
            'goods_id', 'goods__planning_group_sales', 'goods__group', 'goods__brand', *[f'k{i}' for i in range(5)]):
        record = grouped.setdefault((pg, group, brand, goods), {})
        for index, value in enumerate(values):
            record['Директор', index] = float(value)
    book = Workbook()
    sheet = book.active
    sheet.title = 'Коэффициенты'
    sheet.append(['Шаблон текущих коэффициентов: политики, продажи и директор. Значения от -1 до 1. Первые две строки не удалять.'])
    sheet.append(['№', 'Группа планов продаж', 'Группа ERP', 'Марка / бренд',
                  *[f'{kind} {name}' for kind in ('Политики', 'Продажи') for pk, name in segments], *[f'Директор K{i}' for i in range(5)]])
    for index, (key, cells) in enumerate(sorted(grouped.items()), 1):
        sheet.append([index, *key[:3], *[cells.get((kind, segment))
            for kind in ('Политики', 'Продажи') for segment, name in segments], *[cells.get(('Директор', i)) for i in range(5)]])
        # Excel must treat external classification names as literal text.
        for column in range(2, 5):
            sheet.cell(sheet.max_row, column).data_type = 's'
    sheet.freeze_panes = 'E3'
    sheet.auto_filter.ref = f'A2:{sheet.cell(2, 9 + 2 * len(segments)).column_letter}{max(sheet.max_row, 2)}'
    for cell in sheet[2]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='28466B')
        cell.alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[2].height = 32
    for column, width in [('A', 8), ('B', 28), ('C', 32), ('D', 24)]:
        sheet.column_dimensions[column].width = width
    for column in range(5, 10 + 2 * len(segments)):
        sheet.column_dimensions[sheet.cell(2, column).column_letter].width = 18
        for row in range(3, sheet.max_row + 1):
            sheet.cell(row, column).number_format = '0.00%'
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()
