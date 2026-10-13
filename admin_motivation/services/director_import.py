from io import BytesIO
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError
from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook, Workbook
from openpyxl.utils.exceptions import InvalidFileException
from ..models import DirectorEmployee, Subdivision
from ..director_forms import DirectorEmployeeForm


def template():
    book = Workbook()
    book.active.append(['ФИО', 'Должность', 'Филиалы'])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def import_directors(upload):
    if not upload or not upload.name.lower().endswith('.xlsx'):
        raise ValidationError('Выберите файл Excel .xlsx.')
    if upload.size > 5 * 1024 * 1024:
        raise ValidationError('Максимальный размер файла — 5 МБ.')
    raw = upload.read()
    try:
        with ZipFile(BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValidationError('Файл Excel слишком большой после распаковки.')
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
    except (BadZipFile, ValueError, KeyError, OSError, ParseError, InvalidFileException):
        raise ValidationError('Не удалось прочитать файл Excel.')
    try:
        rows = book.active.iter_rows(values_only=True)
        headers = [str(value or '').strip() for value in next(rows, ())]
        if 'ФИО' not in headers or 'Филиалы' not in headers:
            raise ValidationError('Нужны колонки «ФИО» и «Филиалы». Скачайте шаблон.')
        branches = list(Subdivision.objects.all())
        prepared, seen = [], set()
        for number, row in enumerate(rows, 2):
            if number > 2001:
                raise ValidationError('Можно загрузить не больше 2000 строк.')
            if not any(value is not None for value in row):
                continue
            def cell(label):
                index = headers.index(label)
                return str(row[index] or '').strip() if index < len(row) else ''
            name, branch_text = cell('ФИО'), cell('Филиалы')
            if name.casefold() in seen:
                raise ValidationError(f'Строка {number}: ФИО повторяется в файле.')
            seen.add(name.casefold())
            selected = []
            for token in filter(None, (value.strip() for value in branch_text.split(';'))):
                matches = [branch for branch in branches if token.casefold() in (branch.subdivision_key.casefold(), branch.subdivision_name.casefold())]
                if len(matches) != 1:
                    raise ValidationError(f'Строка {number}: филиал «{token}» не найден или неоднозначен.')
                selected.append(matches[0].pk)
            existing = list(DirectorEmployee.objects.filter(full_name__iexact=name)[:2])
            if len(existing) > 1:
                raise ValidationError(f'Строка {number}: несколько директоров с таким ФИО. Измените их вручную.')
            position = cell('Должность') if 'Должность' in headers else (existing[0].position if existing else '')
            form = DirectorEmployeeForm({'full_name': name, 'position': position, 'subdivisions': selected}, instance=existing[0] if existing else None)
            if name.startswith('=') or not form.is_valid():
                raise ValidationError(f'Строка {number}: проверьте ФИО и филиалы.')
            prepared.append(form)
        if not prepared:
            raise ValidationError('В файле нет директоров.')
        with transaction.atomic():
            for form in prepared:
                form.save()
        return len(prepared)
    finally:
        book.close()
