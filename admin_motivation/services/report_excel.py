"""Streaming XLSX exports of the displayed detailed manager report."""
from tempfile import TemporaryFile
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.core.exceptions import ValidationError
from .reports import FILTERS

MONEY = '#,##0.00'
PERCENT = '0.00%'


def report_excel(data, parameters, scope='all'):
    if scope not in ('all', 'summary', 'detail'):
        raise ValidationError('Неизвестная таблица для выгрузки.')
    book = Workbook(write_only=True)
    book.properties.title = 'Мотивация по менеджеру'
    book.properties.creator = 'Рабочий портал'
    currency = data['currency']
    edge = Side(style='thin', color='DCE3EC')

    def cell(sheet, value, *, kind='text', heading=False, total=False, indent=0, fill=None):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub('', value)[:32767]
        result = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            result.data_type = 's'  # External labels are never Excel formulas.
        result.font = Font(name='Calibri', size=11, color='312E81' if total else '1E3A5F', bold=heading or total)
        result.fill = PatternFill('solid', fgColor=fill or ('EDE9FE' if total else 'F5F3FF' if heading else 'FFFFFF'))
        result.alignment = Alignment(horizontal='right' if kind != 'text' else 'left', vertical='center',
                                     wrap_text=heading, indent=min(indent, 15))
        result.border = Border(bottom=edge)
        if kind != 'text':
            result.number_format = PERCENT if kind == 'percent' else MONEY
        return result

    def sheet(name, headers, widths, freeze):
        if len(headers) > 16384:
            raise ValidationError('Слишком много столбцов для Excel. Уточните годы отчёта.')
        target = book.create_sheet(name)
        target.sheet_view.showGridLines = False
        target.freeze_panes = freeze
        target.sheet_properties.pageSetUpPr.fitToPage = True
        target.page_setup.orientation = 'landscape'
        target.page_setup.fitToWidth = 1
        target.page_setup.fitToHeight = 0
        target.print_title_rows = '1:2'
        for index, width in enumerate(widths, 1):
            target.column_dimensions[get_column_letter(index)].width = width
        target.row_dimensions[1].height = 30
        target.row_dimensions[2].height = 65
        target.append([cell(target, f'{name} · {currency}', heading=True)])
        target.append([cell(target, title, heading=True) for title in headers])
        return target

    def append(target, values, kinds, *, index, total=False, indent=0):
        if index > 1048576:
            raise ValidationError('Слишком много строк для Excel. Уточните фильтры.')
        target.row_dimensions[index].height = 22
        target.row_dimensions[index].outlineLevel = min(indent, 7)
        target.append([cell(target, value, kind=kind, total=total, indent=indent if column == 0 else 0)
                       for column, (value, kind) in enumerate(zip(values, kinds))])

    try:
        if scope in ('all', 'summary'):
            summary = data['period_summary']
            headers = ['Подразделение', 'Автор сделки', *[column['label'] for column in summary['columns']], 'Доплаты за квартал', 'Итого']
            target = sheet('Итого по мотивации', headers, [28, 42] + [18] * (len(headers) - 2), 'C3')
            kinds = ['text', 'text'] + ['money'] * (len(headers) - 2)
            totals = summary['totals']
            append(target, ['Итого по компании', '', *totals['values'], totals['remainder'], totals['total']], kinds, index=3, total=True)
            index = 3
            for group in summary['groups']:
                index += 1
                append(target, [group['subdivision'], 'Итого по подразделению', *group['values'], group['remainder'], group['total']], kinds, index=index, total=True)
                for row in group['authors']:
                    index += 1
                    append(target, [row['subdivision'], row['author'], *row['values'], row['remainder'], row['total']], kinds, index=index, indent=1)
            target.auto_filter.ref = f'A2:{get_column_letter(len(headers))}{index}'

        if scope in ('all', 'detail'):
            periods = [*data['periods'], 'Итого']
            titles = [header['label'] + (' (ср. взв.)' if index < 5 else '') for index, header in enumerate(data['headers'])]
            headers = ['Подразделение / автор / группа планов / группа товаров / марка',
                       *[f'{period}\n{title}' for period in periods for title in titles]]
            target = sheet('Расширенная мотивация', headers, [48] + [19] * (len(headers) - 1), 'B3')
            kinds = ['text'] + [
                'percent' if index < 5 or index in (12, 16) else 'money'
                for period in periods for index in range(len(titles))]
            index = 2
            def write_row(label, blocks, *, total=False, level=0):
                nonlocal index
                index += 1
                append(target, [label, *[value for block in blocks for value in block]], kinds,
                       index=index, total=total, indent=level)
            if data.get('division_execution'):
                write_row('% выполнения подразделений', data['division_execution'])
            write_row('Итого', data['detail_totals'], total=True)
            def write_nodes(nodes):
                for node in nodes:
                    write_row(node['label'], node['values'], level=node['level'])
                    write_nodes(node['children'])
            write_nodes(data['nodes'])
            target.auto_filter.ref = f'A2:{get_column_letter(len(headers))}{index}'

        if scope == 'all':
            target = sheet('Параметры отчёта', ['Параметр', 'Значение'], [38, 95], 'A3')
            settings = [('Валюта мотивации', currency), ('Годы данных', ', '.join(data['years']) or 'Нет данных')]
            for _, parameter, caption, _, multiple in FILTERS['manager-detail']:
                value = parameters.get(parameter)
                settings.append((caption, ', '.join(value) if isinstance(value, list) and value else str(value) if value else 'Все'))
            settings += [('План и отгрузки', 'USD'), ('K0–K4', 'Σ(Ki × Oi) / Σ(Oi), вес — отгрузки соответствующего сегмента'),
                         ('Годовое выполнение', 'План и факт из строк «Год»; мотивация и остаток — по всем периодам')]
            for index, row in enumerate(settings, 3):
                append(target, row, ['text', 'text'], index=index)
        output = TemporaryFile(mode='w+b')
        try:
            book.save(output)
            output.seek(0)
            return output
        except Exception:
            output.close()
            raise
    finally:
        book.close()
