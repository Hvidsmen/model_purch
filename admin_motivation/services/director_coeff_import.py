"""Import blocks of director names and K0–K4 from a matrix workbook."""
from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import unicodedata
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from ..models import DirectorEmployee, Goods, SubdivisionDirectorCoeff
from .directors import KEYS, lock_editable


def label(value):
    return ' '.join(unicodedata.normalize('NFC', str(value or '')).replace('\\n', ' ').split())


def identity(value):
    return label(value).casefold()


def coefficient(value):
    if isinstance(value, bool) or value is None:
        raise ValueError
    text = str(value).strip().replace(',', '.')
    result = Decimal(text.rstrip('%')) / (100 if text.endswith('%') else 1)
    if not result.is_finite() or not -1 <= result <= 1:
        raise ValueError
    return result.quantize(Decimal('.000001'), rounding=ROUND_HALF_UP)


def read_director_blocks(upload):
    if not upload or not upload.name.lower().endswith('.xlsx'):
        raise ValidationError('Выберите файл Excel .xlsx.')
    if upload.size > 5 * 1024 * 1024:
        raise ValidationError('Максимальный размер файла — 5 МБ.')
    raw = upload.read()
    book = None
    try:
        with ZipFile(BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValidationError('Файл слишком большой после распаковки.')
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        if 'Коэффициенты' not in book.sheetnames:
            raise ValidationError('В файле нужен лист «Коэффициенты».')
        sheet = book['Коэффициенты']
        if sheet.max_row > 5000 or sheet.max_column > 200:
            raise ValidationError('Допускается до 5000 строк и 200 столбцов.')
        rows = list(sheet.iter_rows(values_only=True))
        header = next((index for index, row in enumerate(rows[:50])
                       if len(row) >= 4 and label(row[1]) == 'Группа планов продаж'
                       and label(row[2]) in ('Группа ERP', 'Группа товаров') and label(row[3]) == 'Марка / бренд'), None)
        if header is None:
            raise ValidationError('Не найдена шапка товарной матрицы.')
        segments = next((index for index in range(header, min(header + 3, len(rows)))
                         if [label(value).upper() for value in rows[index][4:9]] == ['K0', 'K1', 'K2', 'K3', 'K4']), None)
        if segments is None:
            raise ValidationError('После шапки нужны блоки K0–K4.')
        blocks = []
        errors = []
        for column in range(4, len(rows[segments]) - 4, 5):
            if [label(value).upper() for value in rows[segments][column:column + 5]] != ['K0', 'K1', 'K2', 'K3', 'K4']:
                continue
            names = list({identity(row[column]): label(row[column]) for row in rows[:header] if len(row) > column and label(row[column])}.values())
            entries, seen = [], set()
            for row_number, row in enumerate(rows[segments + 1:], segments + 2):
                raw_values = row[column:column + 5]
                if not any(value is not None and str(value).strip() for value in raw_values):
                    continue
                matrix = tuple(identity(value) for value in row[1:4])
                if len(matrix) != 3 or not all(matrix):
                    errors.append(f'Строка {row_number}: заполните группу планов, группу ERP и марку.')
                    continue
                if matrix in seen:
                    errors.append(f'Строка {row_number}: товарная строка повторяется в блоке.')
                seen.add(matrix)
                try:
                    values = tuple(coefficient(value) for value in raw_values)
                    if len(values) != 5:
                        raise ValueError
                except (ValueError, InvalidOperation):
                    errors.append(f'Строка {row_number}, блок {column + 1}: заполните все K0–K4 числами от -1 до 1. Формулы сохраните как значения.')
                    continue
                entries.append((matrix, values, row_number))
            if entries and not names:
                errors.append(f'Над блоком со столбца {column + 1} нет ФИО директоров.')
            if entries:
                blocks.append((names, entries))
        if errors:
            raise ValidationError(errors[:30])
        if not blocks:
            raise ValidationError('В файле нет заполненных блоков коэффициентов директоров.')
        return blocks
    except ValidationError:
        raise
    except Exception as error:
        raise ValidationError('Не удалось прочитать Excel. Выберите корректный файл .xlsx.') from error
    finally:
        if book:
            book.close()


def import_director_coefficients(upload, version_id):
    blocks = read_director_blocks(upload)
    with transaction.atomic():
        version = lock_editable(version_id)
        employees, goods = defaultdict(list), defaultdict(list)
        for employee in DirectorEmployee.objects.filter(is_archived=False).prefetch_related('subdivisions'):
            employees[identity(employee.full_name)].append(employee)
        for good in Goods.objects.all():
            goods[tuple(identity(value) for value in (good.planning_group_sales, good.group, good.brand))].append(good)
        errors, prepared, seen, imported_employees = [], {}, set(), set()
        for names, entries in blocks:
            selected = []
            for name in names:
                key = identity(name)
                if key in seen:
                    errors.append(f'Директор «{name}» указан в нескольких блоках.')
                    continue
                seen.add(key)
                matches = employees[key]
                if len(matches) != 1:
                    errors.append(f'ФИО «{name}»: ' + ('директор не найден.' if not matches else 'найдено несколько сотрудников.'))
                    continue
                employee = matches[0]
                subdivisions = list(employee.subdivisions.all())
                if not subdivisions:
                    errors.append(f'У директора «{name}» нет привязанных подразделений.')
                    continue
                selected.append((employee, subdivisions))
            for matrix, values, row_number in entries:
                matches = goods[matrix]
                if len(matches) != 1:
                    errors.append(f'Строка {row_number}: товарная матрица «{" / ".join(matrix)}» не найдена или неоднозначна.')
                    continue
                for employee, subdivisions in selected:
                    imported_employees.add(employee.pk)
                    for subdivision in subdivisions:
                        prepared[subdivision.pk, employee.pk, matches[0].pk] = values
        if errors:
            raise ValidationError(list(dict.fromkeys(errors))[:30])
        existing = {(row.subdivision_id, row.employee_id, row.goods_id): row
                    for row in SubdivisionDirectorCoeff.objects.filter(version=version, employee_id__in=imported_employees)}
        created, updated = [], []
        for (sub_id, employee_id, goods_id), values in prepared.items():
            row = existing.get((sub_id, employee_id, goods_id))
            if row and tuple(getattr(row, key) for key in KEYS) == values:
                continue
            if row is None:
                row = SubdivisionDirectorCoeff(version=version, subdivision_id=sub_id, employee_id=employee_id, goods_id=goods_id)
                created.append(row)
            else:
                updated.append(row)
            for key, value in zip(KEYS, values):
                setattr(row, key, value)
        SubdivisionDirectorCoeff.objects.bulk_create(created, batch_size=100)
        SubdivisionDirectorCoeff.objects.bulk_update(updated, KEYS, batch_size=100)
        return {'directors': len(imported_employees), 'created': len(created), 'updated': len(updated),
                'unchanged': len(prepared) - len(created) - len(updated)}


def director_coefficient_template(version):
    """Blank workbook accepted by read_director_blocks, with real lookup lists."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.utils import get_column_letter

    goods = Goods.objects.order_by('planning_group_sales', 'group', 'brand', 'pk')
    if goods.count() > 4995:
        raise ValidationError('В справочнике больше 4995 товарных строк. Для загрузки разделите матрицу на файлы.')
    book = Workbook()
    sheet = book.active
    sheet.title = 'Коэффициенты'
    sheet.freeze_panes = 'E6'

    def text(cell, value):
        cell.value = str(value or '')
        cell.data_type = 's'
        cell.alignment = Alignment(vertical='center', wrap_text=True)

    for row, instruction in enumerate([
        '1. Выберите ФИО в жёлтых ячейках справа — один директор на строку.',
        '2. Заполните все K0–K4 для нужных товаров. Пример: 0,2% = 0,002.',
        'Пустой блок строки не меняет ставки. Подробности — на листе «Инструкция».',
    ], 1):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        text(sheet.cell(row, 1), instruction)
        sheet.row_dimensions[row].height = 34
        sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=9)
        sheet.cell(row, 5).fill = PatternFill('solid', fgColor='FFF4CC')
    for column, caption in enumerate(['№', 'Группа планов продаж', 'Группа ERP', 'Марка / бренд'], 1):
        text(sheet.cell(4, column), caption)
        sheet.cell(4, column).font = Font(bold=True, color='FFFFFF')
        sheet.cell(4, column).fill = PatternFill('solid', fgColor='51428C')
    sheet.merge_cells('E4:I4')
    text(sheet['E4'], 'Коэффициенты директоров')
    sheet['E4'].font = Font(bold=True, color='FFFFFF')
    for index in range(5):
        sheet.cell(4, 5 + index).fill = PatternFill('solid', fgColor='51428C')
        cell = sheet.cell(5, 5 + index, f'K{index}')
        cell.font = Font(bold=True, color='51428C')
        cell.fill = PatternFill('solid', fgColor='F0ECFF')
    sheet.row_dimensions[4].height = 30
    for number, (pg, group, brand) in enumerate(goods.values_list('planning_group_sales', 'group', 'brand'), 1):
        row = number + 5
        sheet.cell(row, 1, number)
        for column, value in enumerate((pg, group, brand), 2):
            text(sheet.cell(row, column), value)
        for column in range(5, 10):
            sheet.cell(row, column).number_format = '0.0000%'
            sheet.cell(row, column).fill = PatternFill('solid', fgColor='F7F5FF')
        sheet.row_dimensions[row].height = 26
    for column, width in [(1, 6), (2, 28), (3, 38), (4, 26)]:
        sheet.column_dimensions[get_column_letter(column)].width = width
    for column in range(5, 10):
        sheet.column_dimensions[get_column_letter(column)].width = 13
    coefficients = DataValidation(type='decimal', operator='between', formula1='-1', formula2='1', allow_blank=True)
    coefficients.showErrorMessage = True
    coefficients.errorTitle = 'Недопустимая ставка'
    coefficients.error = 'Введите ставку от −100% до 100%, например 0,2%.'
    sheet.add_data_validation(coefficients)
    coefficients.add(f'E6:I{max(6, sheet.max_row)}')

    directory = book.create_sheet('Директора')
    directory.append(['ФИО для выбора', 'Привязанные подразделения'])
    employees = DirectorEmployee.objects.filter(is_archived=False).prefetch_related('subdivisions').order_by('full_name', 'pk')
    for row, employee in enumerate(employees, 2):
        text(directory.cell(row, 1), employee.full_name)
        text(directory.cell(row, 2), ', '.join(str(branch) for branch in employee.subdivisions.all()))
    directory.column_dimensions['A'].width = 48
    directory.column_dimensions['B'].width = 90
    directory.freeze_panes = 'A2'
    if directory.max_row >= 2:
        book.defined_names.add(DefinedName('DirectorNames', attr_text=f"'Директора'!$A$2:$A${directory.max_row}"))
        names = DataValidation(type='list', formula1='=DirectorNames', allow_blank=True)
        names.showErrorMessage = True
        names.errorTitle = 'Выберите ФИО из справочника'
        names.error = 'Сначала добавьте директора и его подразделения в справочник портала.'
        sheet.add_data_validation(names)
        names.add('E1:E3')

    guide = book.create_sheet('Инструкция')
    instructions = [
        ('Версия на портале', str(version)),
        ('ФИО', 'Выберите директора в E1, E2 или E3. Несколько имён означают одинаковые ставки для всех этих директоров и всех их привязанных подразделений.'),
        ('Товары', 'Товарные группы и марки заполнены из справочника портала. Не меняйте их написание. Заполняйте только нужные строки.'),
        ('Ставки', 'Заполните все пять K0–K4 или оставьте все пять ячеек строки пустыми. 0,2% равно 0,002; число 0,2 означает 20%. Ноль означает обнуление ставки.'),
        ('Разные ставки', 'Для другого набора ставок скопируйте весь блок E:I в J:N. Выберите других директоров над новым блоком. Один директор не должен повторяться в разных блоках.'),
        ('Ограничения', 'Не более 200 столбцов и 5000 строк. Названия листа «Коэффициенты» и заголовков K0–K4 сохраняйте. Формулы замените значениями.'),
        ('Проверка', 'Перед загрузкой проверьте лист «Директора»: импорт использует действующие привязки портала, а не изменения этого справочного листа.'),
        ('Сохранение', 'Загрузка доступна в последний редактируемый черновик. Изменяются только заполненные ставки директоров. Политики, продажи и глобальные ставки сохраняются.'),
    ]
    for row, (topic, instruction) in enumerate(instructions, 1):
        text(guide.cell(row, 1), topic)
        guide.cell(row, 1).font = Font(bold=True, color='51428C')
        text(guide.cell(row, 2), instruction)
        guide.row_dimensions[row].height = 48
    guide.column_dimensions['A'].width = 25
    guide.column_dimensions['B'].width = 110
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()
