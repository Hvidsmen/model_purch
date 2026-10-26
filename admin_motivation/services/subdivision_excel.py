"""Export the same grouped subdivision report used by the browser."""
from tempfile import TemporaryFile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from django.core.exceptions import ValidationError
from .reports import PILOT_BONUS_NOTE, FILTERS, PREMIUM_PERIODS

METRICS = [*[f'Отгрузки O{i}, USD' for i in range(5)], 'План периода, USD', 'Факт периода, USD',
    '% выполнения периода', 'Возможная мотивация, USD', 'План мотивации, USD', 'Факт мотивации, USD',
    '% выполнения мотивации', 'С учётом выполнения, USD', 'Доплата за кварталы, USD', 'Итого, USD']


def subdivision_excel(data, parameters, table='all', scope='subdivisions'):
    if table not in ('all', 'summary', 'detail'):
        raise ValidationError('Неизвестная таблица для выгрузки.')
    book = Workbook(write_only=True)
    book.properties.title = 'Мотивация по подразделениям'
    book.properties.creator = 'Рабочий портал'
    years = ', '.join(data.get('years', [])) or 'Нет данных'
    periods = dict(PREMIUM_PERIODS)
    selection = ', '.join(periods[value] for value in parameters.get('PremiumPeriods', [])) or 'Все'
    caption = f'Год расчёта: {years} · Период премии: {selection} · USD'

    def cell(target, value, heading=False, total=False, percent=False, indent=0):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub('', value)[:32767]
        result = WriteOnlyCell(target, value=value)
        if isinstance(value, str) and 'доплат' in value.casefold():
            result.comment = Comment(PILOT_BONUS_NOTE, 'Портал')
        if isinstance(value, str):
            result.data_type = 's'
        result.font = Font(name='Calibri', size=11, bold=heading or total, color='24364B')
        result.fill = PatternFill('solid', fgColor='EDE9FE' if total else 'F3F5FA' if heading else 'FFFFFF')
        result.alignment = Alignment(wrap_text=heading, vertical='center', indent=min(indent, 15),
            horizontal='left' if isinstance(value, str) else 'right')
        if not isinstance(value, str):
            result.number_format = '0.00%' if percent else '#,##0.00'
        return result

    def sheet(name, headers, first_width=36, freeze='C4'):
        if len(headers) > 16384:
            raise ValidationError('Слишком много столбцов для Excel. Уточните периоды отчёта.')
        target = book.create_sheet(name)
        target.freeze_panes = freeze
        target.sheet_view.showGridLines = False
        target.print_title_rows = '1:3'
        target.page_setup.orientation = 'landscape'
        target.sheet_properties.pageSetUpPr.fitToPage = True
        target.page_setup.fitToWidth = 1
        for i in range(1, len(headers) + 1):
            target.column_dimensions[get_column_letter(i)].width = first_width if i == 1 else 22
        target.row_dimensions[1].height = 28
        target.row_dimensions[2].height = 34
        target.row_dimensions[3].height = 65
        for row in (1, 2):
            target.merged_cells.add(f'A{row}:{get_column_letter(len(headers))}{row}')
        target.append([cell(target, name, heading=True)])
        target.append([cell(target, caption, heading=True)])
        target.append([cell(target, value, heading=True) for value in headers])
        return target

    try:
        if table in ('all', 'summary'):
            headers = ['Подразделение', 'Тип коэффициента']
            for period in data['periods']:
                headers += [f'{period}\nВозможная мотивация, USD', f'{period}\nС учётом выполнения, USD']
            headers += ['Итого\nВозможная мотивация, USD', 'Итого\nС учётом выполнения, USD', 'Доплата за кварталы, USD', 'Итого, USD']
            target = sheet('Мотивация подразделений', headers)
            target.append([cell(target, value, total=True) for value in ['Итого', 'Все типы', *data['summary_totals']]])
            count = 4
            for row in data['summary']:
                target.append([cell(target, value, total=True) for value in [row['label'], 'Итого по подразделению', *row['values']]])
                count += 1
                for child in row['children']:
                    count += 1
                    target.row_dimensions[count].outlineLevel = 1
                    target.append([cell(target, value) for value in [row['label'], child['label'], *child['values']]])
            target.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{count}'
        if table in ('all', 'detail'):
            periods_all = [*data['periods'], 'Итого']
            metrics = [(index, title) for index, title in enumerate(METRICS)]
            columns = [(period, index, title) for p, period in enumerate(periods_all) for index, title in metrics
                       if p == len(periods_all)-1 or index != 13]
            headers = ['Товарная матрица' if scope == 'matrix' else 'Подразделение / товарная матрица',
                       *[f'{period}\n{title}' for period, _, title in columns]]
            target = sheet('Детализация мотивации', headers, 54, 'B4')
            count = 3
            def append(label, blocks, level=0, total=False):
                nonlocal count
                count += 1
                if count > 1048576:
                    raise ValidationError('Слишком много строк для Excel. Уточните фильтры.')
                target.row_dimensions[count].outlineLevel = min(level, 7)
                result = [cell(target, label, total=total, indent=level)]
                for p, block in enumerate(blocks):
                    for index, value in enumerate(block):
                        if p < len(blocks)-1 and index == 13:
                            continue
                        result.append(cell(target, value, total=total, percent=index in (7, 11)))
                target.append(result)
            append('Итого', data['detail_totals'], total=True)
            def walk(nodes):
                for node in nodes:
                    append(node['label'], node['values'], node['level'])
                    walk(node['children'])
            walk(data['nodes'])
            target.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{count}'
        if table == 'all':
            target = sheet('Параметры отчёта', ['Параметр', 'Значение'], 38, 'A4')
            settings = [('Год расчёта мотивации', years), ('Период премии', selection),
                ('Детализация', 'По товарной матрице' if scope == 'matrix' else 'По подразделениям')]
            for _, parameter, label, _, _ in FILTERS['subdivisions']:
                if parameter == 'PremiumPeriods':
                    continue
                values = parameters.get(parameter) or []
                settings.append((label, ', '.join(values) or 'Все'))
            settings.append(('Итоги', 'План и факт итога — из строк «Год»; мотивация и доплаты — по выбранным периодам.'))
            for pair in settings:
                target.append([cell(target, value) for value in pair])
        output = TemporaryFile(mode='w+b')
        try:
            book.save(output)
            output.seek(0)
            return output
        except Exception:
            output.close()
            raise
    finally:
        book.close()
