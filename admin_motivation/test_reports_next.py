from contextlib import nullcontext
from datetime import timedelta
from decimal import Decimal
from uuid import uuid4
from unittest.mock import Mock, patch
from django.test import TestCase
from django.http import QueryDict
from django.urls import reverse
from django.utils import timezone
from django.core.exceptions import ValidationError
from django.contrib.auth.models import User, Group
from openpyxl import load_workbook
from base.testing import authorize_admin
from base.access import FINANCIER, POLITICIAN, BRANCH_DIRECTOR
from base.models import MotivationReportAccess, PortalReportGroup
from .models import MotivationExecution, SalesPlanScenario, TypeCoeff
from .services import reports_next
from .services.manager_detail import manager_detail_report
from .test_manager_detail import detail_rows


class ReportsNextTests(TestCase):
    def setUp(self):
        from .test_sales_plans import SalesPlanTests
        from .services.approval import review_approval, approve_review
        from .services.workflow import export_snapshot
        fixture = SalesPlanTests(); fixture.setUp(); fixture.load()
        self.fixture = fixture
        self.actor = authorize_admin(self.client)
        approve_review(review_approval(fixture.version)['token'])
        self.run = MotivationExecution.objects.create(run_id=uuid4(), job_key=uuid4(), state='completed',
            started_at=timezone.now(), completed_at=timezone.now(), parameters={'fact_year': 2026,
            'plan_version': fixture.scenario.source_version, 'quarter_share': '80', 'year_share': '20', 'payout_threshold': '.8'},
            export_snapshot=export_snapshot())

    def test_new_tab_and_old_reports_remain_available(self):
        response = self.client.get(reverse('motivation_reports'))
        self.assertContains(response, 'Текущие отчёты')
        self.assertContains(response, 'Новые отчёты')
        response = self.client.get(reverse('motivation_reports_next'))
        self.assertContains(response, 'Сводный отчёт: план премий')
        self.assertContains(response, 'Реестр премий к выплате')
        self.assertContains(response, 'Моя премия')

    def test_plan_defaults_to_run_plan_and_compares_frozen_coefficients(self):
        self.assertEqual(reports_next.default_plan(self.run), self.fixture.scenario)
        query = QueryDict('year=2026&period=all')
        nodes = reports_next.planned_budget(self.run, self.fixture.scenario, query)
        self.assertEqual(nodes[0]['global'], Decimal(225))
        self.assertEqual(nodes[0]['local'], Decimal(345))
        self.assertEqual(nodes[0]['difference'], Decimal(120))
        annual = reports_next.planned_budget(self.run, self.fixture.scenario, QueryDict('period=5'))
        self.assertEqual(annual[0]['global'], Decimal(45))
        fourth = reports_next.planned_budget(self.run, self.fixture.scenario, QueryDict('period=4'))
        self.assertEqual(fourth[0]['global'], Decimal(180))
        self.fixture.good.brand = 'Changed classification'; self.fixture.good.save()
        with self.assertRaises(ValidationError):
            reports_next.planned_budget(self.run, self.fixture.scenario, query)

    def test_executive_page_uses_plan_selection_and_excel_numbers(self):
        url = reverse('motivation_report_next', args=['executive'])
        response = self.client.get(url, {'run': '1', 'period': 'all'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['plan'], self.fixture.scenario)
        self.assertContains(response, 'Разница между вариантами')
        self.assertContains(response, '120,00')
        exported = self.client.get(url, {'run': '1', 'period': 'all', 'export': 'xlsx'})
        from io import BytesIO
        book = load_workbook(BytesIO(b''.join(exported.streaming_content)))
        self.assertEqual(book['Премии']['B3'].value, 225)
        self.assertEqual(book['Премии']['C3'].value, 345)
        book.close(); exported.close()

    def test_current_run_guard_never_mislabels_another_or_running_result(self):
        cursor, connection = Mock(), Mock()
        cursor.fetchone.side_effect = [(0,), (self.run.run_id, 'completed')]
        with reports_next.current_run(self.run, lambda *args: (connection, cursor)):
            pass
        connection.close.assert_called_once()
        cursor.fetchone.side_effect = [(0,), (uuid4(), 'completed')]
        with self.assertRaises(ValidationError):
            with reports_next.current_run(self.run, lambda *args: (connection, cursor)):
                self.fail('Must not read another result')
        cursor.fetchone.side_effect = [(-1,)]
        with self.assertRaises(ValidationError):
            with reports_next.current_run(self.run, lambda *args: (connection, cursor)):
                self.fail('Must not read during a calculation')

    def test_actuals_filter_period_and_keep_bonus_separate(self):
        query = QueryDict('year=2026&type=Политики&period=5')
        data = manager_detail_report([row for row in detail_rows() if row['КварталМотивации'] == 'Год'])
        with patch('admin_motivation.services.reports_next.load_report', return_value=data) as load:
            result = reports_next.actual_reports(self.actor, self.run, query)
        self.assertEqual(result['total'], Decimal(95))
        self.assertEqual(len(result['rows']), 2)
        self.assertEqual(load.call_args.kwargs['premium_period'], '5')
        self.assertTrue(load.call_args.kwargs['summary_only'])
        self.assertEqual(result['sections'], [])

    def test_finance_registry_does_not_claim_payment_approval(self):
        user = User.objects.create_user('next-finance')
        user.groups.add(Group.objects.get(name=FINANCIER)); self.client.force_login(user)
        data = manager_detail_report(detail_rows())
        with patch('admin_motivation.services.reports_next.current_run', return_value=nullcontext()), patch('admin_motivation.services.reports_next.load_report', return_value=data):
            response = self.client.get(reverse('motivation_report_next', args=['registry']), {'run': '1', 'type': 'Политики'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'ещё не подтверждены к выплате')
        self.assertContains(response, 'Как рассчитано')

    def test_personal_scope_blocks_foreign_author_branch_type_and_privileged_reports(self):
        user = User.objects.create_user('next-employee')
        user.groups.add(Group.objects.get(name=POLITICIAN))
        MotivationReportAccess.objects.create(user=user, subdivision=self.fixture.subs[0], deal_author='Own Author')
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse('motivation_report_next', args=['registry'])).status_code, 403)
        self.assertEqual(self.client.get(reverse('motivation_report_next', args=['executive'])).status_code, 403)
        own = dict(detail_rows()[1], Подразделение='A', АвторСделки='Own Author')
        data = manager_detail_report([own])
        with patch('admin_motivation.services.reports_next.current_run', return_value=nullcontext()), patch('admin_motivation.services.reports_next.load_report', return_value=data) as load:
            response = self.client.get(reverse('motivation_report_next', args=['mine']), {'run': '1', 'author': 'Foreign', 'subdivision': 'B', 'type': 'Политики'})
        self.assertEqual(response.status_code, 200)
        values = load.call_args.args[1]
        self.assertEqual(values['DealAuthor'], ['Own Author'])
        self.assertEqual(values['Subdivision'], ['A'])
        self.assertEqual(self.client.get(reverse('motivation_report_next', args=['mine']), {'period': '1'}).status_code, 403)

    def test_summary_query_does_not_load_coefficients_tree_or_department_inputs(self):
        from .services.reports import load_report
        row = detail_rows()[1]
        fields = ['Подразделение', 'АвторСделки', 'Год', 'КварталМотивации',
                  'СуммаМотивацииUSD', 'СуммаМотивацииUSDУчетПроцента', 'Остаток',
                  'СуммаМотивацииRUR', 'СуммаМотивацииRURУчетПроцента', 'ОстатокRUR']
        cursor, connection = Mock(), Mock()
        cursor.description = [(field,) for field in fields]
        cursor.fetchmany.side_effect = [[tuple(row[field] for field in fields)], []]
        values = {'TypeCoeff_': 'Политики', 'Subdivision': ['A'], 'DealAuthor': [],
                  'gp_sales': [], 'GPGoods': [], 'Brend': [], 'Years': ['2026'], 'Currency': 'USD'}
        with patch('admin_motivation.services.reports.connect_database', return_value=(connection, cursor)):
            data = load_report('manager-detail', values, summary_only=True, premium_period='5')
        self.assertEqual(cursor.execute.call_count, 1)
        self.assertNotIn('K0WeightedSum', cursor.execute.call_args.args[0])
        self.assertEqual(data['period_summary']['totals']['total'], 95)
        self.assertNotIn('nodes', data)
        connection.close.assert_called_once()
        self.assertEqual(reports_next.period_code(None), '')

    def test_ledger_pagination_does_not_change_total_or_excel(self):
        rows = [dict(detail_rows()[1], АвторСделки=f'Employee {index}') for index in range(125)]
        data = manager_detail_report(rows)
        url = reverse('motivation_report_next', args=['registry'])
        with patch('admin_motivation.services.reports_next.current_run', return_value=nullcontext()), patch('admin_motivation.services.reports_next.load_report', return_value=data):
            response = self.client.get(url, {'run': '1', 'type': 'Политики', 'page': '2'})
            self.assertEqual(len(response.context['ledger_page']), 25)
            self.assertEqual(response.context['ledger']['total'], 125 * 95)
            exported = self.client.get(url, {'run': '1', 'type': 'Политики', 'page': '2', 'export': 'xlsx'})
            from io import BytesIO
            book = load_workbook(BytesIO(b''.join(exported.streaming_content)))
            self.assertEqual(book['Премии'].max_row, 127)
            self.assertEqual(sum(row[6].value for row in book['Премии'].iter_rows(min_row=3)), 125 * 95)
            book.close(); exported.close()

    def test_admin_can_preview_mine_but_must_choose_employee(self):
        user = User.objects.create_user('preview-employee')
        user.groups.add(Group.objects.get(name=POLITICIAN))
        MotivationReportAccess.objects.create(user=user, subdivision=self.fixture.subs[0], deal_author='Own Author')
        url = reverse('motivation_report_next', args=['mine'])
        self.assertContains(self.client.get(url), 'Сотрудник для предпросмотра')
        with patch('admin_motivation.services.reports_next.load_report') as load:
            response = self.client.get(url, {'run': '1'})
            self.assertContains(response, 'Выберите сотрудника для предпросмотра')
            load.assert_not_called()
        own = dict(detail_rows()[1], Подразделение='A', АвторСделки='Own Author')
        with patch('admin_motivation.services.reports_next.current_run', return_value=nullcontext()), patch('admin_motivation.services.reports_next.load_report', return_value=manager_detail_report([own])) as load:
            response = self.client.get(url, {'run': '1', 'preview_user': user.pk, 'subdivision': 'B', 'author': 'Foreign', 'type': 'Политики'})
        self.assertContains(response, 'Own Author')
        self.assertEqual(load.call_args.args[1]['Subdivision'], ['A'])
        self.assertEqual(load.call_args.args[1]['DealAuthor'], ['Own Author'])
        self.assertIn('preview_user=', response.context['ledger']['rows'][0]['detail_url'])
        self.assertEqual(response.context['preview_user'].pk, user.pk)

    def test_employee_cannot_use_admin_preview_to_change_identity(self):
        user = User.objects.create_user('own-preview-employee')
        user.groups.add(Group.objects.get(name=POLITICIAN))
        MotivationReportAccess.objects.create(user=user, subdivision=self.fixture.subs[0], deal_author='Own Author')
        self.client.force_login(user)
        own = dict(detail_rows()[1], Подразделение='A', АвторСделки='Own Author')
        with patch('admin_motivation.services.reports_next.current_run', return_value=nullcontext()), patch('admin_motivation.services.reports_next.load_report', return_value=manager_detail_report([own])) as load:
            response = self.client.get(reverse('motivation_report_next', args=['mine']), {'run': '1', 'preview_user': self.actor.pk})
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Сотрудник для предпросмотра')
        self.assertEqual(load.call_args.args[1]['DealAuthor'], ['Own Author'])

    def test_presentation_groups_people_without_changing_calculation_or_export_rows(self):
        from .services.reports_next_presenter import present
        ledger = {'total': Decimal(15), 'rows': [
            {'subdivision': 'A', 'author': 'Own', 'unallocated': False, 'kind': 'Политики', 'code': '4', 'period': '4 кв.', 'amount': Decimal(10), 'bonus': Decimal(0), 'total': Decimal(10), 'detail_url': '/detail/?period=4'},
            {'subdivision': 'A', 'author': 'Own', 'unallocated': False, 'kind': 'Продажи', 'code': '5', 'period': 'Год', 'amount': Decimal(3), 'bonus': Decimal(2), 'total': Decimal(5), 'detail_url': '/detail/?period=5'}]}
        display = present('team', ledger, ledger)
        self.assertEqual(display['people'], 1)
        self.assertEqual(display['total'], Decimal(15))
        self.assertEqual(display['employees'][0]['total'], Decimal(15))
        self.assertEqual(len(ledger['rows']), 2)
        self.assertEqual(sum(row['total'] for row in display['period_cards']), Decimal(15))

    def test_budget_shares_are_percentage_fields_from_calculation_job(self):
        query = QueryDict('period=4')
        self.assertEqual(self.run.parameters['quarter_share'], '80')
        result = reports_next.planned_budget(self.run, self.fixture.scenario, query)
        self.assertEqual(result[0]['global'], Decimal(180))
        self.run.parameters['year_share'] = '30'
        with self.assertRaises(ValidationError):
            reports_next.planned_budget(self.run, self.fixture.scenario, query)

    def test_plan_uses_active_branch_bindings_and_excludes_archived_directors_from_both_budgets(self):
        from datetime import date
        from .models import DirectorEmployee, SubdivisionDirectorCoeff, GlobalDirectorCoeff
        from .services.versions import create_version
        from .services.approval import review_approval, approve_review
        from .services.workflow import export_snapshot
        draft = create_version(date(2026, 1, 1), source_version=self.fixture.version.pk)
        active = DirectorEmployee.objects.create(full_name='Active Moscow director')
        moved = DirectorEmployee.objects.create(full_name='Removed Moscow director')
        archived = DirectorEmployee.objects.create(full_name='Archived director')
        for person in (active, moved, archived):
            person.subdivisions.add(self.fixture.subs[0])
        moved.subdivisions.add(self.fixture.subs[1])
        GlobalDirectorCoeff.objects.create(version=draft, goods=self.fixture.good, **{f'k{i}': '.01' for i in range(5)})
        for person, branch, rate in [(active, self.fixture.subs[0], '.02'), (moved, self.fixture.subs[0], '.03'),
                                     (archived, self.fixture.subs[0], '.04'), (moved, self.fixture.subs[1], '.03')]:
            SubdivisionDirectorCoeff.objects.create(version=draft, goods=self.fixture.good, employee=person,
                subdivision=branch, **{f'k{i}': rate for i in range(5)})
        approve_review(review_approval(draft)['token'])
        self.run.export_snapshot = export_snapshot()
        self.run.save()
        moved.subdivisions.remove(self.fixture.subs[0])
        archived.is_archived = True
        archived.save(update_fields=['is_archived'])
        query = QueryDict('period=all&type=Директора')
        nodes = reports_next.planned_budget(self.run, self.fixture.scenario, query)
        self.assertEqual(nodes[0]['global'], Decimal(15))
        self.assertEqual(nodes[0]['local'], Decimal(35))
        moscow = next(node for node in nodes if node['name'] == 'A')
        self.assertEqual([child['name'] for child in moscow['children'][0]['children']], [active.full_name])
        other = next(node for node in nodes if node['name'] == 'B')
        self.assertEqual(other['children'][0]['children'][0]['name'], moved.full_name)
        historical = reports_next.planned_budget(self.run, self.fixture.scenario, QueryDict('period=all&type=Директора&director_scope=coefficients'))
        self.assertEqual(historical[0]['global'], Decimal(35))
        self.assertEqual(historical[0]['local'], Decimal(105))
        self.assertEqual(SubdivisionDirectorCoeff.objects.filter(version=draft).count(), 4)

    def test_personal_goals_products_and_unavailable_reward_at_100_are_explicit(self):
        from .services.reports_next_presenter import personal_details
        from .services.pilot_formula import GLOBAL_SUBDIVISION
        rows = [dict(row, VariationCalculate=GLOBAL_SUBDIVISION,
                     КварталМотивации='4 Квартал(1-4 кв)' if row['КварталМотивации'] != 'Год' else 'Год') for row in detail_rows()]
        performance = [dict(row, СуммаПланUSD=1000, СуммаОтгрузокUSD=600) for row in rows]
        data = manager_detail_report(rows, performance_rows=performance)
        ledger = {'currency': 'USD', 'sections': [{'kind': 'Политики', 'data': data}]}
        result = personal_details(ledger, QueryDict('year=2026&period=all'))
        self.assertEqual(len(result['goals']), 2)
        for goal in result['goals']:
            self.assertEqual(goal['remaining'], Decimal(400))
            self.assertEqual(goal['ratio'], Decimal('.6'))
            self.assertIsNone(goal['reward_at_target'])
            self.assertIn('будущих личных отгрузок', goal['reward_reason'])
            self.assertEqual(goal['possible'], 100)
        leaf = result['product_sections'][0]['groups'][0]['children'][0]['children'][0]['children'][0]
        self.assertEqual(leaf['label'], '<Brand>')
        self.assertIn('planning=1.+RAC', leaf['detail_url'])
        self.assertIn('brand=%3CBrand%3E', leaf['detail_url'])
        self.assertTrue(leaf['detail_url'].endswith('#next-deep-calculation'))
        self.assertIn('start=2026-01-01', leaf['deals_url'])
        self.assertEqual(leaf['values'][-1], Decimal(190))

    def test_mine_excel_includes_product_premiums_and_goal_explanation(self):
        user = User.objects.create_user('goal-employee')
        user.groups.add(Group.objects.get(name=POLITICIAN))
        MotivationReportAccess.objects.create(user=user, subdivision=self.fixture.subs[0], deal_author='Иванов')
        from .services.pilot_formula import GLOBAL_SUBDIVISION
        rows = [dict(row, VariationCalculate=GLOBAL_SUBDIVISION) for row in detail_rows()]
        data = manager_detail_report(rows, performance_rows=rows)
        with patch('admin_motivation.services.reports_next.current_run', return_value=nullcontext()), patch('admin_motivation.services.reports_next.load_report', return_value=data):
            response = self.client.get(reverse('motivation_report_next', args=['mine']), {'run': '1', 'preview_user': user.pk, 'export': 'xlsx'})
        from io import BytesIO
        book = load_workbook(BytesIO(b''.join(response.streaming_content)))
        self.assertIn('Цели и потенциал', book.sheetnames)
        self.assertIn('Товарная детализация', book.sheetnames)
        self.assertIsNone(book['Цели и потенциал']['H2'].value)
        self.assertIn('будущих личных отгрузок', book['Цели и потенциал']['I2'].value)
        self.assertEqual(sum(row[7].value for row in book['Товарная детализация'].iter_rows(min_row=2)), 190)
        book.close(); response.close()

    def test_employee_period_columns_and_type_totals_reconcile(self):
        from .services.reports_next_presenter import present
        rows = []
        for kind, code, amount, bonus in [('Политики', '4', 10, 0), ('Политики', '5', 20, 5), ('Продажи', '4', 7, 0)]:
            rows.append({'subdivision': 'A', 'author': 'Own', 'unallocated': False, 'kind': kind, 'code': code,
                'period': 'Год' if code == '5' else '4 кв.', 'amount': Decimal(amount), 'bonus': Decimal(bonus),
                'total': Decimal(amount + bonus), 'detail_url': '/detail/?type=' + kind + '&period=' + code})
        ledger = {'kinds': ['Политики', 'Продажи'], 'rows': rows, 'total': Decimal(42)}
        data = present('team', ledger, ledger)
        self.assertEqual([column['label'] for column in data['period_columns']], ['4 кв.', 'Год'])
        self.assertEqual(len(data['employees']), 1)
        person = data['employees'][0]
        self.assertEqual(person['period_values'], [Decimal(17), Decimal(20)])
        self.assertEqual(sum(person['period_values']) + person['bonus'], person['total'])
        self.assertEqual([item['total'] for item in data['type_totals']], [Decimal(35), Decimal(7)])
        self.assertEqual(len(data['registry_rows']), 2)
        self.assertEqual(sum(item['total'] for item in data['registry_rows']), ledger['total'])
