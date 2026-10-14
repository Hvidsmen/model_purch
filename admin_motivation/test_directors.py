from datetime import date
from decimal import Decimal
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from base.testing import authorize_test_case
from .models import (GlobalCoeffVersion, Goods, Chanel, Subdivision, DirectorEmployee,
                     GlobalDirectorCoeff, SubdivisionDirectorCoeff)
from .services.directors import copy_directors
from .services.versions import create_version, delete_draft
from .services.approval import review_approval, approve_review
from .director_forms import GlobalDirectorAddForm


class DirectorCoefficientTests(TestCase):
    def test_branch_selects_first_director_by_default_and_keeps_explicit_choice(self):
        url = reverse('coeff_subdivisions_admin_motivation', args=[self.sub.pk])
        response = self.client.get(url)
        self.assertEqual(response.context['employee'], response.context['director_employees'].first())
        response = self.client.get(url, {'employee': self.second.pk})
        self.assertEqual(response.context['employee'], self.second)
        self.second.subdivisions.remove(self.sub)
        self.assertEqual(self.client.get(url).context['employee'], self.first)
        self.first.subdivisions.remove(self.sub)
        self.assertIsNone(self.client.get(url).context['employee'])

    def test_admin_copies_approved_globals_to_branch_with_audit(self):
        from django.db import models
        from django.contrib.auth import get_user_model
        from .models import CoefficientGroupCopyLog, GlobalCoeff, SubdivisionCoeff, TypeCoeff, SegmentCoeff, VariationCalculate
        from .services.versions import apply_to_subdivisions
        kind = TypeCoeff.objects.create(type_coeff_name='Политики')
        segment = SegmentCoeff.objects.create(segment_name='K0')
        variation = VariationCalculate.objects.create(variation_name='Copy history level')
        GlobalCoeff.objects.create(version=self.version, goods=self.goods, type_coeff=kind, segment=segment,
            motivation_coeff=.02, manager_coeff=1, variation_calculate=variation)
        SubdivisionDirectorCoeff.objects.create(version=self.version, goods=self.goods, subdivision=self.sub, employee=self.first, k0='.9')
        models.QuerySet(model=GlobalCoeffVersion).filter(pk=self.version.pk).update(status='approved')
        self.version.refresh_from_db()
        user = get_user_model().objects.get(pk=self.client.session['_auth_user_id'])
        with self.assertRaises(ValidationError):
            apply_to_subdivisions(self.version, [self.sub], groups=['directors'], employee=self.first)
        apply_to_subdivisions(self.version, [self.sub], groups=['policies', 'directors'], employee=self.first, user=user)
        self.assertEqual(SubdivisionCoeff.objects.get(version=self.version, subdivision=self.sub).motivation_coeff, .02)
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(employee=self.first).k0, Decimal('.01'))
        self.assertEqual(CoefficientGroupCopyLog.objects.filter(version=self.version, subdivision=self.sub).count(), 1)
        self.assertIn('0.900000', str(CoefficientGroupCopyLog.objects.get().changes))
        self.version.refresh_from_db()
        self.assertEqual(self.version.status, 'approved')
        response = self.client.post(reverse('gb_act'), {'version': self.version.pk, 'action_button': 'apply_sub',
            'copy_groups': ['directors'], 'selected_subdivisions': [self.sub.pk], 'director_employee': self.second.pk})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(SubdivisionDirectorCoeff.objects.filter(employee=self.second).exists())
        response = self.client.post(reverse('sub_act', args=[self.sub.pk]), {'version': self.version.pk,
            'action_button': 'copy_from_global', 'copy_groups': ['directors'], 'director_employee': self.first.pk})
        self.assertEqual(response.status_code, 302)

    def test_global_copy_skips_branches_without_directors(self):
        from .services.versions import apply_to_subdivisions
        skipped = apply_to_subdivisions(self.version, [self.sub, self.other_sub], groups=['directors'],
            employees_by_subdivision={self.sub.pk: [self.first, self.second], self.other_sub.pk: []})
        self.assertEqual(skipped, [str(self.other_sub)])
        self.assertEqual(SubdivisionDirectorCoeff.objects.filter(subdivision=self.sub).count(), 2)
        self.assertFalse(SubdivisionDirectorCoeff.objects.filter(subdivision=self.other_sub).exists())
        with self.assertRaises(ValidationError):
            apply_to_subdivisions(self.version, [self.sub], groups=['directors'], employees_by_subdivision={self.sub.pk: []})
        with self.assertRaises(ValidationError):
            apply_to_subdivisions(self.version, [self.other_sub], groups=['directors'], employees_by_subdivision={self.other_sub.pk: []})

    def test_global_copy_multiple_employees_per_branch_and_atomic_validation(self):
        from .services.versions import apply_to_subdivisions
        self.second.subdivisions.add(self.other_sub)
        third = DirectorEmployee.objects.create(full_name='Не выбран')
        third.subdivisions.add(self.sub)
        preserved = SubdivisionDirectorCoeff.objects.create(version=self.version, subdivision=self.sub,
            employee=third, goods=self.goods, k0=Decimal('.9'))
        recipients = {self.sub.pk: [self.first, self.second], self.other_sub.pk: [self.second]}
        apply_to_subdivisions(self.version, [self.sub, self.other_sub], groups=['directors'], employees_by_subdivision=recipients)
        self.assertEqual(SubdivisionDirectorCoeff.objects.filter(version=self.version, employee=self.second).count(), 2)
        preserved.refresh_from_db()
        self.assertEqual(preserved.k0, Decimal('.9'))
        self.global_row.k0 = Decimal('.7')
        self.global_row.save()
        with self.assertRaises(ValidationError):
            apply_to_subdivisions(self.version, [self.sub, self.other_sub], groups=['directors'],
                employees_by_subdivision={self.sub.pk: [self.first], self.other_sub.pk: [self.first]})
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(employee=self.first).k0, Decimal('.01'))
        response = self.client.post(reverse('gb_act'), {'version': self.version.pk, 'action_button': 'apply_sub',
            'copy_group_selection': 'yes', 'copy_groups': ['directors'], 'branch_employee_selection': 'yes',
            'selected_subdivisions': [self.sub.pk, self.other_sub.pk],
            f'director_employees_{self.sub.pk}': [self.first.pk, self.second.pk],
            f'director_employees_{self.other_sub.pk}': [self.second.pk]})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(employee=self.first).k0, Decimal('.7'))
        page = self.client.get(reverse('global_coeff_admin_motivation'), {'version': self.version.pk})
        self.assertContains(page, 'copy-branch-search')
        self.assertContains(page, f'name="director_employees_{self.sub.pk}"')

    def test_admin_group_copy_approved_audited_and_guarded(self):
        from django.db import models
        from django.contrib.auth import get_user_model
        from django.contrib.auth.models import Group
        from .models import GlobalCoeff, TypeCoeff, SegmentCoeff, VariationCalculate, CoefficientGroupCopyLog
        from .services.group_copy import copy_group
        from django.core.exceptions import PermissionDenied
        policies = TypeCoeff.objects.create(type_coeff_name='Политики')
        sales = TypeCoeff.objects.create(type_coeff_name='Продажи')
        variation = VariationCalculate.objects.create(variation_name='Группа планов продаж')
        for i in range(5):
            segment = SegmentCoeff.objects.create(segment_name=f'K{i}')
            for kind in (policies, sales):
                GlobalCoeff.objects.create(version=self.version, goods=self.goods, type_coeff=kind,
                    segment=segment, motivation_coeff=(0.0015999999999999999 if kind == sales and i == 0 else (i + 1) / (50 if kind == sales else 100)), manager_coeff=.5, variation_calculate=variation)
        models.QuerySet(model=GlobalCoeffVersion).filter(pk=self.version.pk).update(status='approved')
        self.version.refresh_from_db()
        user = get_user_model().objects.get(pk=self.client.session['_auth_user_id'])
        count = copy_group(user, self.version.pk, self.version.revision, 'sales', 'directors')
        self.assertEqual(count, 1)
        self.global_row.refresh_from_db()
        self.assertEqual(self.global_row.k4, Decimal('.10'))
        self.assertEqual(self.global_row.k0, Decimal('0.001600'))
        self.assertEqual(CoefficientGroupCopyLog.objects.count(), 1)
        log = CoefficientGroupCopyLog.objects.get()
        self.assertEqual(log.changes[0]['before']['k0'], '0.010000')
        self.assertEqual(log.changes[0]['after']['k0'], '0.001600')
        with self.assertRaises(ValidationError):
            copy_group(user, self.version.pk, self.version.revision, 'sales', 'directors')
        self.version.refresh_from_db()
        copy_group(user, self.version.pk, self.version.revision, 'directors', 'policies')
        updated = GlobalCoeff.objects.get(version=self.version, goods=self.goods, type_coeff=policies, segment__segment_name='K4')
        self.assertEqual(updated.manager_coeff, .5)
        self.assertEqual(updated.motivation_coeff, .10)
        self.version.refresh_from_db()
        copy_group(user, self.version.pk, self.version.revision, 'policies', 'sales')
        director_user = get_user_model().objects.create_user('role-director')
        director_user.groups.add(Group.objects.get_or_create(name='Генеральный директор')[0])
        self.version.refresh_from_db()
        with self.assertRaises(PermissionDenied):
            copy_group(director_user, self.version.pk, self.version.revision, 'sales', 'directors')
        response = self.client.get(reverse('motivation_copy_group'), {'version': self.version.pk})
        self.assertContains(response, 'Изменяется историческая версия')

    def test_directory_name_and_branch_filters(self):
        url = reverse('motivation_director_directory')
        response = self.client.get(url, {'name': 'ПЕРВЫЙ'})
        self.assertEqual([item.pk for item in response.context['all_employees']], [self.first.pk])
        self.first.subdivisions.add(self.other_sub)
        response = self.client.get(url, {'branch': self.other_sub.pk, 'name': 'директор'})
        self.assertEqual([item.pk for item in response.context['all_employees']], [self.first.pk])
        response = self.client.get(url, {'branch': self.other_sub.pk, 'name': 'Второй'})
        self.assertContains(response, 'Директоры не найдены')
        self.assertEqual(len(self.client.get(url).context['all_employees']), 2)

    def test_director_position_form_and_excel(self):
        from io import BytesIO
        from openpyxl import Workbook, load_workbook
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .services.director_import import import_directors, template
        url = reverse('motivation_director_directory')
        response = self.client.post(url, {'edit_employee': self.first.pk,
            'full_name': self.first.full_name, 'position': 'Директор филиала', 'subdivisions': [self.sub.pk]})
        self.assertRedirects(response, url)
        self.first.refresh_from_db()
        self.assertEqual(self.first.position, 'Директор филиала')
        self.assertContains(self.client.get(url), 'Директор филиала')
        def upload(with_position):
            book = Workbook()
            book.active.append(['ФИО', 'Должность', 'Филиалы'] if with_position else ['ФИО', 'Филиалы'])
            book.active.append([self.first.full_name, 'Региональный директор', 'A'] if with_position else [self.first.full_name, 'A'])
            output = BytesIO()
            book.save(output)
            return SimpleUploadedFile('directors.xlsx', output.getvalue())
        import_directors(upload(True))
        import_directors(upload(False))
        self.first.refresh_from_db()
        self.assertEqual(self.first.position, 'Региональный директор')
        book = load_workbook(BytesIO(template()))
        self.assertIn('Должность', [cell.value for cell in book.active[1]])
        book.close()

    def test_excel_import_atomic_update_and_auto_code(self):
        from io import BytesIO
        from openpyxl import Workbook
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .services.director_import import import_directors
        def upload(rows):
            book = Workbook()
            book.active.append(['ФИО', 'Филиалы'])
            for row in rows:
                book.active.append(row)
            output = BytesIO()
            book.save(output)
            return SimpleUploadedFile('directors.xlsx', output.getvalue())
        self.assertEqual(import_directors(upload([['Импорт директор', 'A;B']])), 1)
        employee = DirectorEmployee.objects.get(full_name='Импорт директор')
        code = employee.employee_code
        self.assertTrue(code.startswith('DIR-'))
        self.assertEqual(employee.subdivisions.count(), 2)
        import_directors(upload([['Импорт директор', 'B']]))
        employee.refresh_from_db()
        self.assertEqual(employee.employee_code, code)
        self.assertEqual(list(employee.subdivisions.values_list('pk', flat=True)), [self.other_sub.pk])
        with self.assertRaises(ValidationError):
            import_directors(upload([['Не сохранять', 'A'], ['Ошибка', 'unknown']]))
        self.assertFalse(DirectorEmployee.objects.filter(full_name='Не сохранять').exists())
        url = reverse('motivation_director_directory')
        self.assertEqual(self.client.get(url + '?template=1')['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response = self.client.post(url, {'action': 'import', 'file': upload([['Из формы', 'A']])})
        self.assertRedirects(response, url)
        self.assertTrue(DirectorEmployee.objects.filter(full_name='Из формы').exists())

    def test_standalone_directory_create_edit_and_validation(self):
        url = reverse('motivation_director_directory')
        response = self.client.get(url)
        self.assertContains(response, 'Справочник директоров')
        self.assertContains(response, 'Первый директор')
        response = self.client.post(url, {'full_name': 'Новый директор', 'employee_code': 'D3', 'subdivisions': [self.sub.pk, self.other_sub.pk]})
        self.assertRedirects(response, url)
        employee = DirectorEmployee.objects.get(full_name='Новый директор')
        self.assertEqual(employee.subdivisions.count(), 2)
        response = self.client.post(url, {'edit_employee': employee.pk, 'full_name': 'Обновлён', 'employee_code': 'D3', 'subdivisions': [self.sub.pk]})
        self.assertRedirects(response, url)
        employee.refresh_from_db()
        self.assertEqual(employee.full_name, 'Обновлён')
        self.assertEqual(employee.subdivisions.count(), 1)
        response = self.client.post(url, {'full_name': ''})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['employee_form'].errors)
        self.assertTrue(employee.employee_code.startswith('DIR-'))

    def setUp(self):
        authorize_test_case(self)
        self.version = GlobalCoeffVersion.objects.get(effective_from=date(2001, 1, 1))
        channel = Chanel.objects.create(chanel_name='Channel')
        self.sub = Subdivision.objects.create(subdivision_key='A', subdivision_name='A', subdivision_global='A', chanel=channel)
        self.other_sub = Subdivision.objects.create(subdivision_key='B', subdivision_name='B', subdivision_global='B', chanel=channel)
        self.first = DirectorEmployee.objects.create(full_name='Первый директор', employee_code='D1')
        self.second = DirectorEmployee.objects.create(full_name='Второй директор', employee_code='D2')
        self.first.subdivisions.add(self.sub)
        self.second.subdivisions.add(self.sub)
        self.goods = Goods.objects.create(goods_key='Test', planning_group_sales='1. RAC', group='Group', brand='Brand')
        self.global_row = GlobalDirectorCoeff.objects.create(version=self.version, goods=self.goods,
            k0=Decimal('.01'), k1=Decimal('.02'), k2=Decimal('.03'), k3=Decimal('.04'), k4=Decimal('.05'))
        self.global_url = reverse('motivation_directors_global')
        self.local_url = reverse('motivation_directors_local')

    def local_post(self, action, employee=None, **data):
        return self.client.post(self.local_url, {'version': self.version.pk, 'subdivision': self.sub.pk,
            'employee': (employee or self.first).pk, 'action': action, **data})

    def test_percent_form_and_single_coefficient_set(self):
        form = GlobalDirectorAddForm({'goods': self.goods.pk, **{f'k{i}': '12.3456' for i in range(5)}},
            instance=GlobalDirectorCoeff(version=self.version))
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data['k0'], Decimal('.123456'))
        self.assertEqual(GlobalDirectorAddForm(instance=self.global_row).initial['k1'], Decimal('2'))
        bad = GlobalDirectorAddForm({'goods': self.goods.pk, **{f'k{i}': '101' for i in range(5)}},
            instance=GlobalDirectorCoeff(version=self.version))
        self.assertFalse(bad.is_valid())

    def test_global_and_local_forms_render(self):
        global_page = self.client.get(self.global_url, follow=True)
        self.assertContains(global_page, 'Директор K0')
        self.assertContains(global_page, f'director_coeff={self.goods.pk}=0')
        page = self.client.get(self.local_url, {'subdivision': self.sub.pk}, follow=True)
        self.assertContains(page, 'Первый директор')
        self.assertContains(page, 'Второй директор')
        self.assertContains(page, 'id="director-employee-choice"')

    def test_each_employee_has_independent_coefficients(self):
        self.local_post('copy_global', confirm_copy='yes')
        self.local_post('copy_global', self.second, confirm_copy='yes')
        first = SubdivisionDirectorCoeff.objects.get(employee=self.first)
        first.k0 = Decimal('.5')
        first.save()
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(employee=self.second).k0, Decimal('.01'))
        self.assertEqual(GlobalDirectorCoeff.objects.get().k0, Decimal('.01'))
        self.assertEqual(self.first.subdivisions.count(), 1)
        self.assertEqual(self.sub.directors.count(), 2)

    def test_copy_requires_confirmation_and_assigned_employee(self):
        self.local_post('copy_global')
        self.assertFalse(SubdivisionDirectorCoeff.objects.exists())
        with self.assertRaises(ValidationError):
            copy_directors(self.version.pk, subdivision=self.other_sub, employee=self.first, from_global=True)
        self.assertFalse(SubdivisionDirectorCoeff.objects.exists())

    def test_empty_source_does_not_remove_current_rows(self):
        self.local_post('copy_global', confirm_copy='yes')
        self.global_row.delete()
        with self.assertRaises(ValidationError):
            copy_directors(self.version.pk, subdivision=self.sub, employee=self.first, from_global=True)
        self.assertEqual(SubdivisionDirectorCoeff.objects.count(), 1)

    def test_new_version_clones_director_rows_and_old_version_is_frozen(self):
        self.local_post('copy_global', confirm_copy='yes')
        next_version = create_version(date(2026, 1, 1), 'Next', self.version.pk)
        self.assertEqual(GlobalDirectorCoeff.objects.get(version=next_version).k4, Decimal('.05'))
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(version=next_version).employee, self.first)
        with self.assertRaises(ValidationError):
            self.global_row.k0 = Decimal('.9')
            self.global_row.save()
        with self.assertRaises(ValidationError):
            copy_directors(self.version.pk, next_version.pk)
        copied = GlobalDirectorCoeff.objects.get(version=next_version)
        copied.k0 = Decimal('.8')
        copied.save()
        copy_directors(next_version.pk, self.version.pk)
        self.assertEqual(GlobalDirectorCoeff.objects.get(version=next_version).k0, Decimal('.01'))
        copy_directors(next_version.pk, self.version.pk, self.sub, self.first)
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(version=next_version).k4, Decimal('.05'))

    def test_approved_director_version_is_immutable_and_draft_deletion_retains_history(self):
        review = review_approval(self.version)
        approve_review(review['token'])
        self.version.refresh_from_db()
        self.assertEqual(self.version.status, 'approved')
        with self.assertRaises(ValidationError):
            self.global_row.save()
        draft = create_version(date(2026, 1, 1), 'Next', self.version.pk)
        delete_draft(draft.pk)
        self.assertFalse(GlobalDirectorCoeff.objects.filter(version=draft).exists())
        self.assertEqual(GlobalDirectorCoeff.objects.filter(version=self.version).count(), 1)

    def test_partial_save_cannot_modify_other_employee(self):
        self.local_post('copy_global', confirm_copy='yes')
        self.local_post('copy_global', self.second, confirm_copy='yes')
        other = SubdivisionDirectorCoeff.objects.get(employee=self.second)
        response = self.local_post('save', **{'coeff-TOTAL_FORMS': '1', 'coeff-INITIAL_FORMS': '1',
            'coeff-MIN_NUM_FORMS': '0', 'coeff-MAX_NUM_FORMS': '1000', 'coeff-0-id': other.pk,
            **{f'coeff-0-k{i}': '90' for i in range(5)}})
        self.assertEqual(response.status_code, 302)
        other.refresh_from_db()
        self.assertEqual(other.k0, Decimal('.01'))

    def test_save_employee_and_multiple_branch_assignments(self):
        response = self.client.post(self.local_url, {'version': self.version.pk, 'action': 'employee_save',
            'full_name': 'Третий директор', 'employee_code': 'D3', 'subdivisions': [self.sub.pk, self.other_sub.pk]})
        self.assertEqual(response.status_code, 302)
        third = DirectorEmployee.objects.get(full_name='Третий директор')
        self.assertEqual(third.subdivisions.count(), 2)

    def test_new_version_returns_to_director_page(self):
        response = self.client.post(reverse('gb_act'), {'action_button': 'create_version',
            'source_version': self.version.pk, 'effective_from': '2026-01-01', 'title': 'Next',
            'return_directors': 'local', 'return_subdivision': self.sub.pk, 'return_employee': self.first.pk})
        latest = GlobalCoeffVersion.objects.latest('pk')
        self.assertRedirects(response, reverse('coeff_subdivisions_admin_motivation', args=[self.sub.pk]) + f'?employee={self.first.pk}&version={latest.pk}')

    def test_formset_saves_percent_values_and_preserves_other_scopes(self):
        self.local_post('copy_global', confirm_copy='yes')
        row = SubdivisionDirectorCoeff.objects.get(employee=self.first)
        values = {'coeff-TOTAL_FORMS': '1', 'coeff-INITIAL_FORMS': '1', 'coeff-MIN_NUM_FORMS': '0',
                  'coeff-MAX_NUM_FORMS': '1000', 'coeff-0-id': row.pk,
                  **{f'coeff-0-k{i}': '1.0000' for i in range(5)}}
        response = self.local_post('save', **values)
        self.assertEqual(response.status_code, 302)
        row.refresh_from_db()
        self.assertEqual(row.k4, Decimal('.01'))
        self.assertEqual(GlobalDirectorCoeff.objects.get().k4, Decimal('.05'))

    def test_logistician_cannot_open_director_forms(self):
        from base.access import LOGISTICIAN
        from django.contrib.auth.models import User, Group
        user = User.objects.create_user('director-test-logistician')
        user.groups.add(Group.objects.get(name=LOGISTICIAN))
        self.client.force_login(user)
        self.assertEqual(self.client.get(self.global_url).status_code, 403)
        self.assertEqual(self.client.get(self.local_url).status_code, 403)

    def test_seed_matrix_adds_missing_goods_without_overwriting_coefficients(self):
        other = Goods.objects.create(goods_key='Other', planning_group_sales='2. VRF', group='Other', brand='Brand')
        for _ in range(2):
            self.client.post(self.global_url, {'version': self.version.pk, 'action': 'seed_matrix'})
        self.assertEqual(GlobalDirectorCoeff.objects.filter(version=self.version).count(), 2)
        self.assertEqual(GlobalDirectorCoeff.objects.get(goods=other).k0, 0)
        self.assertEqual(GlobalDirectorCoeff.objects.get(goods=self.goods).k0, Decimal('.01'))

    def test_global_directors_save_in_shared_form_and_old_page_redirects(self):
        response = self.client.get(self.global_url)
        self.assertRedirects(response, reverse('global_coeff_admin_motivation'))
        response = self.client.post(reverse('gb_act'), {'version': self.version.pk, 'action_button': 'save_coeff',
            f'director_coeff={self.goods.pk}=0': '25%', f'director_coeff={self.goods.pk}=4': '10%'})
        self.assertEqual(response.status_code, 302)
        self.global_row.refresh_from_db()
        self.assertEqual(self.global_row.k0, Decimal('.25'))
        self.assertEqual(self.global_row.k4, Decimal('.10'))
        page = self.client.get(reverse('global_coeff_admin_motivation'))
        self.assertContains(page, 'data-coefficient-group-toggle="directors"')
        self.assertContains(page, f'director_coeff={self.goods.pk}=0')

    def test_selected_copy_preserves_other_groups_and_other_directors(self):
        from .models import TypeCoeff, SegmentCoeff, VariationCalculate, GlobalCoeff, SubdivisionCoeff
        policy = TypeCoeff.objects.create(type_coeff_name='Политики')
        sales = TypeCoeff.objects.create(type_coeff_name='Продажи')
        segment = SegmentCoeff.objects.create(segment_name='K0')
        variation = VariationCalculate.objects.create(variation_name='Подразделение')
        data = {'version': self.version, 'goods': self.goods, 'segment': segment,
                'manager_coeff': 1, 'variation_calculate': variation}
        GlobalCoeff.objects.create(**data, type_coeff=policy, motivation_coeff=.02)
        GlobalCoeff.objects.create(**data, type_coeff=sales, motivation_coeff=.03)
        local_policy = SubdivisionCoeff.objects.create(**data, subdivision=self.sub, type_coeff=policy, motivation_coeff=.05)
        local_sales = SubdivisionCoeff.objects.create(**data, subdivision=self.sub, type_coeff=sales, motivation_coeff=.07)
        for employee in (self.first, self.second):
            SubdivisionDirectorCoeff.objects.create(version=self.version, goods=self.goods, subdivision=self.sub, employee=employee, k0=Decimal('.4'))
        url = reverse('sub_act', args=[self.sub.pk])
        self.client.post(url, {'version': self.version.pk, 'action_button': 'copy_from_global', 'copy_groups': ['policies']})
        self.assertEqual(SubdivisionCoeff.objects.get(subdivision=self.sub, type_coeff=policy).motivation_coeff, .02)
        local_sales.refresh_from_db()
        self.assertEqual(local_sales.motivation_coeff, .07)
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(employee=self.first).k0, Decimal('.4'))
        self.client.post(url, {'version': self.version.pk, 'action_button': 'copy_from_global', 'copy_groups': ['sales']})
        self.assertEqual(SubdivisionCoeff.objects.get(subdivision=self.sub, type_coeff=sales).motivation_coeff, .03)
        self.client.post(url, {'version': self.version.pk, 'action_button': 'copy_from_global',
            'copy_groups': ['directors'], 'director_employee': self.first.pk})
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(employee=self.first).k0, Decimal('.01'))
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(employee=self.second).k0, Decimal('.4'))
        self.assertEqual(SubdivisionCoeff.objects.get(subdivision=self.sub, type_coeff=policy).motivation_coeff, .02)

    def test_global_copy_uses_saved_values_not_unsaved_director_inputs(self):
        self.client.post(reverse('gb_act'), {'version': self.version.pk, 'action_button': 'apply_sub',
            'selected_subdivisions': [self.sub.pk], 'copy_groups': ['directors'], 'director_employee': self.first.pk,
            f'director_coeff={self.goods.pk}=0': '90%'})
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(employee=self.first).k0, Decimal('.01'))
        self.global_row.refresh_from_db()
        self.assertEqual(self.global_row.k0, Decimal('.01'))

    def test_unified_local_form_saves_selected_director_only(self):
        self.local_post('copy_global', confirm_copy='yes')
        self.local_post('copy_global', self.second, confirm_copy='yes')
        url = reverse('sub_act', args=[self.sub.pk])
        response = self.client.post(url, {'version': self.version.pk, 'action_button': 'save_coeff',
            'employee': self.first.pk, f'director_coeff={self.goods.pk}=0': '30%'})
        self.assertRedirects(response, reverse('coeff_subdivisions_admin_motivation', args=[self.sub.pk]) + f'?version={self.version.pk}&employee={self.first.pk}')
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(employee=self.first).k0, Decimal('.3'))
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(employee=self.second).k0, Decimal('.01'))
        page = self.client.get(reverse('coeff_subdivisions_admin_motivation', args=[self.sub.pk]), {'employee': self.first.pk})
        self.assertContains(page, f'director_coeff={self.goods.pk}=0')
        self.assertContains(page, 'Справочник сотрудников-директоров')

    def test_copy_to_other_director_keeps_source_and_other_branch(self):
        self.local_post('copy_global', confirm_copy='yes')
        self.local_post('copy_global', self.second, confirm_copy='yes')
        source = SubdivisionDirectorCoeff.objects.get(employee=self.first)
        source.k0 = Decimal('.6')
        source.save()
        url = reverse('sub_act', args=[self.sub.pk])
        self.client.post(url, {'version': self.version.pk, 'action_button': 'copy_to_director',
            'employee': self.first.pk, 'copy_to_employee': self.second.pk})
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(employee=self.second).k0, Decimal('.01'))
        response = self.client.post(url, {'version': self.version.pk, 'action_button': 'copy_to_director',
            'employee': self.first.pk, 'copy_to_employee': self.second.pk, 'confirm_copy': 'yes'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(employee=self.second).k0, Decimal('.6'))
        source.refresh_from_db()
        self.assertEqual(source.k0, Decimal('.6'))

    def test_director_version_copy_preserves_other_employee(self):
        from .services.versions import copy_from_version
        self.local_post('copy_global', confirm_copy='yes')
        self.local_post('copy_global', self.second, confirm_copy='yes')
        draft = create_version(date(2026, 1, 1), 'Next', self.version.pk)
        first = SubdivisionDirectorCoeff.objects.get(version=draft, employee=self.first)
        first.k0 = Decimal('.8')
        first.save()
        second = SubdivisionDirectorCoeff.objects.get(version=draft, employee=self.second)
        second.k0 = Decimal('.9')
        second.save()
        copy_from_version(draft.pk, self.version.pk, self.sub, employee=self.first)
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(version=draft, employee=self.first).k0, Decimal('.01'))
        self.assertEqual(SubdivisionDirectorCoeff.objects.get(version=draft, employee=self.second).k0, Decimal('.9'))

    def test_shared_directory_adds_employee_and_opens_them(self):
        response = self.client.post(reverse('sub_act', args=[self.sub.pk]), {'version': self.version.pk,
            'action_button': 'employee_save', 'full_name': 'Новый директор', 'employee_code': 'NEW-DIR',
            'subdivisions': [self.sub.pk]})
        employee = DirectorEmployee.objects.get(full_name='Новый директор')
        self.assertRedirects(response, reverse('coeff_subdivisions_admin_motivation', args=[self.sub.pk]) + f'?version={self.version.pk}&employee={employee.pk}')
